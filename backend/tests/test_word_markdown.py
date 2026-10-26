"""Bloc 3 « Le Word dit vrai » — tranche 1 : le Markdown de l'IA devient des
styles Word natifs (titres, puces, numérotation, gras, italique, tableaux).

Avant : `docx_exporter._body` ne gérait que les puces et `**gras**` ; les
`###`, `*italique*` et les tableaux `| a | b |` + `|---|` sortaient bruts dans
le Word et dans le PDF.

Contenu entièrement inventé, au format de ce que produit l'IA.
"""
import io

from docx import Document

from services.docx_exporter import build_memoire_docx

TABLEAU_EFFECTIFS = """\
| Phase | Nb ouvriers | Durée |
|---|---|---|
| Installation | 4 | 1 semaine |
| Gros œuvre | **8** | 6 semaines |
| Finitions | 5 | 3 semaines |
"""

CONTENU = {
    "preambule": (
        "Nous nous engageons sur **trois points** pour ce marché fictif.\n"
        "### Nos engagements\n"
        "- Respect du *délai* contractuel\n"
        "- Chantier propre et **sécurisé**\n"
    ),
    "partie_a": {
        "activites": (
            "## Nos spécialités\n"
            "1. Maçonnerie générale\n"
            "2. Isolation par l'extérieur\n"
            "3. Ravalement de façades\n\n"
            "Texte de liaison avec une *nuance* et un **point fort**.\n"
        ),
    },
    "partie_b": {},
    "partie_c": {
        "effectifs": "Équipe dédiée au chantier :\n\n" + TABLEAU_EFFECTIFS + "\nFin de section.",
        "methodologie": (
            "# Phase de préparation\n"
            "Description de la préparation.\n"
            "### Contrôle des supports\n"
            "* Sondage des supports\n"
            "* Relevé des fissures\n"
            "1. Première étape\n2. Deuxième étape\n"
        ),
    },
}


def _doc() -> Document:
    return Document(io.BytesIO(build_memoire_docx(CONTENU, "Chantier fictif", "ENTREPRISE TEST")))


def _tout_le_texte(doc: Document) -> str:
    morceaux = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for ligne in t.rows:
            morceaux.extend(c.text for c in ligne.cells)
    return "\n".join(morceaux)


def test_aucun_reste_de_syntaxe_markdown():
    texte = _tout_le_texte(_doc())
    for reste in ("###", "## ", "# Phase", "**", "|---|", "|---", "| Phase"):
        assert reste not in texte, f"{reste!r} imprimé tel quel"
    assert "*délai*" not in texte and "*nuance*" not in texte


def test_tableau_markdown_devient_un_vrai_tableau_word():
    doc = _doc()
    tableaux = [t for t in doc.tables if t.rows[0].cells[0].text == "Phase"]
    assert len(tableaux) == 1
    t = tableaux[0]
    assert len(t.rows) == 4 and len(t.columns) == 3
    assert [c.text for c in t.rows[0].cells] == ["Phase", "Nb ouvriers", "Durée"]
    assert [c.text for c in t.rows[2].cells] == ["Gros œuvre", "8", "6 semaines"]
    xml = t._tbl.xml
    assert "w:tblHeader" in xml          # en-tête répété sur chaque page
    assert xml.count("w:cantSplit") == 4  # lignes non coupées
    # le gras d'une cellule est un vrai gras
    assert any(r.bold for r in t.rows[2].cells[1].paragraphs[0].runs)


def test_titres_markdown_en_styles_natifs():
    doc = _doc()
    styles = {p.text: p.style.name for p in doc.paragraphs}
    assert styles["Nos engagements"].startswith("Heading")
    assert styles["Nos spécialités"].startswith("Heading")
    assert styles["Phase de préparation"].startswith("Heading")
    assert styles["Contrôle des supports"].startswith("Heading")
    # sous le titre de section (Heading 2) : toujours plus bas dans la hiérarchie
    niveau = lambda nom: int(styles[nom].split()[-1])  # noqa: E731
    assert niveau("Nos spécialités") > 2 and niveau("Phase de préparation") > 2
    assert niveau("Contrôle des supports") > niveau("Phase de préparation")


def test_listes_puces_et_numerotees():
    doc = _doc()
    par_texte = {p.text: p for p in doc.paragraphs}
    assert par_texte["Respect du délai contractuel"].style.name == "List Bullet"
    assert par_texte["Sondage des supports"].style.name == "List Bullet"
    assert par_texte["Maçonnerie générale"].style.name == "List Number"
    assert par_texte["Première étape"].style.name == "List Number"


def test_numerotation_repart_a_un_pour_chaque_liste():
    doc = _doc()
    par_texte = {p.text: p for p in doc.paragraphs}

    def num_id(p):
        return p._p.pPr.numPr.numId.val if p._p.pPr is not None and p._p.pPr.numPr is not None else None

    # deux listes numérotées distinctes → deux instances de numérotation distinctes
    assert num_id(par_texte["Maçonnerie générale"]) is not None
    assert num_id(par_texte["Maçonnerie générale"]) == num_id(par_texte["Ravalement de façades"])
    assert num_id(par_texte["Maçonnerie générale"]) != num_id(par_texte["Première étape"])


def test_gras_et_italique_sont_de_vrais_runs():
    doc = _doc()
    par_texte = {p.text: p for p in doc.paragraphs}
    p = par_texte["Texte de liaison avec une nuance et un point fort."]
    gras = [r.text for r in p.runs if r.bold]
    ital = [r.text for r in p.runs if r.italic]
    assert gras == ["point fort"]
    assert ital == ["nuance"]


def test_barres_isolees_et_etoiles_non_appariees_ne_laissent_rien():
    contenu = {"preambule": "Un **gras jamais fermé et une étoile * seule, 2 * 3.\nLigne | avec barre.", "partie_a": {}}
    texte = _tout_le_texte(Document(io.BytesIO(build_memoire_docx(contenu, "P", "E"))))
    assert "**" not in texte
    assert "Ligne | avec barre." in texte  # une barre sans ligne de séparation n'est pas un tableau
    assert "2 * 3." in texte
