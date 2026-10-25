"""B3 — script de ré-extraction des .docx déjà en base (backend/scripts/reextraire_docx.py).

Pas de ré-extraction automatique : le texte change (tableaux insérés à leur
place) et les extraits déjà ancrés des AO analysés pourraient ne plus se
retrouver. Le script compte d'abord, et n'écrit que sur demande.

On prouve :
  1. par défaut, il compte les .docx de DCE concernés et l'écart de taille,
     sans rien écrire (trames remplies, PDF et projets supprimés écartés) ;
  2. --ecrire met à jour un document qui change ;
  3. un extrait ancré qui ne se retrouverait plus est compté, et son document
     est laissé tel quel même avec --ecrire ;
  4. la ligne de commande est en comptage seul par défaut.
"""
import importlib.util
import io
import sys
from datetime import datetime
from pathlib import Path

import pytest
from docx import Document

from models.compliance_item import ComplianceItem
from models.project import Project, ProjectDocument

_SCRIPT = Path(__file__).resolve().parent.parent / "scripts" / "reextraire_docx.py"


@pytest.fixture
def script(monkeypatch):
    spec = importlib.util.spec_from_file_location("reextraire_docx", _SCRIPT)
    module = importlib.util.module_from_spec(spec)
    monkeypatch.setitem(sys.modules, spec.name, module)  # requis par @dataclass
    spec.loader.exec_module(module)
    return module


def _docx(*blocs) -> bytes:
    """Fabrique un .docx : une chaîne = un paragraphe, une liste de rangées = un tableau."""
    doc = Document()
    for bloc in blocs:
        if isinstance(bloc, str):
            doc.add_paragraph(bloc)
        else:
            t = doc.add_table(rows=len(bloc), cols=len(bloc[0]))
            for i, rangee in enumerate(bloc):
                for j, valeur in enumerate(rangee):
                    t.cell(i, j).text = valeur
    flux = io.BytesIO()
    doc.save(flux)
    return flux.getvalue()


def _ancien_texte(contenu: bytes) -> str:
    """Le texte d'avant B3 : les paragraphes seuls."""
    return "\n".join(p.text for p in Document(io.BytesIO(contenu)).paragraphs if p.text.strip())


def _deposer(db, racine, projet_id, doc_id, nom, contenu, **champs):
    url = f"/uploads/projects/{projet_id}/dce/{doc_id}-{nom}"
    if contenu is not None:
        chemin = racine / url.removeprefix("/uploads/")
        chemin.parent.mkdir(parents=True, exist_ok=True)
        chemin.write_bytes(contenu)
    if "extracted_text" not in champs:
        champs["extracted_text"] = _ancien_texte(contenu) if contenu else "texte"
    db.add(ProjectDocument(id=doc_id, project_id=projet_id, type="rc",
                           file_url=url, file_name=nom, **champs))
    db.commit()


RC_AVEC_TABLEAU = _docx(
    "Article 5 — Jugement des offres",
    [["Critère", "Pondération"], ["Valeur technique", "60 %"], ["Prix", "40 %"]],
    "Article 6 — Délai de validité des offres",
)


@pytest.fixture
def base(db_session, test_org, tmp_path):
    db = db_session
    db.add(Project(id="p-actif", organization_id=test_org.id, name="AO actif"))
    db.add(Project(id="p-supprime", organization_id=test_org.id, name="AO supprimé",
                   deleted_at=datetime.utcnow()))
    db.commit()
    _deposer(db, tmp_path, "p-actif", "d-tableau", "RC.docx", RC_AVEC_TABLEAU)
    _deposer(db, tmp_path, "p-actif", "d-simple", "AE.docx", _docx("Acte d'engagement", "Article 1"))
    _deposer(db, tmp_path, "p-actif", "d-absent", "CCAP.docx", None)
    _deposer(db, tmp_path, "p-actif", "d-trame", "DC1_signe.docx", RC_AVEC_TABLEAU,
             is_user_completed=True)
    _deposer(db, tmp_path, "p-actif", "d-pdf", "CCTP.pdf", b"%PDF-1.4 x", extracted_text="pdf")
    _deposer(db, tmp_path, "p-supprime", "d-supprime", "RC.docx", RC_AVEC_TABLEAU)
    return db, tmp_path


def _texte_en_base(db, doc_id):
    db.expire_all()
    return db.get(ProjectDocument, doc_id).extracted_text


def test_comptage_par_defaut_sans_rien_ecrire(script, base):
    db, racine = base
    avant = _texte_en_base(db, "d-tableau")
    lignes = []
    bilan = script.examiner(db, racine, detail=True, sortie=lignes.append)

    assert bilan.examines == 3          # RC, AE, CCAP — ni trame remplie, ni PDF, ni projet supprimé
    assert bilan.introuvables == 1      # CCAP absent du disque
    assert bilan.inchanges == 1         # AE sans tableau : même texte qu'avant
    assert bilan.modifies == 1          # RC : le tableau des critères en plus
    assert bilan.caracteres_avant == len(avant)
    assert bilan.caracteres_apres > bilan.caracteres_avant
    assert bilan.ecrits == 0
    assert _texte_en_base(db, "d-tableau") == avant      # rien n'est écrit
    assert len(lignes) == 1 and "RC.docx" in lignes[0]


def test_ecrire_met_a_jour_le_document_qui_change(script, base):
    db, racine = base
    bilan = script.examiner(db, racine, ecrire=True)

    assert bilan.ecrits == 1
    texte = _texte_en_base(db, "d-tableau")
    assert "Valeur technique\t60 %" in texte
    assert texte.index("Article 5") < texte.index("Critère") < texte.index("Article 6")
    assert "Critère" not in _texte_en_base(db, "d-trame")      # trame remplie intacte
    assert "Critère" not in _texte_en_base(db, "d-supprime")   # projet supprimé intact


def test_extrait_ancre_qui_ne_se_retrouverait_plus(script, base):
    """L'extrait ancré chevauche deux paragraphes que le tableau sépare
    désormais : il est compté, et le document reste tel quel."""
    db, racine = base
    contenu = _docx(
        "Les candidats remettent leur offre",
        [["Pièce", "Format"], ["Mémoire technique", "PDF"]],
        "avant le 12 mars à midi.",
    )
    _deposer(db, racine, "p-actif", "d-chevauche", "RC_lot2.docx", contenu)
    db.add(ComplianceItem(project_id="p-actif", exigence_text="Date limite de remise",
                          source_document="RC_lot2.docx",
                          source_excerpt="remettent leur offre\navant le 12 mars"))
    db.add(ComplianceItem(project_id="p-actif", exigence_text="Pondération",
                          source_document="RC.docx", source_excerpt="Article 5 — Jugement des offres"))
    db.commit()
    avant = _texte_en_base(db, "d-chevauche")

    bilan = script.examiner(db, racine, ecrire=True)

    assert bilan.modifies == 2
    assert bilan.ancres == 2            # un extrait par document modifié
    assert bilan.ancres_perdus == 1     # celui que le tableau coupe
    assert bilan.gardes == ["d-chevauche"]
    assert _texte_en_base(db, "d-chevauche") == avant
    assert "Valeur technique" in _texte_en_base(db, "d-tableau")   # l'autre est écrit


def test_ligne_de_commande_en_comptage_seul(script, base, monkeypatch, capsys):
    import services.file_storage as fs

    db, racine = base
    monkeypatch.setattr(fs, "UPLOADS_ROOT", racine)
    avant = _texte_en_base(db, "d-tableau")

    assert script.main([]) == 0
    sortie = capsys.readouterr().out
    assert "rien n'est écrit" in sortie
    assert "texte modifié : 1" in sortie
    assert _texte_en_base(db, "d-tableau") == avant

    with pytest.raises(SystemExit):
        script.main(["--dry-run", "--ecrire"])
