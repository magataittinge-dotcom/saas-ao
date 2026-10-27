"""/tmp de LibreOffice (bloc 6, revue de sécurité) : soffice crée son socket ``OSL_PIPE_<uid>_…`` dans ``/tmp`` (chemin
codé en dur dans sal : ``TMPDIR`` n'y change rien). Sous systemd (``PrivateTmp=true``), ce ``/tmp`` est celui de l'API :
LibreOffice y listait et y effaçait ses fichiers temporaires.

MESURÉ (outils-secu/essai_tmp.py, vrai soffice) : ``MAKE_SOCK`` SEUL fait échouer le démarrage (code 1, 0,1 s) car
soffice efface le socket avant de le créer (``unlink``, EACCES sans ``REMOVE_FILE`` même si le fichier n'existe pas) ;
``MAKE_SOCK | REMOVE_FILE`` suffit, ``READ_DIR`` est inutile.

Deux niveaux : ``prive`` (espace de noms de montage : le dossier jetable est monté sur ``/tmp``, le vrai ``/tmp`` n'est
plus visible) et ``partage`` (repli : ``MAKE_SOCK | REMOVE_FILE`` sans ``READ_DIR``, le lanceur de l'API balaie le socket
laissé). Un FAUX soffice en shell tente de lister, lire et effacer un fichier témoin du vrai ``/tmp``."""
from __future__ import annotations

import io
import os
import shutil
import socket
import stat
import tempfile
import uuid
from pathlib import Path

import pytest

from services import landlock
from services import libreoffice_isole as lo
from tests.test_landlock_libreoffice import needs_landlock, needs_soffice

UID = os.getuid()

_FAUX = """#!/bin/sh
while [ $# -gt 1 ]; do if [ "$1" = "--outdir" ]; then OUT="$2"; fi; shift; done
{{
  echo "%PDF-1.4 faux"
  if ls /tmp >/dev/null 2>&1; then
    if ls /tmp 2>/dev/null | grep -q "{nom}"; then echo "liste=VOIT"; else echo "liste=ne_voit_pas"; fi
  else echo "liste=refuse"; fi
  if cat "{temoin}" >/dev/null 2>&1; then echo "lire=LU"; else echo "lire=refuse"; fi
  rm -f "{temoin}" 2>/dev/null
  /usr/bin/python3 -c "import socket; s = socket.socket(socket.AF_UNIX); s.bind('/tmp/OSL_PIPE_{uid}_faux_{nom}')" 2>/dev/null \\
    && echo "socket=cree" || echo "socket=refuse"
}} > "$OUT/entree.pdf"
"""


@pytest.fixture
def decor(tmp_path, monkeypatch):
    nom = "temoin-" + uuid.uuid4().hex
    temoin = Path("/tmp") / nom                       # le VRAI /tmp (TMPDIR des tests ne le déplace pas)
    temoin.write_text("fichier temporaire de l'API\n")
    racine = tmp_path / "lo"
    racine.mkdir()
    faux = tmp_path / "faux_soffice"
    faux.write_text(_FAUX.format(nom=nom, temoin=temoin, uid=UID))
    faux.chmod(0o755)
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_DIR", str(racine))
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_LANDLOCK", "auto")
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_RESEAU", "auto")
    monkeypatch.setattr(lo, "_soffice_binary", lambda: str(faux))
    yield {"nom": nom, "temoin": temoin}
    temoin.unlink(missing_ok=True)
    for reste in Path("/tmp").glob(f"OSL_PIPE_{UID}_faux_{nom}"):
        reste.unlink(missing_ok=True)


def _lire(pdf: bytes) -> dict:
    return dict(l.split("=", 1) for l in pdf.decode().splitlines() if "=" in l)


def _sockets_osl() -> set:
    return {p.name for p in Path("/tmp").glob(f"OSL_PIPE_{UID}_*")}


@needs_landlock
def test_temoin_sans_landlock_le_faux_soffice_liste_lit_et_efface_le_tmp(decor, monkeypatch):
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_LANDLOCK", "aucun")
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_TMP", "partage")
    r = _lire(lo.convertir_en_pdf(b"x", ".docx"))
    assert r["liste"] == "VOIT" and r["lire"] == "LU"
    assert not decor["temoin"].exists()


@needs_landlock
def test_tmp_prive_le_faux_soffice_ne_voit_pas_le_vrai_tmp(decor, monkeypatch, caplog):
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_TMP", "auto")
    caplog.set_level("INFO", logger="services.libreoffice_isole")
    r = _lire(lo.convertir_en_pdf(b"x", ".docx"))
    if "tmp privé" not in caplog.text:
        pytest.skip("espace de noms de montage refusé par ce noyau (AppArmor…) : le niveau « partagé » est testé ci-dessous")
    assert r["liste"] == "ne_voit_pas" and r["lire"] == "refuse"
    assert decor["temoin"].exists()                           # ni listé, ni lu, ni effacé
    assert r["socket"] == "cree"                              # soffice a bien pu créer son socket (dans le /tmp privé)
    assert f"OSL_PIPE_{UID}_faux_{decor['nom']}" not in _sockets_osl()


@needs_landlock
def test_tmp_partage_repli_ni_liste_ni_lecture_et_le_socket_laisse_est_balaye(decor, monkeypatch, caplog):
    """Repli (espace de noms refusé) : MAKE_SOCK | REMOVE_FILE, sans READ_DIR. Reste possible : effacer un fichier dont
    on connaît le NOM EXACT (soffice efface le socket avant de le créer) ; les fichiers de l'API ont des noms
    aléatoires, qu'on ne peut pas deviner sans lister."""
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_TMP", "partage")
    caplog.set_level("INFO", logger="services.libreoffice_isole")
    ancien = Path("/tmp") / f"OSL_PIPE_{UID}_ancien_{decor['nom']}"
    s = socket.socket(socket.AF_UNIX)
    s.bind(str(ancien))
    voisin = Path("/tmp") / f"OSL_PIPE_{UID}_fichier_{decor['nom']}"
    voisin.write_text("pas un socket")
    try:
        r = _lire(lo.convertir_en_pdf(b"x", ".docx"))
        assert "tmp partagé" in caplog.text
        assert r["liste"] == "refuse" and r["lire"] == "refuse"
        assert r["socket"] == "cree"
        assert f"OSL_PIPE_{UID}_faux_{decor['nom']}" not in _sockets_osl()      # balayé par le lanceur de l'API
        assert ancien.exists() and voisin.exists()                              # ni un socket plus ancien, ni un fichier
    finally:
        s.close()
        ancien.unlink(missing_ok=True)
        voisin.unlink(missing_ok=True)


def test_le_balayage_ne_touche_que_les_sockets_osl_pipe_de_l_utilisateur_crees_pendant_la_conversion():
    tmp_path = Path(tempfile.mkdtemp(dir="/tmp"))               # chemin court : limite de 108 octets d'un socket Unix
    avant = lo._sockets_osl_pipe([str(tmp_path)])
    s1 = socket.socket(socket.AF_UNIX)
    s1.bind(str(tmp_path / f"OSL_PIPE_{UID}_a"))
    s2 = socket.socket(socket.AF_UNIX)
    s2.bind(str(tmp_path / f"OSL_PIPE_{UID + 1}_b"))             # un autre utilisateur
    s3 = socket.socket(socket.AF_UNIX)
    s3.bind(str(tmp_path / "autre_socket"))
    (tmp_path / f"OSL_PIPE_{UID}_reg").write_text("x")           # même motif, pas un socket
    (tmp_path / f"OSL_PIPE_{UID}_lien").symlink_to(tmp_path / "autre_socket")
    try:
        lo._balayer_sockets_osl_pipe([str(tmp_path)], avant)
        restes = sorted(p.name for p in tmp_path.iterdir())
        assert restes == sorted([f"OSL_PIPE_{UID + 1}_b", "autre_socket", f"OSL_PIPE_{UID}_reg", f"OSL_PIPE_{UID}_lien"])
    finally:
        for s in (s1, s2, s3):
            s.close()
        shutil.rmtree(tmp_path, ignore_errors=True)


@needs_landlock
@needs_soffice
def test_conversion_reelle_identique_et_aucun_socket_ne_reste(monkeypatch, tmp_path):
    """Un mémoire fictif et un Word déposé : mêmes pages, même texte avec ou sans Landlock, /tmp privé ou partagé ;
    aucun socket OSL_PIPE_ ne reste dans /tmp."""
    from docx import Document
    from PIL import Image
    from services.docx_exporter import build_memoire_docx
    from services.pdf_lecture import lire_pages

    monkeypatch.setenv("SYNORIX_LIBREOFFICE_DIR", str(tmp_path / "lo"))
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_RESEAU", "auto")
    memoire = build_memoire_docx({"preambule": "Texte fictif du préambule.", "partie_a": {}}, "Projet fictif", "Entreprise fictive")
    d = Document()
    d.add_heading("Titre fictif", 1)
    d.add_paragraph("Paragraphe fictif du Word déposé.")
    image = io.BytesIO()
    Image.new("RGB", (300, 150), (20, 90, 160)).save(image, "PNG")
    image.seek(0)
    d.add_picture(image)
    d.add_table(rows=2, cols=2).style = "Table Grid"
    mot = io.BytesIO()
    d.save(mot)

    avant = _sockets_osl()
    for octets in (memoire, mot.getvalue()):
        sorties = {}
        for nom, landlock_mode, tmp in (("sans", "aucun", "partage"), ("prive", "auto", "auto"), ("partage", "auto", "partage")):
            monkeypatch.setenv("SYNORIX_LIBREOFFICE_LANDLOCK", landlock_mode)
            monkeypatch.setenv("SYNORIX_LIBREOFFICE_TMP", tmp)
            sorties[nom] = lire_pages(lo.convertir_en_pdf(octets, ".docx"))
        assert len(sorties["sans"]) >= 1
        assert sorties["sans"] == sorties["prive"] == sorties["partage"]
    assert _sockets_osl() - avant == set()
