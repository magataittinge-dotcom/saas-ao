"""
« Déposé le … par … » (bloc 9, tranche 3 ; B-pieces-a-fournir.md §4 option 1 et §5).

Quand l'utilisateur choisit une pièce au coffre ou dépose un formulaire rempli,
la ligne garde la date et l'auteur du geste : l'équipe (plan Business, 5
utilisateurs) voit qui a fourni quoi. Un retrait efface la trace ; « Actualiser »
la garde avec le choix qu'elle accompagne ; aucune ligne d'une autre organisation
n'est atteignable (isolation multi-tenant). Données inventées.
"""
import io
from datetime import date, datetime, timedelta

import pytest
from openpyxl import Workbook

import routers.candidature as candidature_mod
from models.checklist_item import ChecklistItem
from models.compliance_item import ComplianceItem
from models.document import Document
from models.organization import Organization
from models.project import Project, ProjectDocument

PID = "proj-depose"


@pytest.fixture(autouse=True)
def no_rate_limit(monkeypatch):
    monkeypatch.setattr(candidature_mod.limiter, "enabled", False, raising=False)


@pytest.fixture
def uploads(tmp_path, monkeypatch):
    monkeypatch.setattr(candidature_mod, "UPLOADS_ROOT", tmp_path)
    return tmp_path


def _xlsx() -> bytes:
    wb = Workbook()
    wb.active.append(["Désignation", "Unité", "Quantité", "Prix unitaire HT", "Total HT"])
    wb.active.append(["Enduit", "m²", 10, 5.0, 50.0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _seed(db, org_id):
    db.add(Project(id=PID, organization_id=org_id, name="AO déposé"))
    db.add(Document(id="d-urssaf", organization_id=org_id, type="urssaf", status="valid",
                    expiry_date=date.today() + timedelta(days=90), category="attestations_sociales_fiscales",
                    file_name="urssaf.pdf", file_url="/uploads/o/urssaf.pdf"))
    db.add(ProjectDocument(id="t-dpgf", project_id=PID, type="dpgf_template",
                           file_url="/uploads/x/dpgf.xlsx", file_name="DPGF.xlsx"))
    db.add(ChecklistItem(id="l-urssaf", project_id=PID, document_type_required="urssaf",
                         source_kind="vault", document_group="fournir", status="manquant"))
    db.add(ChecklistItem(id="l-dpgf", project_id=PID, document_type_required="dpgf_template",
                         source_kind="dce_template", document_group="workflow", status="manquant",
                         template_project_doc_id="t-dpgf"))
    db.commit()


def _ligne(client, id_):
    return {l["id"]: l for l in client.get(f"/api/projects/{PID}/checklist").json()}[id_]


def test_une_ligne_jamais_touchee_n_a_pas_de_depot(client, db_session, test_org):
    _seed(db_session, test_org.id)
    l = _ligne(client, "l-urssaf")
    assert l["depose_le"] is None and l["depose_par"] is None


def test_choisir_au_coffre_garde_la_date_et_l_auteur(client, db_session, test_org, test_user):
    _seed(db_session, test_org.id)
    avant = datetime.utcnow() - timedelta(seconds=5)
    r = client.put(f"/api/projects/{PID}/checklist/l-urssaf/link", json={"document_id": "d-urssaf"})
    assert r.status_code == 200, r.text
    corps = r.json()
    assert corps["depose_par"] == test_user.name
    # instant UTC explicite (le « Z » permet à l'écran de dater dans le fuseau de Paris)
    assert corps["depose_le"].endswith("Z")
    instant = datetime.fromisoformat(corps["depose_le"].replace("Z", "+00:00")).replace(tzinfo=None)
    assert avant <= instant <= datetime.utcnow() + timedelta(seconds=5)
    assert _ligne(client, "l-urssaf")["depose_par"] == test_user.name


def test_deposer_un_formulaire_rempli_garde_la_date_et_l_auteur(client, db_session, test_org, test_user, uploads):
    _seed(db_session, test_org.id)
    r = client.post(
        f"/api/projects/{PID}/checklist/l-dpgf/upload-completed",
        files={"file": ("DPGF_rempli.xlsx", _xlsx(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert r.status_code == 200, r.text
    assert r.json()["depose_par"] == test_user.name and r.json()["depose_le"]


def test_retirer_la_piece_efface_la_trace(client, db_session, test_org):
    _seed(db_session, test_org.id)
    client.put(f"/api/projects/{PID}/checklist/l-urssaf/link", json={"document_id": "d-urssaf"})
    r = client.patch(f"/api/projects/{PID}/checklist/l-urssaf", json={"unlink": True})
    assert r.json()["depose_le"] is None and r.json()["depose_par"] is None


def test_actualiser_garde_le_depot_avec_le_choix(client, db_session, test_org, test_user):
    _seed(db_session, test_org.id)
    db_session.add(ComplianceItem(project_id=PID, exigence_text="Fournir une attestation de vigilance URSSAF",
                                  category="candidature", priority="obligatoire", status="non_couvert",
                                  lot="_commun"))
    db_session.commit()
    assert client.post(f"/api/projects/{PID}/checklist/regenerate").status_code == 200
    db_session.expire_all()
    ligne = db_session.query(ChecklistItem).filter_by(project_id=PID, document_type_required="urssaf").one()
    client.put(f"/api/projects/{PID}/checklist/{ligne.id}/link", json={"document_id": "d-urssaf"})
    avant = {l["document_type_required"]: l for l in client.get(f"/api/projects/{PID}/checklist").json()}["urssaf"]

    assert client.post(f"/api/projects/{PID}/checklist/regenerate").status_code == 200
    apres = {l["document_type_required"]: l for l in client.get(f"/api/projects/{PID}/checklist").json()}["urssaf"]
    assert apres["depose_par"] == test_user.name
    assert apres["depose_le"] == avant["depose_le"]


def test_la_ligne_d_une_autre_organisation_reste_inaccessible(client, db_session, test_org):
    db_session.add(Organization(id="org-autre", name="Autre"))
    db_session.add(Project(id="proj-autre", organization_id="org-autre", name="AO autre"))
    db_session.add(ChecklistItem(id="l-autre", project_id="proj-autre", document_type_required="urssaf",
                                 source_kind="vault", document_group="fournir", status="manquant"))
    db_session.commit()
    r = client.put("/api/projects/proj-autre/checklist/l-autre/link", json={"document_id": "d-urssaf"})
    assert r.status_code == 404


def test_migration_0026_ajoute_et_retire_les_colonnes_sans_doublon():
    import importlib.util
    from pathlib import Path

    import sqlalchemy as sa
    from alembic.migration import MigrationContext
    from alembic.operations import Operations

    chemin = Path(__file__).resolve().parent.parent / "alembic/versions/0026_checklist_depose.py"
    spec = importlib.util.spec_from_file_location("m0026", chemin)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    assert m.down_revision == "0025"

    moteur = sa.create_engine("sqlite://")
    with moteur.begin() as c:
        c.execute(sa.text("CREATE TABLE users (id VARCHAR PRIMARY KEY)"))
        c.execute(sa.text("CREATE TABLE checklist_items (id VARCHAR PRIMARY KEY)"))

    def lancer(f):
        with moteur.begin() as c, Operations.context(MigrationContext.configure(c, opts={"render_as_batch": True})):
            f()

    cols = lambda: {c["name"] for c in sa.inspect(moteur).get_columns("checklist_items")}
    lancer(m.upgrade)
    lancer(m.upgrade)  # idempotent
    assert {"depose_le", "depose_par_id"} <= cols()
    lancer(m.downgrade)
    assert not ({"depose_le", "depose_par_id"} & cols())
