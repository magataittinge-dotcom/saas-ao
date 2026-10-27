"""LibreOffice isolé (bloc 6, points 3 à 6) : environnement minimal, profil jetable, une conversion à
la fois ENTRE processus, délai qui tue tout le groupe, plafond de mémoire, aucun réseau.

La plupart des tests passent par un FAUX ``soffice`` (un script Python) : on observe ce que reçoit le
sous-processus sans dépendre de LibreOffice. Les tests « réels » sont sautés sans ``soffice``.
"""
from __future__ import annotations

import json
import os
import shutil
import socket
import subprocess
import sys
import textwrap
import threading
import time
from pathlib import Path

import pytest

from services import libreoffice_isole as lo
from services import libreoffice_lanceur as lanceur
from tests import lo_fabrique

BACKEND = Path(lo.__file__).resolve().parent.parent
needs_soffice = pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")),
                                   reason="LibreOffice absent de l'environnement")
needs_linux = pytest.mark.skipif(not sys.platform.startswith("linux"), reason="Linux seulement")

_ENTETE = """#!{python}
import os, sys, json, time
args = sys.argv[1:]
outdir = args[args.index("--outdir") + 1]
profil = [a for a in args if a.startswith("-env:UserInstallation=")][0].split("=", 1)[1]
"""
_PIED = """
open(os.path.join(outdir, "entree.pdf"), "wb").write(b"%PDF-1.4 faux")
"""


@pytest.fixture
def dossier_lo(tmp_path, monkeypatch):
    d = tmp_path / "lo"
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_DIR", str(d))
    return d


def _faux_soffice(tmp_path, monkeypatch, corps: str, nom: str = "faux_soffice") -> Path:
    script = tmp_path / nom
    script.write_text(_ENTETE.format(python=sys.executable) + textwrap.dedent(corps) + _PIED)
    script.chmod(0o755)
    monkeypatch.setattr(lo, "_soffice_binary", lambda: str(script))
    return script


def _vivant(pid: int) -> bool:
    try:
        with open(f"/proc/{pid}/stat") as f:
            return f.read().rsplit(")", 1)[1].split()[0] != "Z"     # un zombie est mort
    except OSError:
        return False


# ─── Point 3 : environnement minimal ─────────────────────────────────────────

@needs_linux
def test_aucun_secret_n_arrive_au_sous_processus(tmp_path, monkeypatch, dossier_lo):
    secrets = {"ANTHROPIC_API_KEY": "sk-ant-factice-0001", "STRIPE_SECRET_KEY": "sk_live_factice_0002",
               "DATABASE_URL": "url-de-base-factice-0003", "SECRET_KEY": "cle-factice-0004",
               "SYNORIX_VARIABLE_QUELCONQUE": "valeur-factice-0005"}
    for k, v in secrets.items():
        monkeypatch.setenv(k, v)
    sortie = tmp_path / "env.json"
    _faux_soffice(tmp_path, monkeypatch, f"json.dump(dict(os.environ), open({str(sortie)!r}, 'w'))\n")

    lo.convertir_en_pdf(b"PK docx fictif", ".docx")

    recu = json.loads(sortie.read_text())
    assert set(recu) <= {"PATH", "HOME", "LANG", "LC_ALL", "TMPDIR", "FONTCONFIG_FILE"}, sorted(recu)
    assert {"PATH", "HOME", "LC_ALL", "TMPDIR", "FONTCONFIG_FILE"} <= set(recu)
    dump = json.dumps(recu)
    for valeur in secrets.values():
        assert valeur not in dump
    assert recu["HOME"].startswith(str(dossier_lo))          # HOME : le dossier jetable, pas ~
    assert recu["TMPDIR"].startswith(str(dossier_lo))


# ─── Point 4 : profil jetable ────────────────────────────────────────────────

@needs_linux
def test_profil_jetable_par_conversion_et_supprime_apres(tmp_path, monkeypatch, dossier_lo):
    sortie = tmp_path / "profils.jsonl"
    _faux_soffice(tmp_path, monkeypatch, f"""
        p = profil.replace("file://", "")
        xcu = open(os.path.join(p, "user", "registrymodifications.xcu")).read()
        open({str(sortie)!r}, "a").write(json.dumps({{"profil": p, "xcu": xcu}}) + "\\n")
    """)

    lo.convertir_en_pdf(b"PK un", ".docx")
    lo.convertir_en_pdf(b"PK deux", ".docx")

    lignes = [json.loads(l) for l in sortie.read_text().splitlines()]
    assert len(lignes) == 2 and lignes[0]["profil"] != lignes[1]["profil"]
    for l in lignes:
        assert l["profil"].startswith(str(dossier_lo))        # sous le dossier inscriptible, pas dans ~
        assert not os.path.exists(l["profil"])                # supprimé après usage
        assert "DisableMacrosExecution" in l["xcu"] and "BlockUntrustedRefererLinks" in l["xcu"]
    assert [p.name for p in dossier_lo.glob("conv-*")] == []   # rien ne reste


def test_dossier_de_travail_par_defaut_est_sous_le_backend(monkeypatch):
    monkeypatch.delenv("SYNORIX_LIBREOFFICE_DIR", raising=False)
    assert lo.dossier_travail() == BACKEND / "cache" / "libreoffice"


# ─── Point 5 : une conversion à la fois, délai, plafond ──────────────────────

_SCRIPT_PROCESSUS = """
import sys, time, os, logging
sys.path.insert(0, {backend!r})
from services import libreoffice_isole as lo
lo._soffice_binary = lambda: {faux!r}
try:
    lo.convertir_en_pdf(b"PK", ".docx", attente={attente})
    print("ok")
except lo.PdfConversionError as e:
    print("echec", e.raison)
"""


@needs_linux
def test_une_conversion_a_la_fois_entre_deux_processus(tmp_path, monkeypatch, dossier_lo):
    journal = tmp_path / "journal.txt"
    faux = _faux_soffice(tmp_path, monkeypatch, f"""
        debut = time.time()
        time.sleep(1.5)
        open({str(journal)!r}, "a").write(f"{{debut}} {{time.time()}}\\n")
    """)
    code = _SCRIPT_PROCESSUS.format(backend=str(BACKEND), faux=str(faux), attente=60)
    env = {**os.environ, "SYNORIX_LIBREOFFICE_DIR": str(dossier_lo)}
    procs = [subprocess.Popen([sys.executable, "-I", "-c", code], env=env, stdout=subprocess.PIPE, text=True)
             for _ in range(2)]
    sorties = [p.communicate(timeout=60)[0].strip() for p in procs]

    assert sorties == ["ok", "ok"]
    (d1, f1), (d2, f2) = sorted(tuple(map(float, l.split())) for l in journal.read_text().splitlines())
    assert d2 >= f1 - 0.05, "les deux conversions se sont chevauchées"


@needs_linux
def test_l_attente_de_la_place_est_bornee(tmp_path, monkeypatch, dossier_lo):
    journal = tmp_path / "journal.txt"
    faux = _faux_soffice(tmp_path, monkeypatch, f"""
        open({str(journal)!r}, "a").write("tient la place\\n")
        time.sleep(3)
    """)
    code = _SCRIPT_PROCESSUS.format(backend=str(BACKEND), faux=str(faux), attente=60)
    env = {**os.environ, "SYNORIX_LIBREOFFICE_DIR": str(dossier_lo)}
    autre = subprocess.Popen([sys.executable, "-I", "-c", code], env=env, stdout=subprocess.PIPE, text=True)
    try:
        for _ in range(100):
            if journal.exists():
                break
            time.sleep(0.1)
        assert journal.exists()
        debut = time.monotonic()
        with pytest.raises(lo.PdfConversionError) as exc:
            lo.convertir_en_pdf(b"PK", ".docx", attente=0.5)
        assert exc.value.raison == "attente"
        assert time.monotonic() - debut < 2.5
        assert "réessayez" in str(exc.value).lower()
    finally:
        autre.communicate(timeout=30)


@needs_linux
def test_le_delai_tue_tout_le_groupe_de_processus(tmp_path, monkeypatch, dossier_lo):
    pid_fichier = tmp_path / "petit_fils.pid"
    _faux_soffice(tmp_path, monkeypatch, f"""
        import subprocess
        pf = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(120)"])   # comme soffice.bin
        open({str(pid_fichier)!r}, "w").write(str(pf.pid))
        time.sleep(120)
    """)
    debut = time.monotonic()
    with pytest.raises(lo.PdfConversionError) as exc:
        lo.convertir_en_pdf(b"PK", ".docx", delai=1.5)
    assert exc.value.raison == "delai" and time.monotonic() - debut < 15
    pid = int(pid_fichier.read_text())
    for _ in range(40):
        if not _vivant(pid):
            break
        time.sleep(0.1)
    assert not _vivant(pid), "le petit-fils (soffice.bin) a survécu au délai"
    assert [p.name for p in dossier_lo.glob("conv-*")] == []


@needs_linux
def test_plafond_de_memoire_pose_et_respecte(tmp_path, monkeypatch, dossier_lo):
    sortie = tmp_path / "limite.json"
    _faux_soffice(tmp_path, monkeypatch, f"""
        import resource
        json.dump({{"as": resource.getrlimit(resource.RLIMIT_AS)[0], "core": resource.getrlimit(resource.RLIMIT_CORE)[0]}},
                  open({str(sortie)!r}, "w"))
        if os.path.exists({str(tmp_path / "gourmand")!r}):
            bytearray(900 * 1024 * 1024)       # dépasse le plafond de 300 Mo
    """)
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_AS_MO", "300")
    lo.convertir_en_pdf(b"PK", ".docx")
    limite = json.loads(sortie.read_text())
    assert limite == {"as": 300 * 1024 * 1024, "core": 0}

    (tmp_path / "gourmand").write_text("1")
    with pytest.raises(lo.PdfConversionError) as exc:
        lo.convertir_en_pdf(b"PK", ".docx")
    assert exc.value.raison == "echec"
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_AS_MO", "0")      # sans plafond : la même demande passe
    lo.convertir_en_pdf(b"PK", ".docx")


def test_les_messages_sont_positifs_et_sans_detail_technique():
    interdits = ("manquant", "incomplet", "absent", "erreur", "traceback", "exception", "code ")
    for raison, message in lo._MESSAGES.items():
        assert not any(mot in message.lower() for mot in interdits), (raison, message)
        assert "!" not in message
    assert lo._erreur("attente").retry_apres_s < lo._erreur("delai").retry_apres_s


def test_delai_d_un_fichier_depose_croit_avec_la_taille_et_reste_borne():
    assert lo.delai_depot_s(0) == lo.delai_s()
    assert lo.delai_depot_s(50 * 1024 * 1024) > lo.delai_depot_s(1024 * 1024)
    assert lo.delai_depot_s(10 ** 12) == lo.DELAI_MAX_S


# ─── Point 6 : sans réseau ───────────────────────────────────────────────────

class _Ecoute:
    """Un port local ouvert par le test : compte les connexions qui y arrivent."""

    def __enter__(self):
        self.srv = socket.socket()
        self.srv.bind(("127.0.0.1", 0))
        self.srv.listen(16)
        self.srv.settimeout(0.2)
        self.port = self.srv.getsockname()[1]
        self.vues: list = []
        self._stop = threading.Event()
        self._t = threading.Thread(target=self._boucle, daemon=True)
        self._t.start()
        return self

    def _boucle(self):
        while not self._stop.is_set():
            try:
                c, _ = self.srv.accept()
            except socket.timeout:
                continue
            except OSError:
                return
            self.vues.append(1)
            c.close()

    def __exit__(self, *exc):
        time.sleep(0.3)
        self._stop.set()
        self._t.join(timeout=2)
        self.srv.close()

    @property
    def url(self) -> str:
        return f"http://127.0.0.1:{self.port}"


_CLIENT_TCP = """
import socket, sys
try:
    socket.create_connection(("127.0.0.1", {port}), timeout=3).close()
    print("connecte")
except OSError as e:
    print("bloque", type(e).__name__)
a, b = socket.socketpair()                     # un socket Unix doit rester possible (soffice <-> soffice.bin)
a.close(); b.close()
print("unix ok")
"""


def _lancer_client(mode: str, port: int, tmp_path) -> tuple[str, str]:
    """Le lanceur pose la protection `mode`, puis exécute un client TCP : (sortie, état réseau)."""
    etat = tmp_path / f"etat-{mode}.txt"
    cfg = {"cmd": [sys.executable, "-c", _CLIENT_TCP.format(port=port)], "reseau": mode, "fichier_etat": str(etat)}
    r = subprocess.run([sys.executable, "-I", "-S", str(Path(lanceur.__file__)), json.dumps(cfg)],
                       capture_output=True, text=True, timeout=30)
    assert r.returncode == 0, r.stderr
    return r.stdout, etat.read_text()


@needs_linux
def test_temoin_sans_protection_la_connexion_arrive(tmp_path):
    with _Ecoute() as ecoute:
        sortie, etat = _lancer_client("aucun", ecoute.port, tmp_path)
    assert "connecte" in sortie and len(ecoute.vues) == 1, sortie
    assert etat.startswith("aucune")


@needs_linux
def test_seccomp_bloque_toute_connexion_ip_mais_pas_les_sockets_unix(tmp_path):
    with _Ecoute() as ecoute:
        sortie, etat = _lancer_client("seccomp", ecoute.port, tmp_path)
    assert etat == "seccomp", etat
    assert "bloque PermissionError" in sortie and "unix ok" in sortie, sortie
    assert ecoute.vues == []


@needs_linux
def test_espace_reseau_vide_bloque_toute_connexion(tmp_path):
    with _Ecoute() as ecoute:
        sortie, etat = _lancer_client("netns", ecoute.port, tmp_path)
        if not etat.startswith("netns"):
            pytest.skip(f"espace de noms réseau refusé par ce noyau ({etat})")
    assert "bloque" in sortie and "unix ok" in sortie, sortie
    assert ecoute.vues == []


@needs_linux
def test_repli_netns_refuse_puis_seccomp_et_journalisation(tmp_path):
    """Noyau qui refuse l'espace de noms (AppArmor d'Ubuntu 24.04) : le filtre seccomp prend le relais
    et l'état dit pourquoi (c'est ce que l'appelant journalise)."""
    with _Ecoute() as ecoute:
        code = textwrap.dedent(f"""
            import os, sys
            sys.path.insert(0, {str(BACKEND)!r})
            def refuse(*a): raise PermissionError(1, "Operation not permitted")
            os.unshare = refuse
            from services import libreoffice_lanceur as l
            print(l.isoler_reseau("auto"))
            import socket
            try:
                socket.create_connection(("127.0.0.1", {ecoute.port}), timeout=3).close(); print("connecte")
            except OSError as e:
                print("bloque", type(e).__name__)
        """)
        r = subprocess.run([sys.executable, "-I", "-c", code], capture_output=True, text=True, timeout=30)
    assert r.returncode == 0, r.stderr
    etat, resultat = r.stdout.strip().splitlines()
    assert etat.startswith("seccomp") and "refus netns:" in etat, etat
    assert resultat == "bloque PermissionError" and ecoute.vues == []


@needs_linux
def test_isolation_exigee_mais_impossible_refuse_de_lancer(tmp_path):
    code = textwrap.dedent(f"""
        import sys, json
        sys.path.insert(0, {str(BACKEND)!r})
        from services import libreoffice_lanceur as l
        def refuse(): raise OSError(1, "refuse")
        l.poser_espace_reseau_vide = refuse
        l.poser_seccomp = refuse
        sys.argv = ["x", json.dumps({{"cmd": ["/bin/true"], "reseau": "exige", "fichier_etat": {str(tmp_path / "e.txt")!r}}})]
        l.main()
    """)
    r = subprocess.run([sys.executable, "-I", "-c", code], capture_output=True, text=True, timeout=30)
    assert r.returncode == lanceur.CODE_ISOLATION_IMPOSSIBLE, (r.returncode, r.stderr)


@needs_linux
def test_l_etat_reseau_est_journalise_sans_donnee_perso(tmp_path, monkeypatch, dossier_lo, caplog):
    _faux_soffice(tmp_path, monkeypatch, "pass\n")
    with caplog.at_level("INFO", logger="services.libreoffice_isole"):
        lo.convertir_en_pdf(b"PK", ".docx")
    journal = " ".join(r.getMessage() for r in caplog.records)
    assert "réseau :" in journal and ("seccomp" in journal or "netns" in journal)
    assert str(tmp_path) not in journal                  # ni chemin ni nom de fichier

    monkeypatch.setenv("SYNORIX_LIBREOFFICE_RESEAU", "aucun")
    caplog.clear()
    with caplog.at_level("WARNING", logger="services.libreoffice_isole"):
        lo.convertir_en_pdf(b"PK", ".docx")
    assert any("SANS isolation réseau" in r.getMessage() for r in caplog.records)


def _profil_sans_reglage(monkeypatch):
    """LibreOffice tel que livré : les réglages du profil ne comptent pas, seule la protection système."""
    monkeypatch.setattr(lo, "ecrire_profil", lambda p: p.mkdir(parents=True, exist_ok=True))


def _convertir_fichier_distant(genre: str, ecoute: "_Ecoute") -> bool:
    octets = (lo_fabrique.docx_image_distante if genre == "docx" else lo_fabrique.odt_image_distante)(ecoute.url)
    try:
        lo.convertir_en_pdf(octets, "." + genre, delai=100)
        return True
    except lo.PdfConversionError:
        return False


@needs_soffice
@needs_linux
@pytest.mark.parametrize("genre", ["docx", "odt"])
def test_libreoffice_reel_n_ouvre_aucune_connexion(genre, monkeypatch, dossier_lo):
    # Témoin : sans réglage de profil ni protection système, LibreOffice télécharge l'image liée.
    with monkeypatch.context() as m, _Ecoute() as ecoute:
        _profil_sans_reglage(m)
        m.setenv("SYNORIX_LIBREOFFICE_RESEAU", "aucun")
        assert _convertir_fichier_distant(genre, ecoute)
    if not ecoute.vues:
        pytest.skip("le témoin ne se connecte pas avec cette version de LibreOffice : preuve sans objet")

    # 1. Réglages de profil SEULS (BlockUntrustedRefererLinks) : aucune connexion.
    with monkeypatch.context() as m, _Ecoute() as ecoute:
        m.setenv("SYNORIX_LIBREOFFICE_RESEAU", "aucun")
        assert _convertir_fichier_distant(genre, ecoute)
    assert ecoute.vues == [], "le réglage de profil n'a pas suffi"

    # 2. Protection système SEULE (profil sans réglage) : filtre seccomp, puis espace réseau vide.
    for mode in ("seccomp", "netns"):
        with monkeypatch.context() as m, _Ecoute() as ecoute:
            _profil_sans_reglage(m)
            m.setenv("SYNORIX_LIBREOFFICE_RESEAU", mode)
            assert _convertir_fichier_distant(genre, ecoute), mode
        assert ecoute.vues == [], f"connexion malgré le mode {mode}"

    # 3. Tout ensemble (le défaut) : le PDF sort, aucune connexion.
    with _Ecoute() as ecoute:
        assert _convertir_fichier_distant(genre, ecoute)
    assert ecoute.vues == []


@needs_soffice
@needs_linux
def test_conversion_reelle_du_memoire_sous_isolation(dossier_lo):
    from docx import Document
    from services.pdf_lecture import lire_pages

    import io
    d = Document()
    d.add_heading("Titre fictif", 1)
    d.add_paragraph("Phrase fictive pour la conversion.")
    tampon = io.BytesIO()
    d.save(tampon)
    pdf = lo.convertir_en_pdf(tampon.getvalue(), ".docx")
    assert pdf[:5] == b"%PDF-"
    assert "Phrase fictive" in " ".join(lire_pages(pdf))
