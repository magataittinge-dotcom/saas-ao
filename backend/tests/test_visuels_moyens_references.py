"""Moyens humains et matériels, tableau des références (bloc 3, tranche 4) : seulement des
données saisies, rien d'inventé, le bloc disparaît sans donnée. Données inventées."""
import io
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile
from types import SimpleNamespace

import pytest
from docx import Document

from services.visuels import largeur, moyens, references
from services.visuels.couleurs import Theme

TH = Theme("#1E40AF")
NS = "{http://www.w3.org/2000/svg}"
needs_soffice = pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")),
                                   reason="LibreOffice absent de l'environnement")


def _cfg(materiel="", vehicules="", postes=None, gerant="Julien Marchal"):
    return SimpleNamespace(gerant_nom=gerant, gerant_titre="Gérant", postes_cles=postes or [],
                           materiel=materiel, vehicules=vehicules)


# ── Lecture des champs libres ────────────────────────────────────────────────

def test_lignes_de_materiel_lues_sans_rien_ajouter():
    items = moyens.lignes_materiel(
        "2 camions benne 3,5 t; Échafaudage de pied (location)\n- Nacelle articulée 16 m, malaxeur électrique",
        "Fourgon atelier\n3 × Nettoyeur haute pression en propre\nfourgon atelier")
    par_nom = {i["nom"]: i for i in items}
    assert par_nom["camions benne 3,5 t"]["quantite"] == "2"
    assert par_nom["camions benne 3,5 t"]["statut"] == ""                  # pas de statut écrit : rien d'affiché
    assert par_nom["Échafaudage de pied"]["statut"] == "Location"
    assert par_nom["Nettoyeur haute pression en propre"]["quantite"] == "3"
    assert par_nom["Nettoyeur haute pression en propre"]["statut"] == "En propre"
    assert par_nom["Nacelle articulée 16 m"]["picto"] == "nacelle"
    assert sum(1 for i in items if i["nom"].lower() == "fourgon atelier") == 1      # doublon retiré
    assert all(i["quantite"] == "" for i in items if i["nom"] in ("Fourgon atelier", "malaxeur électrique"))


def test_aucun_materiel_aucun_visuel():
    assert moyens.lignes_materiel("", None, "  \n ;") == []
    assert moyens.construire_materiels(_cfg(), None, TH) is None


# ── Humains ──────────────────────────────────────────────────────────────────

POSTES = [
    {"poste": "Conductrice de travaux", "nom": "Sophie Lemaire", "role": ""},
    {"poste": "Compagnon façadier", "nom": "Louis Perrin", "role": ""},
    {"poste": "Compagnon façadier", "nom": "Marc Haddad", "role": ""},
]


def test_humains_par_poste_sans_le_dirigeant_et_ouvriers_saisis():
    cartes = moyens.cartes_humaines(_cfg(postes=POSTES), None, 4)
    assert [(c["role"], c["nombre"]) for c in cartes] == [
        ("Conductrice de travaux", 1), ("Compagnon façadier", 2), ("Ouvriers dédiés au chantier", 4)]
    assert all("Julien Marchal" not in c["noms"] for c in cartes)
    assert [c["role"] for c in moyens.cartes_humaines(_cfg(postes=POSTES), None, None)][-1] == "Compagnon façadier"
    assert moyens.cartes_humaines(_cfg(), None, 0) == []
    assert moyens.construire_humains(_cfg(), None, None, TH) is None


def _verifier_texte_dans_cartes(svg):
    racine = ET.fromstring(svg)
    cartes = [(float(e.get("x")), float(e.get("y")), float(e.get("width")), float(e.get("height")))
              for e in racine.iter(NS + "rect") if e.get("stroke") not in (None, "none") and float(e.get("width")) > 60]
    assert cartes
    for tx in racine.iter(NS + "text"):
        x, y, taille = float(tx.get("x")), float(tx.get("y")), float(tx.get("font-size"))
        if tx.get("text-anchor", "start") != "start" or taille < 8:
            continue
        for cx, cy, cw, ch in cartes:
            if cx <= x <= cx + cw and cy <= y - 3 <= cy + ch:
                assert x + largeur("".join(tx.itertext()), taille, tx.get("font-weight") == "bold") <= cx + cw
                assert y <= cy + ch


def test_synoptiques_texte_dans_leurs_cartes_et_conformes():
    from services.visuels import svg_conforme
    longue = [{"poste": "Chef d'équipe façadier spécialisé en isolation thermique par l'extérieur", "nom": "Marie-Hélène Dubois-Castellane", "role": ""}]
    svg_h, alt_h = moyens.synoptique_humains(moyens.cartes_humaines(_cfg(postes=longue + POSTES), None, 12), TH)
    items = moyens.lignes_materiel("Échafaudage de pied multidirectionnel de 1 400 m² (location)\nNacelle\nCamion benne\n"
                                   "Base vie\nMalaxeur\nAspirateur classe M\nNettoyeur haute pression\nOutil quelconque")
    svg_m, alt_m = moyens.synoptique_materiels(items, TH)
    for svg in (svg_h, svg_m):
        assert svg_conforme(svg) == []
        _verifier_texte_dans_cartes(svg)
        assert max(float(e.get("x", 0)) + float(e.get("width")) for e in ET.fromstring(svg).iter(NS + "rect")
                   if e.get("fill") != "#FFFFFF" or e.get("stroke") != "none") <= 482.5
    assert "Sophie Lemaire" in svg_h and "12" in svg_h and "Moyens humains" in alt_h
    assert "Échafaudage" in alt_m


def test_visuels_placés_en_C2_et_C3_et_absents_sans_donnee():
    from services.visuels.assemblage import construire_visuels
    cfg = _cfg(materiel="Échafaudage de pied\nNacelle", postes=POSTES)
    vis = construire_visuels({"nb_ouvriers": 4}, cfg)
    assert {v.section for v in vis} == {("partie_c", "effectifs"), ("partie_c", "materiels"),
                                        ("partie_b", "demarrage")}     # l'échafaudage et la nacelle écrits font un plan
    assert construire_visuels({}, _cfg(gerant="")) == []        # aucune donnée : aucun bloc


def test_profil_local_du_memoire_prime_sur_le_profil():
    from services.visuels.assemblage import construire_visuels
    cfg = _cfg(materiel="Grue", vehicules="Fourgon")
    vis = construire_visuels({}, cfg, profile_overrides={"materiel": "Nacelle articulée", "vehicules": ""})
    svg = next(v.svg for v in vis if v.section == ("partie_c", "materiels"))
    assert "Nacelle articulée" in svg and "Grue" not in svg and "Fourgon" not in svg


# ── Tableau des références ───────────────────────────────────────────────────

def _ref(**kw):
    base = dict(intitule="", adresse=None, maitre_ouvrage=None, lot=None, montant_ht=None, annee=None)
    base.update(kw)
    return SimpleNamespace(**base)


REFS = [
    _ref(intitule="ITE de l'école Jules-Ferry", adresse="Vesle-le-Haut", maitre_ouvrage="Commune de Vesle-le-Haut",
         lot="Lot 3 — Isolation", montant_ht=410000.0, annee=2025),
    _ref(intitule="Façades de la médiathèque", adresse="Fismes-la-Vallée", maitre_ouvrage="Commune", annee=2023),
    _ref(intitule="Isolation du gymnase", adresse=None, maitre_ouvrage="Communauté de communes", annee=2024,
         montant_ht=320500.5),
]


def test_tableau_trie_formate_et_sans_colonne_vide():
    tab = references.tableau_references(REFS, "#1E40AF")
    assert tab["entetes"] == ["Opération", "Lieu", "Maître d'ouvrage", "Année", "Montant"]
    assert [l[0].split("\n")[0] for l in tab["lignes"]] == [
        "ITE de l'école Jules-Ferry", "Isolation du gymnase", "Façades de la médiathèque"]       # année décroissante
    assert tab["lignes"][0][0] == "ITE de l'école Jules-Ferry\nLot 3 — Isolation"
    assert tab["lignes"][0][4] == "410 000 € HT" and tab["lignes"][2][4] == ""      # jamais de montant inventé
    assert abs(sum(tab["largeurs_cm"]) - 17.0) < 0.05
    sans = references.tableau_references([_ref(intitule="Seul", annee=2022)], "#1E40AF")
    assert sans["entetes"] == ["Opération", "Année"]


def test_aucune_reference_aucun_tableau():
    assert references.construire([], "#1E40AF") is None
    assert references.construire([_ref(intitule="  ")], "#1E40AF") is None


def _document(out):
    return zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml").decode("utf-8")


def test_tableau_natif_en_A10_avec_entete_repetee_et_sans_tableau_markdown():
    from services.docx_exporter import build_memoire_docx
    v = references.construire(REFS, "#1E40AF")
    ia = ("Nos références récentes.\n\n| Intitulé | Lieu | MO |\n|---|---|---|\n| Chantier inventé par l'IA | X | Y |\n\n"
          "Fin du texte.")
    out = build_memoire_docx({"partie_a": {"references": ia, "fournisseurs": "Fournisseurs."}}, "P", "E", visuels=[v])
    xml = _document(out)
    assert "Chantier inventé par l'IA" not in xml and "Nos références récentes." in xml and "Fin du texte." in xml
    doc = Document(io.BytesIO(out))
    assert len(doc.tables) == 1
    tbl = doc.tables[0]
    assert [c.text for c in tbl.rows[0].cells] == ["Opération", "Lieu", "Maître d'ouvrage", "Année", "Montant"]
    assert len(tbl.rows) == 4 and "Jules-Ferry" in tbl.rows[1].cells[0].text
    assert "w:tblHeader" in xml and "w:cantSplit" in xml
    texte = re.sub(r"<[^>]+>", "|", xml)
    assert texte.rindex("10. Références chantiers") < texte.index("Jules-Ferry") < texte.rindex("11. Fournisseurs")


@needs_soffice
def test_moyens_et_references_dans_le_pdf():
    from services.docx_exporter import build_memoire_docx
    from services.pdf_export import docx_to_pdf
    from services.pdf_lecture import lire_pages
    from services.visuels.assemblage import construire_visuels
    cfg = _cfg(materiel="2 camions benne\nÉchafaudage de pied (location)", postes=POSTES)
    vis = construire_visuels({"nb_ouvriers": 4}, cfg, references=REFS)
    contenu = {"partie_a": {"references": "Texte."}, "partie_c": {"effectifs": "Équipe.", "materiels": "Matériels."}}
    texte = " ".join(lire_pages(docx_to_pdf(build_memoire_docx(contenu, "P", "E", visuels=vis))))
    for attendu in ("MOYENS HUMAINS", "Sophie Lemaire", "MOYENS MATÉRIELS", "Échafaudage de pied", "Opération", "gymnase", "320"):
        assert attendu in texte


# ── Export complet : références choisies au pre-flight, isolation entre organisations ─────

def test_export_prend_les_references_choisies_de_son_organisation(db_session, test_org):
    from models.memoire import MemoireTechnique
    from models.organization import Organization
    from models.project import Project
    from models.reference import Reference
    from routers.memoire import build_project_memoire_docx

    autre = Organization(id="autre-org", name="Autre")
    db_session.add(autre)
    db_session.add(Project(id="proj-ref", organization_id=test_org.id, name="AO"))
    db_session.add_all([
        Reference(id="r1", organization_id=test_org.id, intitule="Chantier choisi", annee=2024),
        Reference(id="r2", organization_id=test_org.id, intitule="Chantier non choisi", annee=2023),
        Reference(id="r3", organization_id="autre-org", intitule="Chantier d'une autre entreprise", annee=2024),
    ])
    db_session.commit()
    p = db_session.get(Project, "proj-ref")
    m = MemoireTechnique(project_id="proj-ref", content_json={"partie_a": {"references": "Texte."}},
                         variables={"reference_ids": ["r1", "r3"]})   # r3 : id étranger, ignoré par le scope org
    xml = _document(build_project_memoire_docx(p, m, test_org, db_session))
    assert "Chantier choisi" in xml
    assert "Chantier non choisi" not in xml and "autre entreprise" not in xml

    m.variables = {}              # pas de sélection : toutes les références de l'organisation
    xml = _document(build_project_memoire_docx(p, m, test_org, db_session))
    assert "Chantier choisi" in xml and "Chantier non choisi" in xml and "autre entreprise" not in xml


def test_la_selection_de_references_est_conservee_avec_le_memoire():
    from routers.memoire import _variables_stockees
    v = {"nb_ouvriers": 4}
    assert _variables_stockees(v, ["r1"]) == {"nb_ouvriers": 4, "reference_ids": ["r1"]}
    assert v == {"nb_ouvriers": 4}                     # l'original (envoyé au générateur) reste intact
    assert _variables_stockees(v, None) == {"nb_ouvriers": 4}
