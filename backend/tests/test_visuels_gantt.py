"""Gantt v2 (bloc 3, tranche 3) : lisible imprimé (texte >= 8 pt, libellés entiers, contraste,
noir et blanc), page A4 paysage si besoin, semaines et mois relatifs, jamais de date inventée."""
import io
import re
import shutil
import xml.etree.ElementTree as ET
import zipfile

import pytest

from services.visuels import gantt as g
from services.visuels import largeur, svg_conforme
from services.visuels.couleurs import Theme, contraste, luminance

TH = Theme("#1E40AF")
NS = "{http://www.w3.org/2000/svg}"
needs_soffice = pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")),
                                   reason="LibreOffice absent de l'environnement")

PHASES = [
    {"nom": "Période de préparation (PPSPS, PAQ, plans)", "duree_semaines": 3},
    {"nom": "Installation de chantier et clôtures", "duree_semaines": 1},
    {"nom": "Montage de l'échafaudage de la façade nord", "duree_semaines": 1},
    {"nom": "Pose de l'isolant en polystyrène graphité de 160 mm avec chevillage et harpage aux angles", "duree_semaines": 6},
    {"nom": "Sous-enduit armé et finition", "duree_semaines": 6},
    {"nom": "Démontage, repli, nettoyage", "duree_semaines": 2},
]


def _textes(svg):
    return [(float(e.get("x")), float(e.get("y")), float(e.get("font-size")), "".join(e.itertext()),
             e.get("text-anchor", "start"), e.get("font-weight") == "bold")
            for e in ET.fromstring(svg).iter(NS + "text")]


def test_texte_au_moins_8_pt_une_fois_imprime():
    svg, _, paysage, cm = g.gantt_svg(PHASES, TH)
    nominal = 728.0 if paysage else 482.0
    echelle = (cm / 2.54 * 72) / nominal           # points imprimés par unité du viewBox
    assert all(taille * echelle >= 7.98 for _, _, taille, *_ in _textes(svg))


def test_libelles_entiers_jamais_coupes():
    svg, _, _, _ = g.gantt_svg(PHASES, TH)
    lignes = [tx[3] for tx in _textes(svg)]
    assert not any(lg.endswith("…") for lg in lignes)
    for p in PHASES:   # tous les mots du libellé se retrouvent, dans l'ordre
        mots = p["nom"].split()
        rendu = " ".join(lg for lg in lignes if lg and lg.split()[0] in mots)
        assert all(m in rendu.split() for m in mots), p["nom"]


def test_les_libelles_tiennent_dans_leur_colonne():
    svg, _, paysage, _ = g.gantt_svg(PHASES, TH)
    col = 200.0 if paysage else 158.0
    for x, y, taille, texte, ancre, gras in _textes(svg):
        if x < 10 and ancre == "start" and not texte.startswith("Durée"):
            assert x + largeur(texte, taille, gras) <= col


def test_portrait_puis_paysage_selon_la_duree():
    _, _, paysage, cm = g.gantt_svg(PHASES, TH)          # 19 semaines
    assert (paysage, cm) == (False, 17.0)
    long_ = [{"nom": "Gros oeuvre", "duree_semaines": 20}, {"nom": "Façades", "duree_semaines": 12}]
    _, _, paysage, cm = g.gantt_svg(long_, TH)           # 32 semaines
    assert (paysage, cm) == (True, 25.7)


def test_semaines_mois_et_durees_calculees_sur_les_donnees():
    svg, alt, _, _ = g.gantt_svg(PHASES, TH)
    txt = [tx[3] for tx in _textes(svg)]
    assert "S1" in txt and "Mois 1" in txt and "Mois 5" in txt
    assert "S1 à S3 · 3 sem." in txt and "S5 · 1 sem." in txt and "S6 à S11 · 6 sem." in txt
    assert any("Durée totale du planning : 19 semaines" in x for x in txt)
    assert "Pose de l'isolant" in alt and "19 semaines" in alt


def test_aucune_date_de_calendrier_inventee():
    svg, _, _, _ = g.gantt_svg(PHASES, TH)
    corps = " ".join(tx[3] for tx in _textes(svg))
    assert not re.search(r"\b(20\d\d|janv|févr|mars|avr|mai|juin|juil|août|sept|oct|nov|déc)", corps)


def test_contraste_et_noir_et_blanc():
    for accent in ("#1E40AF", "#38BDF8", "#00B7FF", "#FFD400"):   # y compris cyan clair et jaune
        th = Theme(accent)
        svg, _, _, _ = g.gantt_svg(PHASES, th)
        barres = [e for e in ET.fromstring(svg).iter(NS + "rect") if e.get("fill") == th.accent]
        assert barres
        assert contraste(th.accent, "#FFFFFF") >= 4.5          # barre bien visible sur le fond blanc
        # niveau de gris de la barre : franchement sombre, donc lisible sur une impression N&B
        assert luminance(th.accent) <= 0.18
        # aucun texte clair posé sur un fond clair : tout texte est sombre sur fond blanc/très clair
        for _, _, _, _, _, _ in _textes(svg):
            pass
        fills = {e.get("fill") for e in ET.fromstring(svg).iter(NS + "text")}
        assert all(contraste(f, "#FFFFFF") >= 4.5 for f in fills), fills


def test_svg_conforme_et_hors_bornes():
    svg, _, _, _ = g.gantt_svg(PHASES, TH)
    assert svg_conforme(svg) == []
    assert g.gantt_svg([], TH) is None and g.gantt_svg(None, TH) is None
    assert g.gantt_svg([{"nom": "", "duree_semaines": 3}, {"nom": "X", "duree_semaines": 0}], TH) is None
    assert g.gantt_svg([{"nom": "Pose", "duree_semaines": 10 ** 6}], TH) is None


def test_beaucoup_de_semaines_etiquettes_espacees_sans_chevauchement():
    svg, _, paysage, _ = g.gantt_svg([{"nom": "Chantier long", "duree_semaines": 120}], TH)
    sem = sorted(x for x, y, _, texte, *_ in _textes(svg) if re.fullmatch(r"S\d+", texte))
    assert paysage and len(sem) >= 5
    assert all(b - a >= 15 for a, b in zip(sem, sem[1:]))


def test_gantt_en_a4_paysage_dans_le_word_puis_retour_au_portrait():
    from docx import Document
    from services.docx_exporter import build_memoire_docx
    longues = [{"nom": "Gros oeuvre", "duree_semaines": 20}, {"nom": "Façades", "duree_semaines": 12}]
    v = g.construire(longues, TH)
    assert v.paysage and v.section == ("partie_c", "methodologie")
    out = build_memoire_docx({"partie_c": {"methodologie": "Méthode.", "effectifs": "Équipe."}}, "P", "E", visuels=[v])
    doc = Document(io.BytesIO(out))
    orient = [s.orientation for s in doc.sections]
    assert orient[-2:] == [1, 0]        # WD_ORIENT.LANDSCAPE puis PORTRAIT
    assert doc.sections[-2].page_width > doc.sections[-2].page_height
    assert doc.sections[-1].page_width < doc.sections[-1].page_height
    xml = zipfile.ZipFile(io.BytesIO(out)).read("word/document.xml").decode("utf-8")
    assert "svgBlip" in xml


@needs_soffice
def test_texte_du_gantt_present_dans_le_pdf():
    from services.docx_exporter import build_memoire_docx
    from services.pdf_export import docx_to_pdf
    from services.pdf_lecture import lire_pages
    v = g.construire(PHASES, TH)
    pdf = docx_to_pdf(build_memoire_docx({"partie_c": {"methodologie": "Méthode."}}, "P", "E", visuels=[v]))
    texte = " ".join(lire_pages(pdf))
    for attendu in ("Installation de chantier", "S6 à S11", "Durée totale du planning : 19 semaines"):
        assert attendu in texte


@needs_soffice
def test_page_paysage_dans_le_pdf_et_texte_dans_la_page():
    import pypdfium2 as pdfium
    from services.docx_exporter import build_memoire_docx
    from services.pdf_export import docx_to_pdf
    from services.pdf_lecture import lire_pages
    v = g.construire([{"nom": "Gros oeuvre", "duree_semaines": 20}, {"nom": "Façades et enduits", "duree_semaines": 18}], TH)
    pdf = docx_to_pdf(build_memoire_docx({"partie_c": {"methodologie": "Méthode.", "effectifs": "Équipe."}}, "P", "E",
                                         visuels=[v]))
    doc = pdfium.PdfDocument(pdf)
    formats = [(p.get_width() > p.get_height()) for p in doc]
    textes = lire_pages(pdf)
    i = next(k for k, tx in enumerate(textes) if "Façades et enduits" in tx)
    assert formats[i] is True                       # la page du Gantt est en paysage
    assert any(not f for f in formats[:i])          # le reste du mémoire reste en portrait


def test_le_document_est_en_a4_portrait_et_paysage():
    from docx import Document
    from docx.shared import Cm
    from services.docx_exporter import build_memoire_docx
    v = g.construire([{"nom": "Gros oeuvre", "duree_semaines": 20}, {"nom": "Façades", "duree_semaines": 12}], TH)
    doc = Document(io.BytesIO(build_memoire_docx({"partie_c": {"methodologie": "M."}}, "P", "E", visuels=[v])))
    dims = {(round(s.page_width.cm, 1), round(s.page_height.cm, 1)) for s in doc.sections}
    assert dims == {(21.0, 29.7), (29.7, 21.0)}            # jamais de Letter
