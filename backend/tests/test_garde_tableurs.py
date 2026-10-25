"""Garde contre les bombes de tableur (.xlsx, .ods) et d'archive Office.

Un .xlsx de quelques centaines de Ko peut annoncer des centaines de Mo de chaînes
partagées : openpyxl les charge en entier, même en lecture seule (mesuré : 0,6 Mo
sur le disque, 166 Mo décompressés → 24 s et +363 Mo de mémoire). Avant, rien ne
l'empêchait, à l'extraction du DCE comme au contrôle d'une DPGF remplie ou d'une
trame déposée.

On prouve :
  1. la fin d'archive est lue sans charger le répertoire (ZIP64 compris) : une
     archive de 65 536 entrées est refusée avant que zipfile ne l'ouvre ;
  2. chaînes partagées, styles et XML total démesurés sont refusés ; un .xls trop
     lourd aussi ; un tableur, un .ods ou un .docx normal passe ;
  3. partout où un tableur est ouvert (extraction du DCE, détection des lots,
     liste des ouvrages, DPGF remplie, trame de candidature), la bombe est
     refusée sans qu'openpyxl ne soit appelé.
Tous les fichiers sont fabriqués ici.
"""
import io
import zipfile

import openpyxl
import pytest

from services import garde_archives as garde
from services.garde_archives import ArchiveDemesuree, controler_archive_office, controler_tableur


def _xlsx_normal() -> bytes:
    classeur = openpyxl.Workbook()
    classeur.active.append(["Désignation", "Unité", "Quantité"])
    classeur.active.append(["Enduit gratté", "m2", 12])
    flux = io.BytesIO()
    classeur.save(flux)
    return flux.getvalue()


def _bombe_xlsx(mo: int = 17) -> bytes:
    """Classeur valide dont les chaînes partagées décompressées pèsent `mo` Mo."""
    source = zipfile.ZipFile(io.BytesIO(_xlsx_normal()))
    element = "<si><t>" + "a" * 100 + "</t></si>"
    n = mo * 1024 * 1024 // len(element) + 1
    chaines = ('<?xml version="1.0" encoding="UTF-8"?><sst xmlns="http://schemas.openxmlformats.'
               f'org/spreadsheetml/2006/main" count="{n}" uniqueCount="{n}">'
               + element * n + "</sst>").encode()
    flux = io.BytesIO()
    with zipfile.ZipFile(flux, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as z:
        for info in source.infolist():
            donnees = source.read(info)
            if info.filename == "[Content_Types].xml":
                donnees = donnees.replace(b"</Types>", b'<Override PartName="/xl/sharedStrings.xml" '
                                          b'ContentType="application/vnd.openxmlformats-officedocument.'
                                          b'spreadsheetml.sharedStrings+xml"/></Types>')
            elif info.filename == "xl/_rels/workbook.xml.rels":
                donnees = donnees.replace(b"</Relationships>", b'<Relationship Id="rIdSST" Type="http://'
                                          b'schemas.openxmlformats.org/officeDocument/2006/relationships/'
                                          b'sharedStrings" Target="sharedStrings.xml"/></Relationships>')
            z.writestr(info, donnees)
        z.writestr("xl/sharedStrings.xml", chaines)
    return flux.getvalue()


@pytest.fixture(scope="module")
def bombe() -> bytes:
    contenu = _bombe_xlsx()
    assert len(contenu) < 1024 * 1024   # moins d'1 Mo sur le disque
    return contenu


@pytest.fixture
def openpyxl_interdit(monkeypatch):
    """Ouvertures openpyxl tentées (les appelants avalent les exceptions : on
    compte, on n'ouvre pas)."""
    appels = []

    def _jamais(*a, **kw):
        appels.append(a)
        raise RuntimeError("openpyxl ouvert sur une bombe")
    monkeypatch.setattr(openpyxl, "load_workbook", _jamais)
    return appels


# ── 1. Fin d'archive ─────────────────────────────────────────────────────────────

def test_fin_d_archive_lue_comme_zipfile():
    contenu = _xlsx_normal()
    entrees, repertoire = garde._fin_d_archive(io.BytesIO(contenu))
    assert entrees == len(zipfile.ZipFile(io.BytesIO(contenu)).infolist())
    assert 0 < repertoire < len(contenu)


def test_archive_de_65536_entrees_refusee_avant_ouverture(monkeypatch):
    flux = io.BytesIO()
    with zipfile.ZipFile(flux, "w") as z:      # au-delà de 65 535 : fin d'archive ZIP64
        for i in range(65_536):
            z.writestr(f"x/{i}", b"")
    assert garde._fin_d_archive(io.BytesIO(flux.getvalue()))[0] == 65_536

    def _jamais(*a, **kw):
        raise AssertionError("répertoire chargé par zipfile")
    monkeypatch.setattr(zipfile, "ZipFile", _jamais)
    with pytest.raises(ArchiveDemesuree, match="65536 entrées"):
        controler_archive_office(flux.getvalue())


# ── 2. Plafonds ──────────────────────────────────────────────────────────────────

def test_chaines_partagees_demesurees_refusees(bombe):
    with pytest.raises(ArchiveDemesuree, match="chaînes partagées : 17 Mo, au plus 16 Mo"):
        controler_tableur(bombe)


def test_styles_et_xml_total_demesures_refuses(monkeypatch):
    contenu = _xlsx_normal()
    controler_tableur(contenu)                      # normal : accepté
    monkeypatch.setattr(garde, "_MAX_STYLES", 100)
    with pytest.raises(ArchiveDemesuree, match="styles"):
        controler_tableur(contenu)
    monkeypatch.setattr(garde, "_MAX_STYLES", 8 * garde.MO)
    monkeypatch.setattr(garde, "_MAX_XML", 1000)
    with pytest.raises(ArchiveDemesuree, match="de XML"):
        controler_tableur(contenu)


def test_xls_borne_en_taille(monkeypatch):
    import xlwt
    classeur = xlwt.Workbook()
    classeur.add_sheet("DPGF").write(0, 0, "Désignation")
    flux = io.BytesIO()
    classeur.save(flux)
    controler_tableur(flux.getvalue())              # normal : accepté
    monkeypatch.setattr(garde, "_MAX_XLS", 100)
    with pytest.raises(ArchiveDemesuree, match="tableur démesuré"):
        controler_tableur(flux.getvalue())


def test_ods_et_docx_normaux_acceptes_et_contenu_ods_borne(monkeypatch):
    from docx import Document
    ods = io.BytesIO()
    with zipfile.ZipFile(ods, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", "<office:document-content/>" + " " * 5000)
    docx = io.BytesIO()
    Document().save(docx)
    controler_tableur(ods.getvalue())
    controler_archive_office(docx.getvalue())
    monkeypatch.setattr(garde, "_MAX_XML", 4000)
    with pytest.raises(ArchiveDemesuree):
        controler_tableur(ods.getvalue())


def test_docx_demesure_signale_non_lisible(monkeypatch):
    from docx import Document
    from services.document_processor import DocumentProcessor
    flux = io.BytesIO()
    Document().save(flux)
    monkeypatch.setattr(garde, "_MAX_ENTREES", 3)
    texte, _pages, avertissement = DocumentProcessor().extract_ex(flux.getvalue(), "CCTP.docx")
    assert texte == "" and avertissement is not None


# ── 3. Partout où un tableur est ouvert ──────────────────────────────────────────

@pytest.mark.parametrize("nom", ["DPGF.xlsx", "DPGF renommée.xls"])
def test_extraction_du_dce_refuse_la_bombe(bombe, openpyxl_interdit, nom):
    from services.document_processor import DocumentProcessor
    texte, _pages, avertissement = DocumentProcessor().extract_ex(bombe, nom)
    assert texte == "" and avertissement is not None
    assert openpyxl_interdit == []


def test_controle_de_dpgf_refuse_la_bombe(bombe, openpyxl_interdit, tmp_path):
    from services.dpgf_checker import check_dpgf
    chemin = tmp_path / "DPGF remplie.xlsx"
    chemin.write_bytes(bombe)
    rapport = check_dpgf(chemin)
    assert rapport["valid"] is False
    assert "démesuré" in rapport["warnings"][0]
    assert openpyxl_interdit == []


def test_detection_des_lots_refuse_la_bombe(bombe, openpyxl_interdit, tmp_path):
    from services.lot_detector import detect_lots_from_excel
    chemin = tmp_path / "DPGF lot 05.xlsx"
    chemin.write_bytes(bombe)
    assert detect_lots_from_excel(str(chemin)) == []
    assert openpyxl_interdit == []


def test_liste_des_ouvrages_refuse_la_bombe(bombe, openpyxl_interdit, tmp_path):
    from models.project import TYPE_DPGF
    from services.dpgf_ouvrages import texte_ouvrages_du_lot

    class _Piece:
        file_name, type, file_url = "DPGF lot 05.xlsx", TYPE_DPGF, "/uploads/dce/DPGF lot 05.xlsx"
        related_lots, is_user_completed = ["5"], False

    (tmp_path / "dce").mkdir()
    (tmp_path / "dce" / "DPGF lot 05.xlsx").write_bytes(bombe)
    assert texte_ouvrages_du_lot([_Piece()], "lot5", tmp_path) == ""
    assert openpyxl_interdit == []


def test_depots_de_dpgf_remplie_et_de_trame_refusent_la_bombe(
        client, db_session, test_org, tmp_path, monkeypatch, bombe, openpyxl_interdit):
    import routers.candidature as candidature
    import routers.export as export
    import services.file_storage as fs
    from models.checklist_item import ChecklistItem
    from models.project import TYPE_DPGF, Project, ProjectDocument

    for module in (fs, export, candidature):
        monkeypatch.setattr(module, "UPLOADS_ROOT", tmp_path)
    db_session.add(Project(id="proj-bombe", organization_id=test_org.id, name="AO"))
    db_session.add(ProjectDocument(id="pd-dpgf-b", project_id="proj-bombe", type=TYPE_DPGF,
                                   file_url="/uploads/x.xlsx", file_name="DPGF.xlsx"))
    db_session.add(ChecklistItem(id="cli-dpgf-b", project_id="proj-bombe",
                                 document_type_required=TYPE_DPGF, source_kind="dce_template",
                                 template_project_doc_id="pd-dpgf-b", status="manquant"))
    db_session.commit()
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    resp = client.post("/api/projects/proj-bombe/dpgf-upload",
                       files={"file": ("DPGF remplie.xlsx", bombe, mime)})
    assert resp.status_code == 200, resp.text
    assert resp.json()["verification"]["valid"] is False
    assert "démesuré" in resp.json()["verification"]["warnings"][0]

    resp = client.post("/api/projects/proj-bombe/checklist/cli-dpgf-b/upload-completed",
                       files={"file": ("DPGF signée.xlsx", bombe, mime)})
    assert resp.status_code == 200, resp.text
    assert resp.json()["status"] == "warning"
    assert openpyxl_interdit == []
