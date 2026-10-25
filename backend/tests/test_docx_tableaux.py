"""B3 — le texte d'un .docx garde ses tableaux, dans l'ordre du document.

Avant : `DocumentProcessor._extract_docx` ne lisait que `doc.paragraphs`. Les
tableaux d'un RC, d'un CCAP ou d'un cadre de réponse en Word (critères
pondérés, liste des pièces, tableau des lots, cellules du cadre) n'arrivaient
jamais à l'analyse. `lot_detector` relisait les tableaux de son côté, mais
après tous les paragraphes et en répétant les cellules fusionnées.

On prouve :
  1. paragraphe, tableau, paragraphe : le texte sort dans l'ordre du corps ;
  2. une cellule fusionnée (en largeur ou en hauteur) n'est lue qu'une fois ;
  3. un tableau imbriqué et une zone de texte sont lus (la zone une seule
     fois, malgré sa copie de repli `mc:Fallback`) ;
  4. contrôles de contenu et insertions suivies sont lus, les suppressions
     suivies ne le sont pas ;
  5. `lot_detector` lit exactement le même texte, et ses intitulés de lots
     restent propres (aucune tabulation) ;
  6. une partie principale démesurée (bombe de décompression) n'est pas lue.

Tous les documents sont fabriqués ici : aucun texte de DCE réel.
"""
import io
import zipfile

from docx import Document
from docx.oxml import parse_xml

from services.document_processor import DocumentProcessor

_NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"'
)


def _octets(doc) -> bytes:
    flux = io.BytesIO()
    doc.save(flux)
    return flux.getvalue()


def _texte(contenu: bytes, nom: str = "RC.docx") -> str:
    texte, pages, avertissement = DocumentProcessor().extract_ex(contenu, nom)
    assert avertissement is None
    assert pages is None
    return texte


def _ajouter_au_corps(doc, xml: str) -> None:
    """Insère un bloc XML brut dans le corps, avant les propriétés de section."""
    doc.element.body.sectPr.addprevious(parse_xml(xml))


# ── 1 et 2. ordre du corps + cellules fusionnées ─────────────────────────────

def _rc_avec_tableau_fusionne() -> bytes:
    doc = Document()
    doc.add_paragraph("Article 1 — Objet de la consultation")
    t = doc.add_table(rows=4, cols=3)
    t.cell(0, 0).text = "N° lot"
    t.cell(0, 1).text = "Désignation"
    t.cell(0, 2).text = "Délai d'exécution"
    t.cell(1, 0).text = "01"
    t.cell(1, 1).text = "Gros œuvre"
    t.cell(2, 0).text = "02"
    t.cell(2, 1).text = "Charpente"
    # Fusion en hauteur : un même délai pour les lots 01 et 02.
    t.cell(1, 2).merge(t.cell(2, 2)).text = "Six mois"
    # Fusion en largeur : une ligne qui couvre les trois colonnes.
    t.cell(3, 0).merge(t.cell(3, 2)).text = "Variante autorisée pour tous les lots"
    doc.add_paragraph("Article 2 — Suite du règlement")
    return _octets(doc)


def test_paragraphe_tableau_paragraphe_dans_l_ordre_du_document():
    texte = _texte(_rc_avec_tableau_fusionne())
    assert texte.split("\n") == [
        "Article 1 — Objet de la consultation",
        "N° lot\tDésignation\tDélai d'exécution",
        "01\tGros œuvre\tSix mois",
        "02\tCharpente",
        "Variante autorisée pour tous les lots",
        "Article 2 — Suite du règlement",
    ]


def test_cellule_fusionnee_lue_une_seule_fois():
    texte = _texte(_rc_avec_tableau_fusionne())
    assert texte.count("Six mois") == 1
    assert texte.count("Variante autorisée pour tous les lots") == 1


# ── 3. tableau imbriqué + zone de texte ──────────────────────────────────────

_ZONE_DE_TEXTE = (
    f'<w:r {_NS}><mc:AlternateContent>'
    '<mc:Choice Requires="wps"><w:drawing><wp:anchor><a:graphic><a:graphicData>'
    '<wps:wsp><wps:txbx><w:txbxContent>'
    '<w:p><w:r><w:t>Encadré : visite obligatoire</w:t></w:r></w:p>'
    '</w:txbxContent></wps:txbx></wps:wsp>'
    '</a:graphicData></a:graphic></wp:anchor></w:drawing></mc:Choice>'
    '<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent>'
    '<w:p><w:r><w:t>Encadré : visite obligatoire</w:t></w:r></w:p>'
    '</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>'
    '</mc:AlternateContent></w:r>'
)


def test_tableau_imbrique_et_zone_de_texte_lus():
    doc = Document()
    doc.add_paragraph("Avant le cadre")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Cadre A"
    cellule = t.cell(0, 1)
    cellule.paragraphs[0].text = "Liste des pièces"
    imbrique = cellule.add_table(rows=2, cols=2)
    imbrique.cell(0, 0).text = "Pièce"
    imbrique.cell(0, 1).text = "Signature"
    imbrique.cell(1, 0).text = "DC1"
    imbrique.cell(1, 1).text = "Oui"
    ancre = doc.add_paragraph("Paragraphe porteur de l'encadré")
    ancre._p.append(parse_xml(_ZONE_DE_TEXTE))
    doc.add_paragraph("Après le cadre")

    texte = _texte(_octets(doc), "Cadre_de_reponse.docx")
    assert texte.split("\n") == [
        "Avant le cadre",
        "Cadre A\tListe des pièces",
        "Pièce\tSignature",
        "DC1\tOui",
        "Paragraphe porteur de l'encadré",
        "Encadré : visite obligatoire",
        "Après le cadre",
    ]


# ── 4. contrôles de contenu, insertions et suppressions suivies ──────────────

def test_controles_de_contenu_et_revisions():
    doc = Document()
    p = doc.add_paragraph("Nom du mandataire : ")
    p._p.append(parse_xml(
        f'<w:sdt {_NS}><w:sdtPr><w:alias w:val="Nom"/></w:sdtPr>'
        '<w:sdtContent><w:r><w:t>Entreprise Exemple</w:t></w:r></w:sdtContent></w:sdt>'
    ))
    p._p.append(parse_xml(
        f'<w:ins {_NS} w:id="1" w:author="relecteur"><w:r><w:t> (mandataire solidaire)</w:t></w:r></w:ins>'
    ))
    p._p.append(parse_xml(
        f'<w:del {_NS} w:id="2" w:author="relecteur"><w:r><w:delText> texte retiré</w:delText></w:r></w:del>'
    ))
    _ajouter_au_corps(doc, (
        f'<w:sdt {_NS}><w:sdtPr/><w:sdtContent>'
        '<w:p><w:r><w:t>Bloc saisi dans un contrôle de contenu</w:t></w:r></w:p>'
        '</w:sdtContent></w:sdt>'
    ))

    texte = _texte(_octets(doc), "DC1.docx")
    assert texte.split("\n") == [
        "Nom du mandataire : Entreprise Exemple (mandataire solidaire)",
        "Bloc saisi dans un contrôle de contenu",
    ]
    assert "retiré" not in texte


# ── 5. lot_detector lit le même texte ────────────────────────────────────────

def test_lot_detector_lit_le_meme_texte_et_garde_des_intitules_propres(tmp_path):
    from services.lot_detector import _detect_lots_from_rc_text, _extract_docx_text

    doc = Document()
    doc.add_paragraph("Règlement de la consultation")
    doc.add_paragraph("Article 3 — Allotissement")
    t = doc.add_table(rows=4, cols=3)
    for i, ligne in enumerate([
        ("N° lot", "Désignation des lots", "Estimation"),
        ("01", "Gros œuvre", "120 000 € HT"),
        ("02", "Charpente couverture", "45 000 € HT"),
        ("03", "Menuiseries extérieures", "30 000 € HT"),
    ]):
        for j, valeur in enumerate(ligne):
            t.cell(i, j).text = valeur
    doc.add_paragraph("Article 4 — Délais")
    chemin = tmp_path / "RC.docx"
    chemin.write_bytes(_octets(doc))

    texte = _extract_docx_text(chemin)
    assert texte == _texte(chemin.read_bytes())

    lots = {d.id: d.nom for d in _detect_lots_from_rc_text(texte, "rc")}
    assert set(lots) == {"lot1", "lot2", "lot3"}, lots
    for nom in lots.values():
        assert "\t" not in nom and "  " not in nom, nom
    assert lots["lot2"].startswith("Lot 02 — Charpente couverture")


def test_titre_de_rubrique_de_formulaire_n_est_pas_un_lot(tmp_path):
    """Formulaire DC1 en Word, lu dans l'ordre : l'encadré « D - Présentation du
    candidat » suit la phrase « pour le lot n°… ». Ce titre de rubrique n'est
    pas un lot D (constaté sur un vrai DC1 : seul un faux « Lot D » sortait,
    à la place des lots cités plus haut)."""
    from services.lot_detector import _detect_lots_from_rc_text, _extract_docx_text

    doc = Document()
    doc.add_paragraph("Objet de la consultation : réhabilitation d'une école")
    doc.add_paragraph("Lot 2 : Menuiseries intérieures")
    doc.add_paragraph("Lot 7 : Peinture")
    doc.add_paragraph(
        "pour le lot n°……. ou les lots n°…………… de la procédure de passation du marché public"
    )
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "D - Présentation du candidat"
    doc.add_paragraph("(Cocher la case correspondante.)")
    doc.add_paragraph("Le candidat se présente seul :")
    chemin = tmp_path / "DC1.docx"
    chemin.write_bytes(_octets(doc))

    lots = {d.id: d.nom for d in _detect_lots_from_rc_text(_extract_docx_text(chemin), "autre")}
    assert set(lots) == {"lot2", "lot7"}, lots
    assert not any("Présentation" in nom for nom in lots.values()), lots


# ── 6. bombe de décompression : partie principale démesurée ──────────────────

def test_partie_principale_demesuree_non_lue(monkeypatch):
    import services.document_processor as dp

    monkeypatch.setattr(dp, "_DOCX_PARTIE_MAX", 1024)
    doc = Document()
    for i in range(200):
        doc.add_paragraph(f"Paragraphe {i} du CCTP")
    contenu = _octets(doc)
    with zipfile.ZipFile(io.BytesIO(contenu)) as z:
        assert z.getinfo("word/document.xml").file_size > 1024

    texte, _pages, avertissement = DocumentProcessor().extract_ex(contenu, "CCTP.docx")
    assert texte == ""
    assert avertissement is not None
