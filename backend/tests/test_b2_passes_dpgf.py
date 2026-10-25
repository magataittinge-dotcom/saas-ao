"""B2 — l'AE part en passe 1 ; la DPGF ne part qu'en liste d'ouvrages, sans prix.

Décision de Mohamed (25/10/2026) : « AE oui. DPGF : jamais le tableur brut à
l'IA, et jamais une colonne de prix ou de montant. Synorix extrait sans IA la
liste des ouvrages du lot (désignation, unité, quantité si l'acheteur la donne),
plafonnée à environ 300 lignes, et seule cette liste courte part en passe 2 […].
La DPGF chiffrée par l'entreprise ne part jamais à l'IA. »

Avant : les types hérités « acte_engagement » et « dpgf » ne correspondaient à
aucune pièce (seuls les *_template sont valides) : l'AE n'était jamais analysé,
et le mémoire recevait le texte brut de toutes les pièces, DPGF comprise.

On prouve, en inspectant le texte de CHAQUE appel envoyé à l'IA :
  1. composition : l'AE commun en passe 1 ; en passe 2, le CCTP du lot, l'AE propre
     au lot et la liste des ouvrages, sous « Ouvrages du lot — DPGF de l'acheteur » ;
  2. aucun prix : tableur fabriqué avec colonnes PU, montant, TVA et des montants
     repérables — aucun n'apparaît, ni aucun « € » ;
  3. le tableur brut n'est jamais envoyé (ni son texte extrait) ;
  4. la DPGF remplie et l'AE signé par l'entreprise ne sont jamais envoyés ;
  5. extracteur : plafond de 300 lignes ; xlsx, xls, ods, csv ; onglet du lot choisi
     par son intitulé quand l'acheteur s'est trompé de numéro ; lu hors de la boucle ;
  6. mémoire : ni tableur brut, ni pièce de l'entreprise ; la liste à la place.
Tous les documents sont fabriqués ici : aucun texte de DCE réel.
"""
import asyncio
import io
import threading
import zipfile

import openpyxl
import pytest

from models.project import (
    PROJECT_DOC_TYPES, TYPE_AE, TYPE_DPGF, TYPES_GRILLES_DE_PRIX, TYPES_PIECES_CLES,
    Project, ProjectDocument,
)
from services import dpgf_ouvrages as dpgf

PRIX = [4321.87, 864.37, 38.91, 9552.41, 1910.48, 2500.0, 16374.28, 3274.86, 7777.77,
        12.75, 1530.0]


def _formes(valeur: float) -> list:
    """Écritures d'un montant : 4321.87, 4321,87 ; et 4321, 4 321 dès 1 000."""
    formes = {f"{valeur:.2f}", f"{valeur:.2f}".replace(".", ",")}
    entier = int(valeur)
    if entier >= 1000:
        formes |= {str(entier), f"{entier:,}".replace(",", " ")}
    return sorted(formes)


DPGF_LOT5 = [
    ["DPGF — Lot 05 : Façades"],
    ["N°", "Désignation", "Unité", "Quantité", "Prix unitaire HT", "Montant HT", "TVA 20 %"],
    ["1", "Échafaudage de pied", "ens", 1, 4321.87, 4321.87, 864.37],
    ["2", "Enduit monocouche gratté", "m2", 245.5, 38.91, 9552.41, 1910.48],
    ["3", "Provision pour reprises : 2 500,00 € HT", "ens", None, 2500, 2500, 500],
    [None, "Sous total HT", None, None, None, 16374.28, None],
    [None, "TOTAL LOT 05 HT", None, None, None, 16374.28, 3274.86],
]
DPGF_LOT6 = [
    ["DPGF — Lot 06 : Peinture"],
    ["Désignation", "Unité", "Quantité", "PU", "Montant"],
    ["Peinture acrylique des murs", "m2", 120, 12.75, 1530.0],
]


def _xlsx(onglets: dict) -> bytes:
    classeur = openpyxl.Workbook()
    classeur.remove(classeur.active)
    for nom, lignes in onglets.items():
        feuille = classeur.create_sheet(nom)
        for ligne in lignes:
            feuille.append(ligne)
    flux = io.BytesIO()
    classeur.save(flux)
    return flux.getvalue()


def _brut(onglets: dict) -> str:
    """Le texte « tableur brut » que l'extraction au dépôt aurait stocké."""
    return "\n".join("\t".join(str(c) for c in l if c is not None)
                     for lignes in onglets.values() for l in lignes) + "\nMARQUEUR-TEXTE-BRUT"


# ── Projet complet : analyse lancée, appels à l'IA inspectés ────────────────────

APPELS = []


def _faux_appel(self, dce_text, system_prompt, label="", *a, **kw):
    APPELS.append((label, dce_text))
    return {"requirements": [{
        "exigence": f"Exigence {label}", "source_document": "RC", "source_page": 1,
        "source_excerpt": "Le candidat remet un mémoire technique.",
        "category": "technique", "priority": "obligatoire"}],
        "criteres_jugement": [], "infos_marche": {}}


@pytest.fixture
def ia_espionnee(monkeypatch):
    APPELS.clear()
    from services.ai import checklist_matcher, dce_analyzer
    monkeypatch.setattr(dce_analyzer.DCEAnalyzer, "_sync_call", _faux_appel)
    monkeypatch.setattr(dce_analyzer.DCEAnalyzer, "_demo_mode", False, raising=False)

    async def _aucun(self, *a, **kw):
        return []
    monkeypatch.setattr(checklist_matcher.ChecklistMatcher, "match", _aucun)
    return APPELS


def _deposer(db, racine, pid, nom, type_, contenu=None, texte=None, lots=None, entreprise=False):
    url = f"/uploads/projects/{pid}/dce/{nom}"
    if contenu is not None:
        chemin = racine / url.removeprefix("/uploads/")
        chemin.parent.mkdir(parents=True, exist_ok=True)
        chemin.write_bytes(contenu)
    db.add(ProjectDocument(project_id=pid, type=type_, file_url=url, file_name=nom,
                           extracted_text=texte, related_lots=lots, is_user_completed=entreprise))


@pytest.fixture
def projet(db_session, test_org, tmp_path, monkeypatch):
    import routers.analysis as analysis
    from models.organization import Organization

    monkeypatch.setattr(analysis, "UPLOADS_ROOT", tmp_path)
    db = db_session
    db.get(Organization, test_org.id).plan = "pro"
    pid = "proj-b2"
    db.add(Project(id=pid, organization_id=test_org.id, name="AO B2", selected_lot="lot5",
                   selected_lot_name="Lot 05 — Façades",
                   lots_detectes=[{"id": "lot5", "nom": "Lot 05 — Façades"},
                                  {"id": "lot6", "nom": "Lot 06 — Peinture"}]))
    onglets = {"Lot 05 - Façades": DPGF_LOT5, "Lot 06 - Peinture": DPGF_LOT6}
    _deposer(db, tmp_path, pid, "RC.pdf", "rc", lots=["all"],
             texte="RÈGLEMENT DE LA CONSULTATION. Le candidat remet un mémoire technique. " * 20)
    _deposer(db, tmp_path, pid, "CCAP.pdf", "ccap", lots=["all"],
             texte="CCAP. Les pénalités de retard sont fixées par jour calendaire. " * 20)
    _deposer(db, tmp_path, pid, "AE.pdf", TYPE_AE, lots=["all"],
             texte="ACTE D'ENGAGEMENT MARQUEUR-AE-COMMUN. Le titulaire s'engage. " * 20)
    _deposer(db, tmp_path, pid, "AE lot 06.pdf", TYPE_AE, lots=["6"],
             texte="ACTE D'ENGAGEMENT DU LOT 06 MARQUEUR-AE-LOT6. " * 20)
    _deposer(db, tmp_path, pid, "CCTP lot 05.pdf", "cctp", lots=["5"],
             texte="CCTP LOT 05 MARQUEUR-CCTP5. Les enduits sont conformes au DTU. " * 20)
    _deposer(db, tmp_path, pid, "CCTP lot 06.pdf", "cctp", lots=["6"],
             texte="CCTP LOT 06 MARQUEUR-CCTP6. Peintures. " * 20)
    _deposer(db, tmp_path, pid, "DPGF tous lots.xlsx", TYPE_DPGF, contenu=_xlsx(onglets),
             texte=_brut(onglets), lots=["all"])
    # Pièces de l'entreprise : jamais à l'IA, même si un texte avait été extrait.
    remplie = {"Lot 05": [["Désignation", "Unité", "Quantité", "PU", "Montant"],
                          ["OUVRAGE-SECRET-ENTREPRISE", "m2", 10, 777.0, 7777.77]]}
    _deposer(db, tmp_path, pid, "DPGF remplie.xlsx", TYPE_DPGF, contenu=_xlsx(remplie),
             texte="OUVRAGE-SECRET-ENTREPRISE 7777.77", lots=["5"], entreprise=True)
    _deposer(db, tmp_path, pid, "AE signé.pdf", TYPE_AE, lots=["all"], entreprise=True,
             texte="AE-SIGNE-ENTREPRISE montant de l'offre 7777.77")
    db.commit()
    return pid


def _analyser(client, pid, lots):
    resp = client.post(f"/api/projects/{pid}/analyze", json={"lots": lots})
    assert resp.status_code == 200, resp.text
    for t in threading.enumerate():
        if t.name.startswith("synorix-analysis-"):
            t.join(timeout=30)


def _textes(appels, prefixe):
    return "\n".join(texte for label, texte in appels if label.startswith(prefixe))


def test_composition_des_passes(client, projet, ia_espionnee):
    _analyser(client, projet, ["lot5"])
    passe1, passe2 = _textes(ia_espionnee, "passe1"), _textes(ia_espionnee, "passe2-lot5")

    assert "MARQUEUR-AE-COMMUN" in passe1          # l'AE part en passe 1
    assert "RÈGLEMENT DE LA CONSULTATION" in passe1 and "CCAP." in passe1
    assert "MARQUEUR-AE-LOT6" not in passe1         # un AE propre à un lot part avec ce lot
    assert dpgf.TITRE not in passe1

    assert "MARQUEUR-CCTP5" in passe2
    assert f"=== {dpgf.TITRE} ===" in passe2
    assert "Échafaudage de pied\tens\t1" in passe2
    assert "Enduit monocouche gratté\tm2\t245,5" in passe2
    assert "Provision pour reprises : (montant retiré) HT\tens" in passe2
    for autre_lot in ("MARQUEUR-CCTP6", "MARQUEUR-AE-LOT6", "Peinture acrylique des murs"):
        assert autre_lot not in passe2


def test_l_ae_propre_a_un_lot_part_avec_ce_lot(client, projet, ia_espionnee):
    _analyser(client, projet, ["lot6"])
    assert "MARQUEUR-AE-LOT6" in _textes(ia_espionnee, "passe2-lot6")
    assert "Peinture acrylique des murs\tm2\t120" in _textes(ia_espionnee, "passe2-lot6")


def test_aucun_prix_aucun_tableur_brut_aucune_piece_de_l_entreprise(client, projet, ia_espionnee):
    _analyser(client, projet, ["lot5", "lot6"])
    assert {label.split("-")[0] for label, _ in ia_espionnee} == {"passe1", "passe2"}
    tout = "\n".join(texte for _, texte in ia_espionnee)

    for valeur in PRIX:
        for forme in _formes(valeur):
            assert forme not in tout, f"montant {forme} envoyé à l'IA"
    assert "€" not in tout
    for entete in ("Prix unitaire", "Montant HT", "TVA", "Sous total", "TOTAL LOT"):
        assert entete not in tout
    assert "MARQUEUR-TEXTE-BRUT" not in tout                 # jamais le texte extrait du tableur
    assert "SECRET-ENTREPRISE" not in tout and "AE-SIGNE-ENTREPRISE" not in tout


def test_la_dpgf_est_lue_hors_de_la_boucle(client, projet, ia_espionnee, monkeypatch):
    import routers.analysis as analysis

    lieux = []
    vraie = analysis.texte_ouvrages_du_lot

    def _sonde(*a, **kw):
        try:
            asyncio.get_running_loop()
            lieux.append("boucle")
        except RuntimeError:
            lieux.append("thread")
        return vraie(*a, **kw)

    monkeypatch.setattr(analysis, "texte_ouvrages_du_lot", _sonde)
    _analyser(client, projet, ["lot5"])
    assert lieux == ["thread"]


# ── Extracteur : formats, plafond, onglet du lot ────────────────────────────────

class _Piece:
    def __init__(self, nom, type_=TYPE_DPGF, lots=("all",), entreprise=False):
        self.file_name, self.type = nom, type_
        self.file_url = f"/uploads/dce/{nom}"
        self.related_lots, self.is_user_completed = list(lots), entreprise


def _liste(racine, nom, contenu, lot="lot5", libelle="Lot 05 — Façades", **kw):
    (racine / "dce").mkdir(exist_ok=True)
    (racine / "dce" / nom).write_bytes(contenu)
    return dpgf.texte_ouvrages_du_lot([_Piece(nom, **kw)], lot, racine, libelle)


def _ods(onglets: dict) -> bytes:
    def cellule(v):
        if v is None:
            return "<table:table-cell/>"
        if isinstance(v, (int, float)):
            return (f'<table:table-cell office:value-type="float" office:value="{v}">'
                    f'<text:p>{v}</text:p></table:table-cell>')
        return f'<table:table-cell office:value-type="string"><text:p>{v}</text:p></table:table-cell>'
    tables = "".join(
        f'<table:table table:name="{nom}">'
        # chaque ligne finit par 16 000 cellules vides répétées, comme dans un vrai .ods
        + "".join("<table:table-row>" + "".join(cellule(v) for v in l)
                  + '<table:table-cell table:number-columns-repeated="16000"/></table:table-row>'
                  for l in lignes)
        + '<table:table-row table:number-rows-repeated="1048000"><table:table-cell/></table:table-row>'
        + "</table:table>" for nom, lignes in onglets.items())
    contenu = (
        '<?xml version="1.0" encoding="UTF-8"?><office:document-content '
        'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
        'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
        f'<office:body><office:spreadsheet>{tables}</office:spreadsheet></office:body>'
        '</office:document-content>')
    flux = io.BytesIO()
    with zipfile.ZipFile(flux, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", contenu)
    return flux.getvalue()


def _xls(onglets: dict) -> bytes:
    import xlwt
    classeur = xlwt.Workbook()
    for nom, lignes in onglets.items():
        feuille = classeur.add_sheet(nom)
        for i, ligne in enumerate(lignes):
            for j, v in enumerate(ligne):
                if v is not None:
                    feuille.write(i, j, v)
    flux = io.BytesIO()
    classeur.save(flux)
    return flux.getvalue()


def _csv(lignes) -> bytes:
    return "\n".join(";".join("" if v is None else str(v) for v in l) for l in lignes).encode("cp1252")


@pytest.mark.parametrize("nom,fabrique", [
    ("dpgf.xlsx", lambda: _xlsx({"Lot 05 - Façades": DPGF_LOT5, "Lot 06 - Peinture": DPGF_LOT6})),
    ("dpgf.xls", lambda: _xls({"Lot 05 - Façades": DPGF_LOT5, "Lot 06 - Peinture": DPGF_LOT6})),
    ("dpgf.ods", lambda: _ods({"Lot 05 - Façades": DPGF_LOT5, "Lot 06 - Peinture": DPGF_LOT6})),
    ("dpgf lot 05.csv", lambda: _csv(DPGF_LOT5)),
])
def test_formats_lus_sans_prix(tmp_path, nom, fabrique):
    lots = ["5"] if nom.endswith(".csv") else ["all"]
    texte = _liste(tmp_path, nom, fabrique(), lots=lots)
    assert "Échafaudage de pied\tens\t1" in texte
    assert "Enduit monocouche gratté\tm2\t245,5" in texte
    assert "Peinture acrylique des murs" not in texte
    for valeur in PRIX:
        for forme in _formes(valeur):
            assert forme not in texte
    assert "€" not in texte and "Sous total" not in texte


def test_plafond_de_300_lignes(tmp_path):
    lignes = [["Désignation", "Unité", "Quantité", "Prix unitaire", "Montant"]]
    lignes += [[f"Ouvrage numéro {i}", "u", i + 1, 10.0, 10.0 * (i + 1)] for i in range(450)]
    texte = _liste(tmp_path, "dpgf lot 05.xlsx", _xlsx({"DPGF": lignes}), lots=["5"])
    ouvrages = [l for l in texte.splitlines() if l.startswith("Ouvrage numéro")]
    assert len(ouvrages) == dpgf.MAX_LIGNES == 300
    assert "(… 150 lignes de plus, non reprises : liste plafonnée à 300 lignes.)" in texte


def test_onglet_choisi_par_l_intitule_du_lot(tmp_path):
    """Constaté sur un vrai DCE : l'acheteur avait titré « LOT 05 » l'onglet des
    menuiseries (lot 04) comme celui des façades (lot 05)."""
    entete = ["", "UNITÉ", "QUANT.", "PRIX UNIT.", "MONTANT"]
    classeur = _xlsx({
        "MEXT ALU": [["LOT 05 : MENUISERIES EXTÉRIEURES ALU", *entete[1:]], ["Châssis aluminium", "u"]],
        "FAÇADES": [["LOT 05 : REVÊTEMENT DE FAÇADE + ITE", *entete[1:]], ["Enduit gratté", "m2"]],
    })
    texte = _liste(tmp_path, "lots archi.xlsx", classeur, type_="autre",
                   libelle="Lot 05 — Revêtement de façade – ITE")
    assert "Enduit gratté\tm2" in texte and "Châssis" not in texte
    lot4 = dpgf.texte_ouvrages_du_lot([_Piece("lots archi.xlsx", type_="autre")], "lot4", tmp_path,
                                      "Lot 04 — Menuiseries extérieures alu – Serrurerie")
    assert "Châssis aluminium\tu" in lot4 and "Enduit" not in lot4


def test_plusieurs_quantites_seul_le_total(tmp_path):
    lignes = [
        ["", "", "BÂTIMENT A", "", "", "BÂTIMENT B", "", "", "TOTAL", "", ""],
        ["N°", "Désignation", "U", "Qté", "PU", "U", "Qté", "PU", "U", "Qté", "Montant"],
        ["1", "Dalle béton", "m3", 12, 101.5, "m3", 8, 101.5, "m3", 20, 2030.0],
    ]
    texte = _liste(tmp_path, "dpgf lot 05.xlsx", _xlsx({"DPGF": lignes}), lots=["5"])
    assert "Dalle béton\tm3\t20" in texte          # la quantité totale, jamais une partielle
    assert "101" not in texte and "2030" not in texte


def test_sans_en_tete_rien_n_est_extrait(tmp_path):
    lignes = [["Échafaudage", 4321.87], ["Enduit", 9552.41]]
    assert _liste(tmp_path, "dpgf lot 05.xlsx", _xlsx({"DPGF": lignes}), lots=["5"]) == ""


def test_jamais_la_dpgf_de_l_entreprise(tmp_path):
    texte = _liste(tmp_path, "dpgf lot 05.xlsx", _xlsx({"DPGF": DPGF_LOT5}), lots=["5"],
                   entreprise=True)
    assert texte == ""


# ── Types uniques et étiquettes de lot ──────────────────────────────────────────

def test_types_uniques_valides():
    assert {TYPE_AE, TYPE_DPGF} <= set(PROJECT_DOC_TYPES)
    assert TYPES_GRILLES_DE_PRIX <= set(PROJECT_DOC_TYPES)
    assert TYPES_PIECES_CLES <= set(PROJECT_DOC_TYPES)


def test_etiquettes_ae_et_dpgf():
    from services.document_tagger import assign_document_lots
    assert assign_document_lots("AE.pdf", TYPE_AE) == ["all"]
    assert assign_document_lots("AE lot 06.pdf", TYPE_AE) == ["6"]
    assert assign_document_lots("DPGF tous lots.xlsx", TYPE_DPGF) == ["all"]
    # « Plomberie » n'est pas un diagnostic plomb : la DPGF du lot 10 n'est plus « info ».
    assert assign_document_lots("DPGF_lot 10 Plomberie Sanitaires_DCE.xlsx", TYPE_DPGF) == ["10"]


# ── Mémoire ──────────────────────────────────────────────────────────────────────

def test_memoire_sans_tableur_brut_ni_piece_de_l_entreprise():
    from services.ai.memoire_generator import texte_dce_pour_memoire

    class _Doc:
        def __init__(self, type_, nom, texte, entreprise=False):
            self.type, self.file_name, self.extracted_text = type_, nom, texte
            self.is_user_completed = entreprise

    docs = [
        _Doc("autre", "Notice.pdf", "NOTICE-AUTRE"),
        _Doc(TYPE_DPGF, "DPGF.xlsx", "DPGF-BRUTE 4321.87"),
        _Doc("bpu_template", "BPU.pdf", "BPU-BRUT 38.91"),
        _Doc("autre", "Tableau.xls", "TABLEUR-BRUT 9552.41"),
        _Doc(TYPE_AE, "AE signé.pdf", "AE-SIGNE-ENTREPRISE", entreprise=True),
        _Doc(TYPE_AE, "AE.pdf", "AE-COMMUN"),
        _Doc("cctp", "CCTP.pdf", "CCTP-LOT"),
        _Doc("rc", "RC.pdf", "RC-TEXTE"),
    ]
    ouvrages = f"=== {dpgf.TITRE} ===\nEnduit gratté\tm2"
    texte = texte_dce_pour_memoire(docs, ouvrages)
    for absent in ("DPGF-BRUTE", "BPU-BRUT", "TABLEUR-BRUT", "AE-SIGNE", "4321", "38.91", "9552"):
        assert absent not in texte
    positions = [texte.index(m) for m in ("CCTP-LOT", "RC-TEXTE", "AE-COMMUN", dpgf.TITRE, "NOTICE-AUTRE")]
    assert positions == sorted(positions)   # CCTP, RC, AE, ouvrages, puis le reste


def test_le_job_memoire_passe_la_liste_des_ouvrages(client, db_session, projet, monkeypatch):
    import routers.memoire as memoire_mod
    import services.file_storage as fs

    import routers.analysis as analysis
    monkeypatch.setattr(fs, "UPLOADS_ROOT", analysis.UPLOADS_ROOT)
    monkeypatch.setattr(memoire_mod.limiter, "enabled", False, raising=False)
    db_session.get(Project, projet).current_step = 5
    db_session.commit()
    recus = {}

    async def _capture(self, **kwargs):
        recus.update(kwargs)
        return {"preambule": "ok", "partie_a": {"implantation": "ok"}, "partie_b": {}, "partie_c": {}}

    monkeypatch.setattr(memoire_mod.MemoireGenerator, "generate", _capture)
    resp = client.post(f"/api/projects/{projet}/memoire/generate", json={})
    assert resp.status_code == 200, resp.text
    for t in threading.enumerate():
        if t.name.startswith("synorix-memoire-"):
            t.join(timeout=30)

    assert recus["ouvrages_dpgf"].startswith(f"=== {dpgf.TITRE} ===")
    assert "Échafaudage de pied\tens\t1" in recus["ouvrages_dpgf"]
    assert "SECRET-ENTREPRISE" not in recus["ouvrages_dpgf"]
