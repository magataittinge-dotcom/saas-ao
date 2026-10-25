"""
Export « 0/N » — décision de Mohamed du 25/10 (« confirmé par le patron »).

  • Plus jamais « 0/N » : le détail d'export donne N exigences relevées (celles
    du lot du dossier), combien ont leur extrait retrouvé dans le DCE, et M
    confirmées par l'utilisateur. Le statut « couvert » n'est jamais posé tout
    seul (ni par l'analyse, ni par mots-clés) : seul le geste de l'utilisateur.
  • Confirmer une exigence (PATCH) ou tout un groupe (POST). Groupes → catégories
    réelles : candidature ← candidature ; offre technique ← offre + critères de
    notation ; exécution ← technique + planning. Isolation par organisation.
  • Rien de ce compteur dans ce qui part chez l'acheteur (ZIP : noms, mémoire).
  • Les avertissements du ZIP sont lisibles par le front (CORS) et dits en
    formulation positive ; le gate d'export reçoit le groupe de chaque pièce.
"""
import io
import json
import re
import zipfile
from pathlib import Path

import pytest

from models.audit_log import AuditLog
from models.checklist_item import ChecklistItem
from models.compliance_item import ComplianceItem
from models.memoire import MemoireTechnique
from models.organization import Organization
from models.project import Project, ProjectDocument


def _exigence(db, pid, id_, cat, lot="_commun", extrait="Extrait inventé du règlement.", statut="non_couvert"):
    db.add(ComplianceItem(
        id=id_, project_id=pid, exigence_text=f"Exigence inventée {id_}", category=cat,
        priority="obligatoire", status=statut, lot=lot, source_document="RC", source_page=2,
        source_excerpt=extrait,
    ))


def _projet(db, org_id, pid="proj-conf", selected_lot=None):
    db.add(Project(id=pid, organization_id=org_id, name="AO Gymnase", selected_lot=selected_lot))
    for id_, cat in (("c1", "candidature"), ("c2", "candidature"), ("o1", "offre"),
                     ("n1", "criteres_notation"), ("t1", "technique"), ("t2", "technique"),
                     ("p1", "planning")):
        _exigence(db, pid, f"{pid}-{id_}", cat)
    db.commit()
    return pid


def _statuts(db, pid):
    db.expire_all()
    return {i.id.split("-")[-1]: i.status for i in
            db.query(ComplianceItem).filter(ComplianceItem.project_id == pid).all()}


# ─── Détail d'export : plus de « couvertes » ────────────────────────────────

def test_detail_n_exigences_relevees_aucune_confirmee(client, db_session, test_org):
    pid = _projet(db_session, test_org.id)
    _exigence(db_session, pid, f"{pid}-sans", "technique", extrait=None)
    db_session.commit()

    d = client.get(f"/api/projects/{pid}/export/detail").json()
    assert "compliance_covered" not in d
    assert d["compliance_total"] == 8
    assert d["compliance_confirmed"] == 0
    assert d["compliance_with_source"] == 7
    items = {i["id"].split("-")[-1]: i for i in d["compliance_items"]}
    assert items["c1"]["groupe"] == "candidature"
    assert items["n1"]["groupe"] == "offre_technique"
    assert items["p1"]["groupe"] == "execution"
    assert items["c1"]["source_verifiee"] is True and items["sans"]["source_verifiee"] is False
    assert items["c1"]["source_document"] == "RC" and items["c1"]["source_page"] == 2


def test_detail_ne_compte_que_le_lot_du_dossier(client, db_session, test_org):
    pid = _projet(db_session, test_org.id, selected_lot="lot1")
    _exigence(db_session, pid, f"{pid}-l1", "technique", lot="lot1")
    _exigence(db_session, pid, f"{pid}-l2", "technique", lot="lot2")
    db_session.commit()

    d = client.get(f"/api/projects/{pid}/export/detail").json()
    ids = {i["id"].split("-")[-1] for i in d["compliance_items"]}
    assert "l1" in ids and "l2" not in ids
    assert d["compliance_total"] == 8


# ─── Confirmer une exigence, puis un groupe ─────────────────────────────────

def test_confirmer_une_exigence_puis_reprendre(client, db_session, test_org):
    pid = _projet(db_session, test_org.id)
    r = client.patch(f"/api/projects/{pid}/compliance/{pid}-c1", json={"confirmee": True})
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "couvert"
    assert client.get(f"/api/projects/{pid}/export/detail").json()["compliance_confirmed"] == 1
    assert db_session.query(AuditLog).filter(AuditLog.action == "compliance.confirmer").count() == 1

    r = client.patch(f"/api/projects/{pid}/compliance/{pid}-c1", json={"confirmee": False})
    assert r.json()["status"] == "non_couvert"
    assert client.get(f"/api/projects/{pid}/export/detail").json()["compliance_confirmed"] == 0


@pytest.mark.parametrize("groupe,attendues", [
    ("candidature", {"c1", "c2"}),
    ("offre_technique", {"o1", "n1"}),
    ("execution", {"t1", "t2", "p1"}),
])
def test_tout_confirmer_un_groupe(client, db_session, test_org, groupe, attendues):
    pid = _projet(db_session, test_org.id)
    r = client.post(f"/api/projects/{pid}/compliance/confirmer-groupe", json={"groupe": groupe})
    assert r.status_code == 200, r.text
    assert r.json()["exigences"] == len(attendues)
    statuts = _statuts(db_session, pid)
    assert {k for k, v in statuts.items() if v == "couvert"} == attendues


def test_tout_confirmer_respecte_le_lot_du_dossier(client, db_session, test_org):
    pid = _projet(db_session, test_org.id, selected_lot="lot1")
    _exigence(db_session, pid, f"{pid}-l1", "technique", lot="lot1")
    _exigence(db_session, pid, f"{pid}-l2", "technique", lot="lot2")
    db_session.commit()

    client.post(f"/api/projects/{pid}/compliance/confirmer-groupe", json={"groupe": "execution"})
    statuts = _statuts(db_session, pid)
    assert statuts["l1"] == "couvert" and statuts["t1"] == "couvert"
    assert statuts["l2"] == "non_couvert"


@pytest.mark.parametrize("corps", [
    {"groupe": "pieces"}, {"groupe": "tout"}, {}, {"groupe": "execution", "inconnu": 1},
])
def test_groupe_invalide_422(client, db_session, test_org, corps):
    pid = _projet(db_session, test_org.id)
    r = client.post(f"/api/projects/{pid}/compliance/confirmer-groupe", json=corps)
    assert r.status_code == 422, r.text
    assert "couvert" not in _statuts(db_session, pid).values()


def test_confirmation_isolee_par_organisation(client, db_session, test_org):
    db_session.add(Organization(id="org-autre", name="Autre BTP"))
    db_session.commit()
    pid_autre = _projet(db_session, "org-autre", pid="proj-autre")
    pid = _projet(db_session, test_org.id)

    assert client.patch(f"/api/projects/{pid_autre}/compliance/{pid_autre}-c1",
                        json={"confirmee": True}).status_code == 404
    assert client.post(f"/api/projects/{pid_autre}/compliance/confirmer-groupe",
                       json={"groupe": "execution"}).status_code == 404
    # une exigence d'un autre projet, par mon projet : introuvable
    assert client.patch(f"/api/projects/{pid}/compliance/{pid_autre}-c1",
                        json={"confirmee": True}).status_code == 404
    assert "couvert" not in _statuts(db_session, pid_autre).values()


def test_couvert_n_est_pose_que_par_le_geste_de_l_utilisateur():
    """Garde-fou : la valeur « couvert » n'apparaît en clair que dans le modèle ;
    le code la pose seulement via STATUT_CONFIRMEE, dans la route de confirmation."""
    racine = Path(__file__).resolve().parent.parent
    litteral = re.compile(r"""["']couvert["']""")
    fautes, poseurs = [], []
    for f in racine.rglob("*.py"):
        rel = f.relative_to(racine).as_posix()
        if rel.startswith(("venv/", "tests/", "alembic/")) or "/site-packages/" in rel:
            continue
        code = f.read_text(encoding="utf-8")
        if litteral.search(code) and rel != "models/compliance_item.py":
            fautes.append(rel)
        if re.search(r"status\s*[:=]\s*STATUT_CONFIRMEE|ComplianceItem\.status:\s*", code):
            poseurs.append(rel)
    assert fautes == [], fautes
    assert poseurs == ["routers/compliance.py"], poseurs


# ─── Ce qui part chez l'acheteur : aucun compteur ───────────────────────────

def test_zip_sans_compteur_d_exigences(client, db_session, test_org, tmp_path, monkeypatch):
    import routers.export as export_mod
    from services.pdf_export import PdfConversionError

    monkeypatch.setattr(export_mod, "UPLOADS_ROOT", tmp_path)

    def _pas_de_pdf(_):
        raise PdfConversionError("indisponible en test")
    monkeypatch.setattr(export_mod, "docx_to_pdf", _pas_de_pdf)  # le Word part dans le ZIP

    pid = _projet(db_session, test_org.id)
    db_session.add(MemoireTechnique(project_id=pid, content_json={
        "preambule": "Préambule inventé.", "partie_a": {"implantation": "Texte inventé."},
        "partie_b": {}, "partie_c": {}}))
    db_session.commit()
    client.post(f"/api/projects/{pid}/compliance/confirmer-groupe", json={"groupe": "candidature"})

    r = client.get(f"/api/projects/{pid}/export/zip")
    assert r.status_code == 200, r.text
    motif = re.compile(r"exigence|confirm|couvert|\b\d+\s*/\s*\d+\b", re.IGNORECASE)
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        noms = z.namelist()
        assert not [n for n in noms if motif.search(n)], noms
        memoire = next(n for n in noms if n.endswith("Memoire_technique.docx"))
        from docx import Document as Docx
        doc = Docx(io.BytesIO(z.read(memoire)))
        texte = "\n".join(p.text for p in doc.paragraphs)
        texte += "\n".join(c.text for t in doc.tables for row in t.rows for c in row.cells)
    assert not re.search(r"exigences? (relevée|confirmée|couverte)|confirmées par vous|couvert",
                         texte, re.IGNORECASE), texte[:500]


# ─── Avertissements du ZIP et gate d'export ─────────────────────────────────

def test_avertissements_du_zip_lisibles_et_positifs(client, db_session, test_org):
    db_session.add(Project(id="proj-avert", organization_id=test_org.id, name="AO avertissements"))
    db_session.add(ProjectDocument(project_id="proj-avert", type="dpgf_template",
                                   file_url="/uploads/x/dpgf.xlsx", file_name="DPGF.xlsx"))
    db_session.commit()

    r = client.get("/api/projects/proj-avert/export/zip", headers={"Origin": "http://localhost:3000"})
    assert r.status_code == 200, r.text
    assert "x-export-warnings" in r.headers.get("access-control-expose-headers", "").lower()
    avertissements = json.loads(r.headers["x-export-warnings"])
    assert "Mémoire technique : à générer à l'étape Mémoire" in avertissements
    assert "DPGF remplie : à déposer" in avertissements
    assert not [a for a in avertissements if re.search(r"\bnon (généré|remplie|signé)", a)], avertissements


def test_detail_export_donne_le_groupe_de_chaque_piece(client, db_session, test_org):
    db_session.add(Project(id="proj-gate", organization_id=test_org.id, name="AO gate"))
    for id_, groupe in (("g-fournir", "fournir"), ("g-workflow", "workflow"), ("g-synorix", "synorix")):
        db_session.add(ChecklistItem(id=id_, project_id="proj-gate", document_type_required="autre",
                                     document_group=groupe, status="manquant"))
    db_session.commit()
    items = client.get("/api/projects/proj-gate/export/detail").json()["checklist_items"]
    assert {i["id"]: i["document_group"] for i in items} == {
        "g-fournir": "fournir", "g-workflow": "workflow", "g-synorix": "synorix"}


def test_chaque_categorie_a_son_groupe():
    from services.exigences import categories_sans_groupe
    assert categories_sans_groupe() == set()
