"""Landlock autour de LibreOffice (bloc 6) : du code natif exécuté dans soffice ne lit ni `.env`, ni les
`uploads/` d'un autre projet, ni les sockets de la base ; il écrit seulement dans le dossier jetable de la conversion.

Un FAUX soffice en shell (`#!/bin/sh`, lisible et exécutable par la règle sur le programme lancé) tente ces
lectures, puis les consigne dans le « PDF » qu'il rend. Sautés sans Landlock dans le noyau."""
from __future__ import annotations

import json
import os
import shutil
import socket
import subprocess
import sys
import textwrap
import threading
import time
from pathlib import Path

import pytest

from services import landlock
from services import libreoffice_isole as lo
from services import libreoffice_lanceur as lanceur

BACKEND = Path(lo.__file__).resolve().parent.parent
# Le dépôt sur un disque 9p (WSL /mnt/c) rend Landlock indisponible (tests/test_landlock_9p.py) : ces tests
# vérifient Landlock POSÉ, ils ne valent que sur ext4 & co (copies de travail, CI, VPS).
needs_landlock = pytest.mark.skipif(
    landlock.abi_version() < 1 or landlock.type_systeme_fichiers(str(Path(__file__).resolve().parent)) == "9p",
    reason="Landlock absent de ce noyau, ou dépôt sur un disque 9p (WSL)")
needs_soffice = pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")),
                                   reason="LibreOffice absent de l'environnement")

_FAUX = """#!/bin/sh
# arguments : ... --outdir DIR FICHIER
while [ $# -gt 1 ]; do if [ "$1" = "--outdir" ]; then OUT="$2"; fi; shift; done
essai() {{ if cat "$2" >/dev/null 2>&1; then echo "$1=LU"; else echo "$1=refuse"; fi; }}
{{
  echo "%PDF-1.4 faux"
  essai env {env}
  essai autre_projet {autre}
  essai cote_de_l_espace {cote}
  essai entree "$OUT/entree.docx"
  if echo ecrit > "$OUT/ecrit.txt" 2>/dev/null; then echo "ecriture_dossier=ok"; else echo "ecriture_dossier=refuse"; fi
  if echo x > {hors} 2>/dev/null; then echo "ecriture_hors=ECRIT"; else echo "ecriture_hors=refuse"; fi
}} > "$OUT/entree.pdf"
"""


@pytest.fixture
def decor(tmp_path, monkeypatch):
    """Un faux serveur : `backend/.env`, un autre projet, un fichier à côté du dossier de travail, et le faux soffice."""
    (tmp_path / "backend").mkdir()
    env = tmp_path / "backend" / ".env"
    env.write_text("CLE_FACTICE=valeur-factice-landlock\n")
    autre = tmp_path / "uploads" / "projet-B" / "piece.pdf"
    autre.parent.mkdir(parents=True)
    autre.write_bytes(b"%PDF-1.4 piece d'un autre projet")
    racine = tmp_path / "lo"
    racine.mkdir()
    cote = racine / ".env"                                # à côté du dossier jetable de la conversion
    cote.write_text("CLE_FACTICE=valeur-factice-cote\n")
    hors = tmp_path / "hors.txt"
    faux = tmp_path / "faux_soffice"
    faux.write_text(_FAUX.format(env=env, autre=autre, cote=cote, hors=hors))
    faux.chmod(0o755)
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_DIR", str(racine))
    monkeypatch.setattr(lo, "_soffice_binary", lambda: str(faux))
    return {"hors": hors}


def _lire(pdf: bytes) -> dict:
    return dict(l.split("=", 1) for l in pdf.decode().splitlines() if "=" in l)


@needs_landlock
def test_temoin_sans_landlock_le_code_lit_env_et_les_autres_projets(decor, monkeypatch):
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_LANDLOCK", "aucun")
    res = _lire(lo.convertir_en_pdf(b"PK docx fictif", ".docx"))
    assert res["env"] == "LU" and res["autre_projet"] == "LU" and res["cote_de_l_espace"] == "LU"
    assert res["ecriture_hors"] == "ECRIT"


@needs_landlock
def test_avec_landlock_env_et_autres_projets_sont_fermes(decor, monkeypatch):
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_LANDLOCK", "auto")
    res = _lire(lo.convertir_en_pdf(b"PK docx fictif", ".docx"))
    assert res == {"env": "refuse", "autre_projet": "refuse", "cote_de_l_espace": "refuse",
                   "entree": "LU", "ecriture_dossier": "ok", "ecriture_hors": "refuse"}
    assert not decor["hors"].exists()


@needs_landlock
def test_le_niveau_obtenu_est_journalise(decor, monkeypatch, caplog):
    monkeypatch.setenv("SYNORIX_LIBREOFFICE_LANDLOCK", "auto")
    with caplog.at_level("INFO", logger="services.libreoffice_isole"):
        lo.convertir_en_pdf(b"PK", ".docx")
    journal = " ".join(r.getMessage() for r in caplog.records)
    assert "landlock" in journal and ("seccomp+landlock" in journal), journal
    assert str(decor["hors"].parent) not in journal


def test_les_regles_n_ouvrent_ni_le_backend_ni_les_uploads_ni_les_sockets(tmp_path):
    """Aucune règle ne désigne un sensible, ni un dossier qui le contient (Landlock ouvre tout le sous-arbre)."""
    base = tmp_path / "conv-x"
    sensibles = [BACKEND / ".env", BACKEND / "uploads", BACKEND, Path("/run/postgresql"), Path("/run/redis"),
                 Path("/var/run/postgresql"), Path("/srv")]
    for chemin, _ in lo.regles_landlock(base):
        ch = Path(chemin)
        if ch == base:
            continue
        for sensible in sensibles:
            assert ch != sensible and ch not in sensible.parents, (ch, sensible)
    ecritures = [c for c, droits in lo.regles_landlock(base) if droits & landlock.WRITE_FILE]
    assert sorted(ecritures) == sorted([str(base), "/dev/null"])


def _sous_processus(code: str) -> subprocess.CompletedProcess:
    return subprocess.run([sys.executable, "-I", "-c", textwrap.dedent(code)], capture_output=True, text=True, timeout=60)


@pytest.mark.parametrize("mode_reseau, attendu", [("auto", 0), ("exige", 97)])
def test_landlock_absent_repli_journalise_sauf_si_exige(tmp_path, mode_reseau, attendu):
    """Noyau sans Landlock : état « landlock: absent » et repli sur l'existant ; en mode exige, refus (code 97)."""
    etat = tmp_path / "etat.txt"
    r = _sous_processus(f"""
        import sys, json
        sys.path.insert(0, {str(BACKEND / "services")!r})
        sys.path.insert(0, {str(BACKEND)!r})
        import landlock
        landlock.abi_version = lambda: 0                       # noyau sans Landlock
        from services import libreoffice_lanceur as l
        sys.argv = ["x", json.dumps({{"cmd": ["/bin/true"], "reseau": {mode_reseau!r}, "fichier_etat": {str(etat)!r},
                                      "landlock": {{"regles": [["/usr", 13]]}}}})]
        l.main()
    """)
    assert r.returncode == attendu, (r.returncode, r.stderr)
    assert "landlock: absent" in etat.read_text()


@needs_landlock
def test_un_code_python_lance_comme_soffice_ne_lit_pas_le_env(tmp_path):
    """Au niveau du lanceur, avec un vrai interpréteur : lecture refusée hors liste blanche, écriture
    permise dans le seul dossier de travail."""
    (tmp_path / "backend").mkdir()
    (tmp_path / "backend" / ".env").write_text("X=1")
    travail = tmp_path / "travail"
    travail.mkdir()
    (travail / "entree.docx").write_text("contenu")
    code = textwrap.dedent(f"""
        import json
        res = {{}}
        for nom, chemin in (("env", {str(tmp_path / "backend" / ".env")!r}), ("entree", {str(travail / "entree.docx")!r})):
            try:
                open(chemin).read(); res[nom] = "LU"
            except OSError as e:
                res[nom] = "refuse"
        try:
            open({str(travail / "sortie")!r}, "w").write("ok"); res["ecriture"] = "ok"
        except OSError:
            res["ecriture"] = "refuse"
        try:
            open({str(tmp_path / "hors")!r}, "w").write("x"); res["hors"] = "ECRIT"
        except OSError:
            res["hors"] = "refuse"
        print(json.dumps(res))
    """)
    L = landlock
    regles = [["/usr", int(L.LECTURE_EXEC)], ["/lib", int(L.LECTURE_EXEC)], ["/lib64", int(L.LECTURE_EXEC)],
              [os.path.dirname(os.path.dirname(sys.executable)), int(L.LECTURE_EXEC)],
              [sys.base_prefix, int(L.LECTURE_EXEC)],
              [str(travail), int(L.LECTURE | L.ECRITURE)]]
    cfg = {"cmd": [sys.executable, "-I", "-S", "-c", code], "reseau": "aucun", "fichier_etat": str(tmp_path / "etat"),
           "landlock": {"regles": regles}}
    r = subprocess.run([sys.executable, "-I", "-S", str(Path(lanceur.__file__)), json.dumps(cfg)],
                       capture_output=True, text=True, timeout=60)
    assert r.returncode == 0, r.stderr
    assert json.loads(r.stdout.strip().splitlines()[-1]) == {
        "env": "refuse", "entree": "LU", "ecriture": "ok", "hors": "refuse"}
    assert (tmp_path / "etat").read_text().startswith("landlock")


@pytest.mark.skipif(landlock.abi_version() < 4, reason="refus réseau TCP de Landlock : ABI >= 4 (noyau 6.7)")
def test_landlock_refuse_aussi_tcp_sans_seccomp(tmp_path):
    srv = socket.socket()
    srv.bind(("127.0.0.1", 0))
    srv.listen(4)
    port = srv.getsockname()[1]
    code = f"""
import socket
try:
    socket.create_connection(("127.0.0.1", {port}), timeout=3).close(); print("connecte")
except OSError as e:
    print("bloque", type(e).__name__)
"""
    L = landlock
    regles = [["/usr", int(L.LECTURE_EXEC)], ["/lib", int(L.LECTURE_EXEC)], ["/lib64", int(L.LECTURE_EXEC)],
              [os.path.dirname(os.path.dirname(sys.executable)), int(L.LECTURE_EXEC)], [sys.base_prefix, int(L.LECTURE_EXEC)]]
    cfg = {"cmd": [sys.executable, "-I", "-S", "-c", code], "reseau": "aucun", "fichier_etat": str(tmp_path / "e"),
           "landlock": {"regles": regles}}
    r = subprocess.run([sys.executable, "-I", "-S", str(Path(lanceur.__file__)), json.dumps(cfg)],
                       capture_output=True, text=True, timeout=60)
    srv.close()
    assert "bloque" in r.stdout, (r.stdout, r.stderr)


@needs_landlock
@needs_soffice
def test_conversions_reelles_identiques_avec_et_sans_landlock(monkeypatch, tmp_path):
    """Le mémoire (46 pages en production ; ici un mémoire fictif) et un Word déposé avec image et tableau :
    même nombre de pages, même texte, avec et sans Landlock."""
    import io
    from docx import Document
    from PIL import Image
    from services.docx_exporter import build_memoire_docx
    from services.pdf_lecture import lire_pages

    monkeypatch.setenv("SYNORIX_LIBREOFFICE_DIR", str(tmp_path / "lo"))
    memoire = build_memoire_docx({"preambule": "Texte fictif du préambule.", "partie_a": {}}, "Projet fictif", "Entreprise fictive")
    d = Document()
    d.add_heading("Titre fictif", 1)
    d.add_paragraph("Paragraphe fictif du Word déposé.")
    image = io.BytesIO()
    Image.new("RGB", (300, 150), (20, 90, 160)).save(image, "PNG")
    image.seek(0)
    d.add_picture(image)
    d.add_table(rows=2, cols=2).style = "Table Grid"
    mot = io.BytesIO()
    d.save(mot)

    for octets in (memoire, mot.getvalue()):
        sorties = {}
        for mode in ("aucun", "auto"):
            monkeypatch.setenv("SYNORIX_LIBREOFFICE_LANDLOCK", mode)
            sorties[mode] = lire_pages(lo.convertir_en_pdf(octets, ".docx"))
        assert len(sorties["aucun"]) == len(sorties["auto"]) >= 1
        assert sorties["aucun"] == sorties["auto"]
