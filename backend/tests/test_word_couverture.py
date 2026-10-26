"""Bloc 3 « Le Word dit vrai » — tranche 2 : couverture.

Avant : la date de couverture sortait en anglais (« October 2026 », `strftime`
sans locale) et l'en-tête comme « Page 1 / N » étaient imprimés sur la page de
garde (pas de première page distincte). Les tabulations des styles En-tête /
Pied du modèle python-docx (8,25 cm centre, 16,5 cm droite) restaient actives.

Contenu inventé.
"""
import io
import shutil
import subprocess
from datetime import date

import pytest
from docx import Document
from docx.oxml.ns import qn

from services import docx_exporter
from services.docx_exporter import build_memoire_docx

needs_soffice = pytest.mark.skipif(
    not (shutil.which("soffice") or shutil.which("libreoffice")),
    reason="LibreOffice absent de l'environnement",
)

CONTENU = {
    "preambule": "Engagement sur ce marché fictif.",
    "partie_a": {"historique": "Entreprise fictive créée en 2005."},
    "partie_b": {},
    "partie_c": {"methodologie": "Méthode fictive."},
}


def _doc(**kw) -> Document:
    return Document(io.BytesIO(build_memoire_docx(CONTENU, "Chantier fictif", "ENTREPRISE TEST", **kw)))


@pytest.mark.parametrize("jour,attendu", [
    (date(2026, 10, 26), "octobre 2026"),
    (date(2027, 2, 1), "février 2027"),
    (date(2027, 8, 15), "août 2027"),
    (date(2027, 12, 31), "décembre 2027"),
])
def test_date_en_francais(jour, attendu):
    assert docx_exporter._date_fr(jour) == attendu


def test_couverture_porte_la_date_francaise_sans_locale(monkeypatch):
    monkeypatch.setattr(docx_exporter, "_aujourdhui", lambda: date(2026, 10, 26))
    textes = [p.text for p in _doc().paragraphs]
    assert "octobre 2026" in textes
    assert not any("October" in t for t in textes)


def test_premiere_page_distincte_et_vide():
    doc = _doc()
    section = doc.sections[0]
    assert section.different_first_page_header_footer is True
    assert "".join(p.text for p in section.first_page_header.paragraphs) == ""
    assert "".join(p.text for p in section.first_page_footer.paragraphs) == ""
    # En-tête et pied courants : nom de l'entreprise, « Page X / Y » en champs
    assert "ENTREPRISE TEST" in section.header.paragraphs[0].text
    pied = section.footer.paragraphs[0]._p.xml
    assert "PAGE" in pied and "NUMPAGES" in pied


def test_tabulations_des_styles_entete_pied_effacees():
    doc = _doc()
    section = doc.sections[0]
    for para in (section.header.paragraphs[0], section.footer.paragraphs[0]):
        tabs = para._p.pPr.find(qn("w:tabs"))
        assert tabs is not None
        effacees = {t.get(qn("w:pos")) for t in tabs if t.get(qn("w:val")) == "clear"}
        assert {"4680", "9360"} <= effacees


@needs_soffice
def test_pdf_couverture_sans_entete_ni_pied_puis_page_x_sur_y(tmp_path):
    from services.pdf_export import docx_to_pdf

    pdf = docx_to_pdf(build_memoire_docx(CONTENU, "Chantier fictif", "ENTREPRISE TEST"))
    chemin = tmp_path / "m.pdf"
    chemin.write_bytes(pdf)

    def page(n):
        return subprocess.run(["pdftotext", "-f", str(n), "-l", str(n), "-layout", str(chemin), "-"],
                              capture_output=True, text=True).stdout

    assert "Page" not in page(1)
    assert page(1).count("ENTREPRISE TEST") == 1          # le nom du corps, pas d'en-tête en plus
    assert "Page 2 / " in page(2)
    assert "ENTREPRISE TEST" in page(2)
