"""B8 — lecture plafonnée de tous les dépôts de fichiers.

Avant : le logo, l'import d'un mémoire, la DPGF remplie et la trame de
candidature faisaient `await file.read()` avant tout plafond (aucun plafond
pour l'import de mémoire), alors que Nginx accepte 2,2 Go, et sans contrôle du
contenu réel : une requête pouvait monter des centaines de Mo en RAM.

Maintenant, un utilitaire commun (services/upload_validation.recevoir_fichier),
sur le modèle du coffre-fort. On prouve, pour chacune des quatre routes :
  1. un fichier trop gros est refusé (413) sans que la route le lise ;
  2. un Content-Length démesuré est refusé (413) ;
  3. une extension hors liste est refusée (415) ;
  4. un contenu qui ne correspond pas à son extension est refusé (415) ;
  5. un fichier normal passe.
Plus, sur l'utilitaire : la copie s'arrête au plafond même quand la taille est
inconnue ou annoncée à tort, et se fait hors de la boucle d'événements.
"""
import asyncio
import io
from dataclasses import dataclass
from typing import Callable

import openpyxl
import pytest
from starlette.datastructures import UploadFile

from models.checklist_item import ChecklistItem
from models.document import Document as DocumentCoffre
from models.project import Project, ProjectDocument
from services import upload_validation as uv
from tests.pdf_fabrique import pdf_texte

PNG_1PX = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x02\x00\x00\x00\x90wS\xde\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0"
    b"\x00\x00\x00\x03\x00\x01\x18\xdd\x8d\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)
EXECUTABLE = b"MZ\x90\x00" + b"\x00" * 60 + b"PE\x00\x00"


def _pdf_avec_texte() -> bytes:
    return pdf_texte(["Mémoire technique — méthodologie d'exécution des travaux de gros œuvre. " * 3])


def _xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Désignation", "Unité", "Quantité", "Prix unitaire", "Total"])
    ws.append(["Béton de propreté", "m3", 12, 110, 1320])
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()


@dataclass
class Route:
    nom: str
    module: str               # module du routeur (plafond à abaisser)
    plafond: str              # nom de la constante de plafond
    valide: tuple             # (nom de fichier, contenu, type MIME)
    entete_valide: bytes      # octets magiques acceptés, pour fabriquer un gros fichier
    faux: str                 # nom à extension acceptée, pour un contenu d'exécutable
    url: Callable = None
    rien_stocke: Callable = None


def _preparer_projet(db, org):
    db.add(Project(id="proj-b8", organization_id=org.id, name="AO B8"))
    db.add(ProjectDocument(id="pd-dc1", project_id="proj-b8", type="dc1_template",
                           file_url="/uploads/dc1.pdf", file_name="DC1.pdf"))
    db.add(ChecklistItem(id="cli-dc1", project_id="proj-b8", document_type_required="dc1_template",
                         source_kind="dce_template", template_project_doc_id="pd-dc1",
                         status="manquant"))
    db.commit()


def _logo_rien(db, org, racine):
    db.expire_all()
    assert db.get(type(org), org.id).logo_url is None
    assert db.query(DocumentCoffre).count() == 0


def _import_rien(db, org, racine):
    pass  # aucun stockage ; l'appel IA (patché) est vérifié par la fixture


def _dpgf_rien(db, org, racine):
    db.expire_all()
    assert db.get(Project, "proj-b8").dpgf_remplie_url is None
    assert not list(racine.rglob("*dpgf*"))


def _trame_rien(db, org, racine):
    assert db.query(ProjectDocument).filter(ProjectDocument.is_user_completed.is_(True)).count() == 0
    assert not list(racine.rglob("*_DC1_signe*"))


ROUTES = [
    Route("logo", "routers.organizations", "_LOGO_MAX_SIZE",
          ("logo.png", PNG_1PX, "image/png"), b"\x89PNG\r\n\x1a\n", "logo.png",
          url=lambda: "/api/organizations/me/logo", rien_stocke=_logo_rien),
    Route("import-memoire", "routers.memoire_config", "_IMPORT_MAX_SIZE",
          ("memoire.pdf", _pdf_avec_texte(), "application/pdf"), b"%PDF-1.4\n", "memoire.pdf",
          url=lambda: "/api/memoire-config/import", rien_stocke=_import_rien),
    Route("dpgf-remplie", "routers.export", "_DPGF_MAX_SIZE",
          ("dpgf.xlsx", _xlsx(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
          b"PK\x03\x04", "dpgf.xlsx",
          url=lambda: "/api/projects/proj-b8/dpgf-upload", rien_stocke=_dpgf_rien),
    Route("trame-candidature", "routers.candidature", "_COMPLETED_MAX_SIZE",
          ("DC1_signe.pdf", b"%PDF-1.4 formulaire signe", "application/pdf"), b"%PDF-1.4\n",
          "DC1_signe.pdf",
          url=lambda: "/api/projects/proj-b8/checklist/cli-dc1/upload-completed",
          rien_stocke=_trame_rien),
]


@pytest.fixture
def env(client, db_session, test_org, tmp_path, monkeypatch):
    """Stockage dans tmp_path, IA de l'import neutralisée (aucun appel payant)."""
    import routers.candidature as candidature
    import routers.export as export
    import routers.memoire_config as memoire_config
    import services.file_storage as fs

    monkeypatch.setattr(fs, "UPLOADS_ROOT", tmp_path)
    monkeypatch.setattr(export, "UPLOADS_ROOT", tmp_path)
    monkeypatch.setattr(candidature, "UPLOADS_ROOT", tmp_path)
    appels_ia = []

    async def _faux_import(self, texte):
        appels_ia.append(len(texte))
        return {"nom_entreprise": "Exemple BTP"}

    monkeypatch.setattr(memoire_config.MemoireImporter, "extract", _faux_import)
    _preparer_projet(db_session, test_org)
    return client, db_session, test_org, tmp_path, appels_ia


def _poster(client, route, nom, contenu, type_mime="application/octet-stream", **kw):
    return client.post(route.url(), files={"file": (nom, contenu, type_mime)}, **kw)


@pytest.fixture
def espion_lecture(monkeypatch):
    """Compte toute lecture du fichier reçu : `await file.read()` (ancien code)
    et la copie plafonnée de l'utilitaire."""
    lectures = []
    lire = UploadFile.read
    copier = uv.copie_plafonnee

    async def _lire(self, size=-1):
        lectures.append(("UploadFile.read", size))
        return await lire(self, size)

    def _copier(source, plafond):
        lectures.append(("copie_plafonnee", plafond))
        return copier(source, plafond)

    monkeypatch.setattr(UploadFile, "read", _lire)
    monkeypatch.setattr(uv, "copie_plafonnee", _copier)
    return lectures


# ── 1. trop gros : refusé sans être lu ───────────────────────────────────────

@pytest.mark.parametrize("route", ROUTES, ids=lambda r: r.nom)
def test_trop_gros_refuse_sans_etre_lu(route, env, espion_lecture, monkeypatch):
    import importlib

    client, db, org, racine, appels_ia = env
    monkeypatch.setattr(importlib.import_module(route.module), route.plafond, 2048)
    gros = route.entete_valide + b"\x00" * 8192

    resp = _poster(client, route, route.valide[0], gros, route.valide[2])

    assert resp.status_code == 413, resp.text
    assert "trop volumineux" in resp.json()["detail"]
    assert espion_lecture == []          # ni read() complet, ni copie
    assert appels_ia == []
    route.rien_stocke(db, org, racine)


# ── 2. Content-Length démesuré ───────────────────────────────────────────────

@pytest.mark.parametrize("route", ROUTES, ids=lambda r: r.nom)
def test_content_length_demesure_refuse(route, env, espion_lecture):
    client, db, org, racine, appels_ia = env
    nom, contenu, type_mime = route.valide

    resp = _poster(client, route, nom, contenu, type_mime,
                   headers={"content-length": str(3 * 1024 ** 3)})

    assert resp.status_code == 413, resp.text
    assert espion_lecture == []
    assert appels_ia == []
    route.rien_stocke(db, org, racine)


# ── 3. extension hors liste ──────────────────────────────────────────────────

@pytest.mark.parametrize("route", ROUTES, ids=lambda r: r.nom)
def test_extension_hors_liste_refusee(route, env):
    client, db, org, racine, appels_ia = env

    resp = _poster(client, route, "outil.exe", EXECUTABLE)

    assert resp.status_code == 415, resp.text
    assert "Formats acceptés" in resp.json()["detail"]
    assert appels_ia == []
    route.rien_stocke(db, org, racine)


# ── 4. contenu qui ne correspond pas à l'extension ───────────────────────────

@pytest.mark.parametrize("route", ROUTES, ids=lambda r: r.nom)
def test_contenu_ne_correspondant_pas_refuse(route, env):
    client, db, org, racine, appels_ia = env

    resp = _poster(client, route, route.faux, EXECUTABLE)

    assert resp.status_code == 415, resp.text
    assert "ne correspond pas" in resp.json()["detail"]
    assert appels_ia == []
    route.rien_stocke(db, org, racine)


# ── 5. un fichier normal passe ───────────────────────────────────────────────

@pytest.mark.parametrize("route", ROUTES, ids=lambda r: r.nom)
def test_fichier_normal_accepte(route, env):
    client, db, org, racine, appels_ia = env
    nom, contenu, type_mime = route.valide

    resp = _poster(client, route, nom, contenu, type_mime)

    assert resp.status_code == 200, resp.text
    db.expire_all()
    if route.nom == "logo":
        url = db.get(type(org), org.id).logo_url
        assert (racine / url.removeprefix("/uploads/")).read_bytes() == contenu
    elif route.nom == "import-memoire":
        assert appels_ia and resp.json()["fields_count"] == 1
    elif route.nom == "dpgf-remplie":
        url = db.get(Project, "proj-b8").dpgf_remplie_url
        assert (racine / url.removeprefix("/uploads/")).read_bytes() == contenu
        assert "verification" in resp.json()
    else:
        pd = db.query(ProjectDocument).filter(ProjectDocument.is_user_completed.is_(True)).one()
        assert pd.file_size == len(contenu)
        assert (racine / pd.file_url.removeprefix("/uploads/")).read_bytes() == contenu


# ── Utilitaire : copie coupée au plafond ─────────────────────────────────────

class _FluxCompte(io.BytesIO):
    def __init__(self, contenu):
        super().__init__(contenu)
        self.lus = 0

    def read(self, n=-1):
        morceau = super().read(n)
        self.lus += len(morceau)
        return morceau


def test_copie_coupee_au_plafond():
    source = _FluxCompte(b"%PDF-1.4\n" + b"x" * (10 * uv.MO))
    with pytest.raises(uv._Depassement):
        uv.copie_plafonnee(source, 2 * uv.MO)
    assert source.lus <= 2 * uv.MO + uv.VAULT_CHUNK_SIZE   # 3 Mo lus sur 10


@pytest.mark.parametrize("taille_annoncee", [None, 10], ids=["taille-inconnue", "taille-fausse"])
def test_taille_inconnue_ou_fausse_coupee_au_plafond(taille_annoncee):
    from fastapi import HTTPException

    source = _FluxCompte(b"%PDF-1.4\n" + b"x" * (5 * uv.MO))
    fichier = UploadFile(file=source, size=taille_annoncee, filename="plan.pdf")
    regle = uv.RegleDepot(frozenset({".pdf"}), 1 * uv.MO, "PDF")

    with pytest.raises(HTTPException) as refus:
        asyncio.run(uv.recevoir_fichier(None, fichier, regle))
    assert refus.value.status_code == 413
    assert refus.value.detail == "Fichier trop volumineux (max 1 Mo)."
    assert source.lus <= 1 * uv.MO + uv.VAULT_CHUNK_SIZE


def test_copie_hors_de_la_boucle(monkeypatch):
    lieux = []
    copier = uv.copie_plafonnee

    def _sonde(source, plafond):
        try:
            asyncio.get_running_loop()
            lieux.append("boucle")
        except RuntimeError:
            lieux.append("thread")
        return copier(source, plafond)

    monkeypatch.setattr(uv, "copie_plafonnee", _sonde)
    fichier = UploadFile(file=io.BytesIO(b"%PDF-1.4 ok"), size=11, filename="ok.pdf")
    regle = uv.RegleDepot(frozenset({".pdf"}), uv.MO, "PDF")

    async def _recevoir():
        with await uv.recevoir_fichier(None, fichier, regle) as recu:
            return recu.taille, await recu.octets()

    assert asyncio.run(_recevoir()) == (11, b"%PDF-1.4 ok")
    assert lieux == ["thread"]


# ── Enregistrer au coffre une pièce du projet (« .read() sans taille ») ──────
# Avant : `storage.upload(src.read_bytes(), …)` lisait toute la pièce (un plan
# de DCE peut peser des centaines de Mo) en RAM, dans la boucle d'événements.

def _piece_du_projet(db, racine, taille):
    chemin = racine / "projects" / "proj-b8" / "dce" / "Attestation_URSSAF.pdf"
    chemin.parent.mkdir(parents=True, exist_ok=True)
    chemin.write_bytes(b"%PDF-1.4\n" + b"x" * taille)
    db.add(ProjectDocument(id="pd-urssaf", project_id="proj-b8", type="autre",
                           file_url="/uploads/projects/proj-b8/dce/Attestation_URSSAF.pdf",
                           file_name="Attestation_URSSAF.pdf"))
    db.commit()
    return chemin


def test_piece_trop_lourde_pour_le_coffre_refusee_sans_lecture(env, monkeypatch):
    from pathlib import Path

    client, db, org, racine, _ = env
    chemin = _piece_du_projet(db, racine, 4096)
    monkeypatch.setattr(uv, "VAULT_MAX_UPLOAD_SIZE", 1024)
    lus = []
    lire = Path.read_bytes
    monkeypatch.setattr(Path, "read_bytes", lambda self: lus.append(self) or lire(self))

    resp = client.post("/api/documents/from-project-doc", json={"project_doc_id": "pd-urssaf"})

    assert resp.status_code == 413, resp.text
    assert resp.json()["detail"] == "Fichier trop volumineux (max 1 Ko)."
    assert chemin not in lus
    assert db.query(DocumentCoffre).count() == 0
    db.expire_all()
    assert db.get(ProjectDocument, "pd-urssaf").vault_prompt_dismissed is False


def test_piece_copiee_au_coffre_en_flux_hors_de_la_boucle(env, monkeypatch):
    from services.file_storage import FileStorage

    client, db, org, racine, _ = env
    chemin = _piece_du_projet(db, racine, 2048)
    lieux = []
    copier = FileStorage._upload_local_stream

    def _sonde(self, flux, cle):
        try:
            asyncio.get_running_loop()
            lieux.append("boucle")
        except RuntimeError:
            lieux.append("thread")
        return copier(self, flux, cle)

    monkeypatch.setattr(FileStorage, "_upload_local_stream", _sonde)

    resp = client.post("/api/documents/from-project-doc", json={"project_doc_id": "pd-urssaf"})

    assert resp.status_code == 200, resp.text
    assert lieux == ["thread"]
    copie = racine / resp.json()["file_url"].removeprefix("/uploads/")
    assert copie.read_bytes() == chemin.read_bytes()
