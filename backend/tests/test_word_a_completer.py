"""Bloc 3 « Le Word dit vrai » — tranche 3 : plus aucun texte d'instruction imprimé.

Avant : le prompt demandait d'écrire « [À COMPLÉTER : insérer la photo de
l'organigramme] » et « [📍 INSÉREZ ICI UNE CARTE… Google Maps…] » ; ces
consignes ressortaient imprimées dans le Word et le PDF.

Maintenant :
  - le prompt ne demande plus ces deux textes (organigramme et carte sont
    dessinés par le code, un visuel absent fait disparaître le bloc) ;
  - un « [À COMPLÉTER : X] » restant (donnée d'entreprise que l'IA n'avait
    pas) devient un passage surligné « À compléter : X » ; la fonction
    d'export renvoie leur nombre et leurs libellés ;
  - même traitement pour un « [📍 … ] » d'un ancien mémoire.

Contenu inventé.
"""
import io

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from services.ai import prompts
from services.docx_exporter import build_memoire_docx, build_memoire_docx_ex

ANCIEN_MARQUEUR_CARTE = (
    "[📍 INSÉREZ ICI UNE CARTE DE LOCALISATION — Capture d'écran Google Maps montrant "
    "l'emplacement de votre entreprise et la distance jusqu'au chantier.]"
)

CONTENU = {
    "preambule": "Engagement fictif. [À COMPLÉTER : nom du chef de chantier désigné]",
    "partie_a": {
        "implantation": "Siège fictif à Exemple-sur-Mer.\n" + ANCIEN_MARQUEUR_CARTE,
        "roles_missions": (
            "- Conducteur : [À COMPLÉTER : nom du conducteur de travaux]\n"
            "- Qualification : [à compléter : diplôme]\n"
            "| Poste | Nom |\n|---|---|\n| Chef | [À COMPLÉTER : nom du chef d'équipe] |\n"
        ),
    },
    "partie_b": {"securite": "Données fictives. [À COMPLÉTER PAR L'ENTREPRISE]"},
    "partie_c": {},
}


def _doc(octets: bytes) -> Document:
    return Document(io.BytesIO(octets))


def _tout_le_texte(doc: Document) -> str:
    morceaux = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for ligne in t.rows:
            morceaux.extend(c.text for c in ligne.cells)
    return "\n".join(morceaux)


def _runs_surlignes(doc: Document) -> list[str]:
    runs = [r for p in doc.paragraphs for r in p.runs]
    for t in doc.tables:
        for ligne in t.rows:
            for c in ligne.cells:
                runs.extend(r for p in c.paragraphs for r in p.runs)
    return [r.text for r in runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]


def test_le_prompt_ne_demande_plus_les_deux_textes_d_instruction():
    texte = "\n".join(str(v) for v in vars(prompts).values() if isinstance(v, str))
    assert "insérer la photo de l'organigramme" not in texte
    assert "INSÉREZ ICI UNE CARTE" not in texte
    assert "Google Maps" not in texte
    # la règle générale « donnée manquante » reste, elle, demandée
    assert "[À COMPLÉTER : description précise" in texte


def test_a_completer_devient_un_passage_surligne_sans_crochets():
    resultat = build_memoire_docx_ex(CONTENU, "Chantier fictif", "ENTREPRISE TEST")
    doc = _doc(resultat.octets)
    texte = _tout_le_texte(doc)
    assert "[" not in texte and "]" not in texte
    assert "COMPLÉTER" not in texte and "📍" not in texte and "Google" not in texte
    surlignes = _runs_surlignes(doc)
    assert "À compléter : nom du chef de chantier désigné" in surlignes
    assert "À compléter : nom du conducteur de travaux" in surlignes
    assert "À compléter : diplôme" in surlignes
    assert "À compléter : nom du chef d'équipe" in surlignes       # dans un tableau
    assert "À compléter" in surlignes                                  # sans libellé


def test_l_export_renvoie_nombre_et_libelles_dans_l_ordre_du_document():
    resultat = build_memoire_docx_ex(CONTENU, "Chantier fictif", "ENTREPRISE TEST")
    libelles = [p.libelle for p in resultat.a_completer]
    assert libelles == [
        "nom du chef de chantier désigné",
        "carte de localisation",                      # ancien « [📍 … ] » sans carte dessinée
        "nom du conducteur de travaux",
        "diplôme",
        "nom du chef d'équipe",
        "",
    ]
    assert resultat.nombre == 6
    assert resultat.a_completer[0].section == "Préambule"
    assert resultat.a_completer[2].section == "6. Rôles et missions de l'équipe d'encadrement"


def test_marqueur_carte_d_un_ancien_memoire_disparait_quand_la_carte_est_dessinee():
    # 1 pixel PNG : le code dessine la carte, l'ancienne consigne n'a plus lieu d'être
    png = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde"
           b"\x00\x00\x00\x0cIDATx\x9cc\xf8\xcf\xc0\x00\x00\x00\x03\x00\x01\x18\xdd\x8d\xb4\x00\x00\x00\x00IEND\xaeB`\x82")
    resultat = build_memoire_docx_ex(CONTENU, "Chantier fictif", "ENTREPRISE TEST", map_image=png, map_caption="Carte")
    assert "carte de localisation" not in [p.libelle for p in resultat.a_completer]
    assert "📍" not in _tout_le_texte(_doc(resultat.octets))


def test_build_memoire_docx_garde_sa_signature_et_renvoie_des_octets():
    octets = build_memoire_docx(CONTENU, "Chantier fictif", "ENTREPRISE TEST")
    assert isinstance(octets, bytes) and octets[:2] == b"PK"


def test_memoire_sans_passage_a_completer():
    resultat = build_memoire_docx_ex({"preambule": "Rien à signaler.", "partie_a": {}}, "P", "E")
    assert resultat.nombre == 0 and resultat.a_completer == []
