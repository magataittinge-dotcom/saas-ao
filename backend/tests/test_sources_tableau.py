"""Bloc 5 (sujet A), tranche 2 — GET /api/sources/{exigence_id}/tableau : la fenêtre du tableur autour
de la ligne citée. Tableurs inventés. Sécurité : autre organisation (404 identique), pièce du même
projet, chemin sous uploads/, tableur piégé, pièce qui n'est pas un tableur, débit."""
import io

import openpyxl
import pytest
from fastapi import Request
from fastapi.testclient import TestClient

from database import SessionLocal, get_db
from main import app
from models.compliance_item import ComplianceItem
from models.organization import Organization
from models.project import Project, ProjectDocument
from models.user import User
from routers import sources
from routers.auth import get_auth_user
from services import ancres_projet, file_storage, pdfium_processus
from services.pdfium_processus import TravailPdfiumEchoue
from tests.pdf_fabrique import pdf_texte


def _dpgf() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lot 05"
    ws.append(["LOT 05 — FAÇADES"])
    ws.append(["N°", "Désignation", "Unité", "Quantité", "Prix unitaire HT"])
    ws.append([1, "Échafaudage de pied", "ens", 1, 4300.0])
    ws.append([2, "Enduit monocouche gratté", "m2", 240, 38.5])
    for i in range(3, 80):
        ws.append([i, f"Ouvrage courant numéro {i}", "u", i, 10])
    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()


@pytest.fixture(autouse=True)
def _propre(tmp_path, monkeypatch):
    pdfium_processus.arreter()
    monkeypatch.setattr(file_storage, "UPLOADS_ROOT", tmp_path / "uploads")
    yield
    pdfium_processus.arreter()


@pytest.fixture
def dossier(db_session, test_org, tmp_path):
    d = tmp_path / "uploads" / "projects" / "p-tab"
    d.mkdir(parents=True)
    (d / "dpgf.xlsx").write_bytes(_dpgf())
    (d / "cctp.pdf").write_bytes(pdf_texte(["Un texte de CCTP assez long pour exister."]))
    db_session.add(Project(id="p-tab", organization_id=test_org.id, name="AO"))
    db_session.add_all([
        ProjectDocument(id="d-dpgf", project_id="p-tab", type="dpgf_template", file_name="dpgf.xlsx",
                        file_url="/uploads/projects/p-tab/dpgf.xlsx"),
        ProjectDocument(id="d-cctp", project_id="p-tab", type="cctp", file_name="cctp.pdf",
                        file_url="/uploads/projects/p-tab/cctp.pdf"),
    ])
    for id_, doc, extrait in (("t1", "dpgf.xlsx", "Échafaudage de pied"),
                              ("t2", "dpgf.xlsx", "Un poste que le bordereau ne contient pas"),
                              ("t3", "dpgf.xlsx", "Ouvrage courant numéro 40"),
                              ("t4", "cctp.pdf", "Un texte de CCTP assez long pour exister.")):
        db_session.add(ComplianceItem(id=id_, project_id="p-tab", exigence_text="x", source_document=doc,
                                      source_page=None, source_excerpt=extrait))
    db_session.commit()
    return d / "dpgf.xlsx"


@pytest.fixture
def en_tant_que_b(db_session):
    db_session.add(Organization(id="org-B", name="B"))
    user_b = User(id="user-B", email="b@b.fr", name="Bob", organization_id="org-B")
    db_session.add(user_b)
    db_session.commit()

    def _db():
        s = SessionLocal()
        try:
            yield s
        finally:
            s.close()

    def _auth(request: Request):
        request.state.rl_key = "org:org-B"
        return user_b

    def basculer() -> TestClient:
        app.dependency_overrides[get_db] = _db
        app.dependency_overrides[get_auth_user] = _auth
        return TestClient(app)
    try:
        yield basculer
    finally:
        app.dependency_overrides.clear()


# ─── Le contenu ───────────────────────────────────────────────────────────────

def test_fenetre_autour_de_la_ligne_citee(client, dossier):
    r = client.get("/api/sources/t1/tableau")
    assert r.status_code == 200
    assert "no-store" in r.headers["cache-control"]
    f = r.json()
    assert f["feuille"] == "Lot 05" and f["ligne"] == 3
    assert f["entete"] == {"n": 2, "cellules": ["N°", "Désignation", "Unité", "Quantité", "Prix unitaire HT"]}
    cibles = [l for l in f["lignes"] if l["cible"]]
    assert len(cibles) == 1 and cibles[0]["n"] == 3
    assert cibles[0]["cellules"] == ["1", "Échafaudage de pied", "ens", "1", "4300"]
    assert f["colonnes"] == ["A", "B", "C", "D", "E"]


def test_fenetre_bornee_30_lignes_avant_et_apres(client, dossier):
    f = client.get("/api/sources/t3/tableau").json()       # « numéro 40 » : ligne 42
    numeros = [l["n"] for l in f["lignes"]]
    assert f["ligne"] == 42 and min(numeros) == 12 and max(numeros) == 72 and len(numeros) == 61
    assert f["entete"]["n"] == 2


def test_premier_clic_calcule_et_range_l_ancre(client, db_session, dossier):
    assert db_session.get(ComplianceItem, "t1").ancre is None
    client.get("/api/sources/t1/tableau")
    db_session.expire_all()
    a = db_session.get(ComplianceItem, "t1").ancre
    assert a["format"] == "tableur" and a["ligne"] == 3 and a["niveau"] == "passage"


def test_la_route_d_ancre_rend_l_ancre_tableur(client, dossier):
    a = client.get("/api/sources/t1").json()
    assert a["format"] == "tableur" and a["feuille"] == "Lot 05" and a["ligne"] == 3 and a["cellules"] == ["B"]


def test_passage_sans_ligne_trouvee_pas_de_tableau(client, dossier):
    a = client.get("/api/sources/t2").json()
    assert a["format"] == "tableur" and a["niveau"] == "aucun"
    assert client.get("/api/sources/t2/tableau").status_code == 404


def test_niveau_page_montre_le_haut_de_la_feuille_sans_surlignage(client, db_session, dossier):
    db_session.get(ComplianceItem, "t2").ancre = {"version": 1, "format": "tableur", "doc_id": "d-dpgf", "feuille": "Lot 05",
                                                  "ligne": None, "cellules": [], "niveau": "page"}
    db_session.commit()
    f = client.get("/api/sources/t2/tableau").json()
    assert f["ligne"] is None and not any(l["cible"] for l in f["lignes"]) and f["lignes"]


def test_piece_non_tableur_404(client, dossier):
    assert client.get("/api/sources/t4/tableau").status_code == 404
    assert client.get("/api/sources/t4").json().get("format") != "tableur"


def test_tableur_illisible_404_jamais_500(client, db_session, dossier):
    client.get("/api/sources/t1")                      # l'ancre est rangée
    dossier.write_bytes(b"ceci n'est pas un tableur")
    assert client.get("/api/sources/t1/tableau").status_code == 404


def test_tableur_piege_bombe_refusee_404(client, db_session, dossier):
    from tests.test_garde_tableurs import _bombe_xlsx

    client.get("/api/sources/t1")
    dossier.write_bytes(_bombe_xlsx())
    assert client.get("/api/sources/t1/tableau").status_code == 404


def test_tableur_supprime_du_disque_404(client, dossier):
    client.get("/api/sources/t1")
    dossier.unlink()
    assert client.get("/api/sources/t1/tableau").status_code == 404


def test_premier_clic_sur_tableur_disparu_ancre_aucune_pas_d_erreur(client, dossier):
    dossier.unlink()
    r = client.get("/api/sources/t1")
    assert r.status_code == 200 and r.json()["format"] == "tableur" and r.json()["niveau"] == "aucun"


# ─── Sécurité ─────────────────────────────────────────────────────────────────

def test_autre_organisation_404_identique_et_rien_lu(client, en_tant_que_b, dossier, monkeypatch):
    assert client.get("/api/sources/t1/tableau").status_code == 200      # l'ancre est rangée pour l'org A
    autre = en_tant_que_b()
    appels = []
    monkeypatch.setattr(pdfium_processus, "executer", lambda *a, **k: appels.append(1))
    r = autre.get("/api/sources/t1/tableau")
    assert r.status_code == 404 and appels == []
    assert r.json() == autre.get("/api/sources/n-existe-pas/tableau").json()      # indistinguable


def test_projet_supprime_inaccessible(client, db_session, dossier):
    from datetime import datetime
    db_session.get(Project, "p-tab").deleted_at = datetime.utcnow()
    db_session.commit()
    assert client.get("/api/sources/t1/tableau").status_code == 404


def test_ancre_pointant_une_piece_d_un_autre_projet_refusee(client, db_session, test_org, dossier):
    db_session.add(Project(id="p-autre", organization_id=test_org.id, name="Autre"))
    db_session.add(ProjectDocument(id="d-autre", project_id="p-autre", type="dpgf_template", file_name="dpgf.xlsx",
                                   file_url="/uploads/projects/p-tab/dpgf.xlsx"))
    db_session.get(ComplianceItem, "t1").ancre = {"version": 1, "format": "tableur", "doc_id": "d-autre",
                                                  "feuille": "Lot 05", "ligne": 3, "cellules": ["B"], "niveau": "passage"}
    db_session.commit()
    assert client.get("/api/sources/t1/tableau").status_code == 404


@pytest.mark.parametrize("url", ["/uploads/../../etc/passwd.xlsx", "/uploads/projects/../../../secret.xlsx",
                                 "/etc/passwd.xlsx", "/uploads/projects/p-tab/../../../../hors.xlsx"])
def test_traversee_de_chemin_refusee(client, db_session, tmp_path, dossier, url):
    (tmp_path / "hors.xlsx").write_bytes(_dpgf())
    (tmp_path / "secret.xlsx").write_bytes(_dpgf())
    db_session.get(ProjectDocument, "d-dpgf").file_url = url
    db_session.get(ComplianceItem, "t1").ancre = {"version": 1, "format": "tableur", "doc_id": "d-dpgf",
                                                  "feuille": "Lot 05", "ligne": 3, "cellules": ["B"], "niveau": "passage"}
    db_session.commit()
    assert client.get("/api/sources/t1/tableau").status_code == 404


def test_le_client_ne_choisit_ni_la_feuille_ni_la_ligne(client, dossier):
    r = client.get("/api/sources/t1/tableau?feuille=Autre&ligne=70&chemin=/etc/passwd")
    assert r.status_code == 200 and r.json()["feuille"] == "Lot 05" and r.json()["ligne"] == 3


def test_non_authentifie_refuse(dossier):
    app.dependency_overrides.clear()
    assert TestClient(app).get("/api/sources/t1/tableau").status_code in (401, 403)


@pytest.mark.parametrize("raison", ["attente", "indisponible"])
def test_pas_de_place_503_retry_after(client, dossier, monkeypatch, raison):
    client.get("/api/sources/t1")            # ancre rangée

    def occupe(*a, **k):
        raise TravailPdfiumEchoue(raison, "")
    monkeypatch.setattr(pdfium_processus, "executer", occupe)
    r = client.get("/api/sources/t1/tableau")
    assert r.status_code == 503 and int(r.headers["retry-after"]) > 0


@pytest.mark.parametrize("raison", ["delai", "memoire", "plantage"])
def test_travail_interrompu_404_pas_de_503_perpetuel(client, dossier, monkeypatch, raison):
    client.get("/api/sources/t1")

    def lent(*a, **k):
        raise TravailPdfiumEchoue(raison, "")
    monkeypatch.setattr(pdfium_processus, "executer", lent)
    assert client.get("/api/sources/t1/tableau").status_code == 404


def test_limite_de_debit(client, dossier):
    from services.rate_limit import limiter
    assert "routers.sources.tableau" in limiter._route_limits
    limiter.enabled = True
    try:
        limiter.reset()
        n = int(sources.LIMITE_TABLEAU.split("/")[0])
        codes = [client.get("/api/sources/t2/tableau").status_code for _ in range(n + 2)]
    finally:
        limiter.enabled = False
        limiter.reset()
    assert 429 in codes[-2:] and 429 not in codes[:n]
