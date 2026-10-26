"""Bloc 3 — revue de sécurité de `memoire_markdown` : pas de ReDoS ni de
récursion non bornée, tableaux plafonnés (rien n'est perdu en silence).

Le texte du mémoire vient de l'IA mais aussi de l'éditeur : il est traité
comme une entrée non fiable. Tout contenu ci-dessous est inventé.
"""
import io
import time

from docx import Document

from services import memoire_markdown as md
from services.docx_exporter import build_memoire_docx


def _texte(blocs) -> str:
    morceaux = []
    for b in blocs:
        if isinstance(b, md.Paragraphe):
            morceaux.append("".join(s.texte for s in b.segments))
    return "\n".join(morceaux)


def test_200_ko_d_etoiles_alternees_se_lisent_en_moins_d_une_seconde():
    entree = "*a " * 70_000                      # ≈ 200 ko sur une seule ligne
    t0 = time.perf_counter()
    blocs = md.analyser(entree)                 # ni RecursionError ni lenteur
    assert time.perf_counter() - t0 < 1.0
    assert blocs


def test_soulignes_et_etoiles_melanges_restent_lineaires():
    for entree in ("_a *b " * 40_000, "**a *b _c " * 20_000, "*" * 200_000, "_" * 200_000, "[" * 200_000,
                   "`" * 100_000, "<b " * 60_000, " " * 200_000 + "#", "# " + " " * 200_000):
        t0 = time.perf_counter()
        md.analyser(entree)
        assert time.perf_counter() - t0 < 1.0, entree[:10]


def test_ligne_tres_longue_ne_perd_aucun_mot():
    mots = [f"mot{i}" for i in range(30_000)]
    blocs = md.analyser(" ".join(mots))
    texte = _texte(blocs)
    assert texte.split() == mots


def test_mise_en_forme_courante_inchangee():
    segs = md.analyser_inline("a **gras** b *ital* c ***les deux*** d 2 * 3 e nom_de_variable f")
    assert [(s.texte, s.gras, s.italique) for s in segs] == [
        ("a ", False, False), ("gras", True, False), (" b ", False, False), ("ital", False, True),
        (" c ", False, False), ("les deux", True, True),
        (" d 2 * 3 e nom_de_variable f", False, False),
    ]


def _grand_tableau(lignes: int, colonnes: int) -> str:
    entete = "| " + " | ".join(f"C{j}" for j in range(colonnes)) + " |"
    sep = "|" + "---|" * colonnes
    corps = ["| " + " | ".join(f"v{i}-{j}" for j in range(colonnes)) + " |" for i in range(lignes)]
    return "\n".join([entete, sep, *corps])


def test_tableau_1000_lignes_100_colonnes_est_borne_et_garde_son_contenu():
    t0 = time.perf_counter()
    blocs = md.analyser(_grand_tableau(1000, 100))
    assert time.perf_counter() - t0 < 5.0
    assert not [b for b in blocs if isinstance(b, md.Tableau)]   # trop large : lignes de texte
    texte = _texte(blocs)
    assert "v999-99" in texte and "v0-0" in texte                # rien n'est perdu


def test_tableau_trop_long_garde_les_premieres_lignes_et_le_reste_en_paragraphes():
    blocs = md.analyser(_grand_tableau(1000, 3))
    tableaux = [b for b in blocs if isinstance(b, md.Tableau)]
    assert len(tableaux) == 1 and len(tableaux[0].lignes) == md.MAX_LIGNES_TABLEAU == 500
    assert "v999-2" in _texte(blocs) and "v500-0" in _texte(blocs)
    assert "v499-0" not in _texte(blocs)


def test_cellules_en_trop_ne_sont_pas_perdues():
    blocs = md.analyser("| A | B |\n|---|---|\n| 1 | 2 | 3 | 4 |\n")
    tableau = blocs[0]
    assert isinstance(tableau, md.Tableau) and len(tableau.lignes[0]) == 2
    derniere = "".join(s.texte for s in tableau.lignes[0][1])
    assert "2" in derniere and "3" in derniere and "4" in derniere


def test_tableau_de_500_lignes_s_exporte_en_temps_borne():
    contenu = {"preambule": _grand_tableau(500, md.MAX_COLONNES_TABLEAU), "partie_a": {}}
    t0 = time.perf_counter()
    doc = Document(io.BytesIO(build_memoire_docx(contenu, "P", "E")))
    assert time.perf_counter() - t0 < 30.0
    assert len(doc.tables[0].rows) == 501


def test_grand_tableau_exporte_sans_tableau_word_mais_sans_perte():
    contenu = {"preambule": _grand_tableau(1000, 100), "partie_a": {}}
    t0 = time.perf_counter()
    doc = Document(io.BytesIO(build_memoire_docx(contenu, "P", "E")))
    assert time.perf_counter() - t0 < 30.0
    assert not doc.tables
    assert "v999-99" in "\n".join(p.text for p in doc.paragraphs)
