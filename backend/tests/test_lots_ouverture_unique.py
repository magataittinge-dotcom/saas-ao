"""Détection des lots plus rapide, résultat inchangé (nuit du 23/10).

- Un classeur DPGF n'est ouvert qu'UNE fois (avant : 3 lectures complètes,
  44 s sur les 6 DPGF de Gueux pendant l'extraction de fond).
- La progression n'est plus signalée quand rien n'a changé (avant : un rappel
  par fichier au scan des noms, chacun écrit en base).
"""
from types import SimpleNamespace

import openpyxl

import services.lot_detector as ld


def _classeur(chemin):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lot 05 — Façade ITE"
    ws.append(["Désignation", "Tranche ferme : école", "Tranche optionnelle : salle"])
    ws.append(["Isolant laine de roche 140 mm", "m²", 1250])
    ws2 = wb.create_sheet("Page de garde")
    ws2.append(["Lot 05 — Revêtement de façade – ITE"])
    wb.save(chemin)


def test_un_classeur_dpgf_est_ouvert_une_seule_fois(tmp_path, monkeypatch):
    chemin = tmp_path / "DPGF lot 05.xlsx"
    _classeur(chemin)
    ouvertures = []
    vrai = openpyxl.load_workbook

    def compter(*a, **k):
        ouvertures.append(a[0] if a else k.get("filename"))
        return vrai(*a, **k)

    monkeypatch.setattr(openpyxl, "load_workbook", compter)
    lots = ld.detect_lots_from_excel(str(chemin))

    assert len(ouvertures) == 1
    # les onglets sans lot donnent des lots « lot_sheet_* », écartés ensuite par la fusion
    reels = [l for l in lots if not l["id"].startswith("lot_sheet_")]
    assert [l["id"] for l in reels] == ["lot5"]
    assert reels[0]["tranches"] == ["Tranche ferme : école", "Tranche optionnelle : salle"]


def test_classeur_ouvert_par_l_appelant_meme_resultat(tmp_path):
    chemin = tmp_path / "DPGF.xlsx"
    _classeur(chemin)
    seul = ld._detect_tranches_in_excel(str(chemin))
    wb = openpyxl.load_workbook(str(chemin), read_only=True, data_only=True)
    try:
        partage = ld._detect_tranches_in_excel(str(chemin), wb=wb)
        cellules = ld._scan_excel_cell_content_for_lots(str(chemin), wb.sheetnames, wb=wb)
    finally:
        wb.close()
    assert partage == seul
    assert [d.id for d in cellules] == [d.id for d in ld._scan_excel_cell_content_for_lots(str(chemin), ["Lot 05 — Façade ITE", "Page de garde"])]


def test_progression_signalee_seulement_quand_elle_change():
    docs = [SimpleNamespace(file_name=f"plan-{i:03d}.dwg", type="autre", extracted_text="", file_url=None, related_lots=None)
            for i in range(93)]
    rappels = []
    ld.detect_all_lots(docs, uploads_root=None, on_progress=lambda pct, detail: rappels.append((pct, detail)))
    # 93 fichiers au scan des noms → 6 pourcentages distincts (0 à 5 %), pas 93
    assert len(rappels) <= 12, rappels
    assert len(set(rappels)) == len(rappels)
    assert rappels[-1] == (95, "Consolidation des lots...")
