"""Ancre d'une exigence dans un tableur (DPGF, BPU, DQE) — sujet A, bloc 5 (« Voir dans le DCE »).

Un tableur n'a pas de PDF : l'écran montre un tableau HTML dont la ligne citée est surlignée.
L'ancre, un petit JSON rangé avec l'exigence (``compliance_items.ancre``) ::

    {"version": 1, "format": "tableur", "doc_id": "...", "feuille": "Lot 05", "ligne": 12,
     "cellules": ["B"], "niveau": "passage" | "page" | "aucun"}

``ligne`` : numéro de ligne du tableur (à partir de 1) ; ``cellules`` : lettres des colonnes
qui portent le passage. Niveau « passage » : la ligne est surlignée. « page » : la feuille
est citée, sans surlignage (``ligne`` et ``cellules`` vides). « aucun » : rien à montrer.

Le texte que l'analyse a lu est celui de ``DocumentProcessor._extract_xlsx`` (cellules
non vides d'une ligne, séparées par une tabulation, feuilles à la suite). Ici, la même
lecture donne à chaque cellule sa position, et la comparaison se fait sur le flux
alphanumérique (lettres et chiffres sans accents ni casse, comme ``pdf_ancre``).

Règles d'honnêteté : jamais de devinette.
- On ne surligne une ligne que si le passage cité s'y retrouve EN ENTIER (d'une cellule à
  l'autre si l'extrait en couvre plusieurs) et qu'aucune autre ligne du classeur ne le porte.
- Un passage de 4 à 11 lettres (un code, une unité) n'est retenu que s'il est le contenu
  entier d'une cellule ; en dessous de 4 lettres, rien.
- Passage sur plusieurs lignes du classeur, ou lecture arrêtée par les bornes : niveau
  « page » si toutes les occurrences sont dans une seule feuille, sinon « aucun ».
- Aucun calcul, aucun commentaire : les montants sont des cellules comme les autres.

Bornes (tableur piégé) : le fichier passe la garde anti-bombe (``garde_archives``) avant
ouverture, openpyxl lit au fil de l'eau (``read_only``), au plus ``MAX_LIGNES`` lignes et
``MAX_COLONNES`` colonnes par feuille et ``MAX_CELLULES`` cellules par classeur. Le travail
s'exécute dans le processus plafonné en mémoire et à délai (``pdfium_processus``).
"""
from __future__ import annotations

import datetime
import re
import zipfile
from pathlib import Path
from typing import Any, Iterator, NamedTuple, Optional, Sequence

from services.garde_archives import controler_tableur
from services.pdf_ancre import MIN_FLUX_MOT, MIN_FLUX_PASSAGE
from services.pdf_positions import forme_flux

__all__ = [
    "EXTENSIONS", "FORMAT", "MAX_CELLULES", "MAX_COLONNES", "MAX_LIGNES", "VERSION",
    "ancre_aucune", "calculer_ancres", "est_tableur", "fenetre", "lettre",
]

VERSION = 1
FORMAT = "tableur"
EXTENSIONS = (".xlsx", ".xlsm", ".xls", ".ods")

MAX_LIGNES = 20_000          # lignes lues par feuille
MAX_COLONNES = 64            # colonnes lues par ligne
MAX_CELLULES = 300_000       # cellules non vides lues par classeur
_LARGEUR_LUE = 256           # colonnes regardées pour savoir si la borne de colonnes coupe des données (limite du .xls)

_ZONE_ENTETE = 40            # l'en-tête du tableau est cherché dans les 40 premières lignes
_TAILLE_CELLULE = 300        # caractères d'une cellule renvoyée à l'écran


def est_tableur(nom_fichier: Optional[str]) -> bool:
    return (nom_fichier or "").lower().endswith(EXTENSIONS)


def lettre(index: int) -> str:
    """Lettre de colonne d'un index à partir de 0 : 0 → « A », 26 → « AA »."""
    index += 1
    sortie = ""
    while index:
        index, reste = divmod(index - 1, 26)
        sortie = chr(65 + reste) + sortie
    return sortie


def ancre_aucune(doc_id: Optional[str]) -> dict:
    return {"version": VERSION, "format": FORMAT, "doc_id": doc_id, "feuille": None,
            "ligne": None, "cellules": [], "niveau": "aucun"}


def _ancre_page(doc_id: str, feuille: str) -> dict:
    return {"version": VERSION, "format": FORMAT, "doc_id": doc_id, "feuille": feuille,
            "ligne": None, "cellules": [], "niveau": "page"}


def _ancre_passage(doc_id: str, feuille: str, ligne: int, cellules: list[str]) -> dict:
    return {"version": VERSION, "format": FORMAT, "doc_id": doc_id, "feuille": feuille,
            "ligne": ligne, "cellules": cellules, "niveau": "passage"}


# ─── Lecture bornée ───────────────────────────────────────────────────────────

class _Bilan:
    """Ce que la lecture a vu : ``tronque`` si une borne a coupé des données non vides."""

    def __init__(self) -> None:
        self.cellules = 0
        self.tronque = False


def _vide(valeur: Any) -> bool:
    return valeur is None or (isinstance(valeur, str) and not valeur.strip())


def _lignes_bornees(brutes: Iterator[Sequence], bilan: _Bilan) -> Iterator[tuple[int, list]]:
    """(numéro à partir de 1, valeurs) ; s'arrête aux bornes de lignes, de colonnes et de cellules.
    ``brutes`` fournit une ligne de plus que la borne et ``_LARGEUR_LUE`` colonnes, pour savoir si
    la borne a coupé des données."""
    for numero, ligne in enumerate(brutes, 1):
        valeurs = list(ligne)
        if numero > MAX_LIGNES:
            if any(not _vide(v) for v in valeurs):
                bilan.tronque = True
            return
        if len(valeurs) > MAX_COLONNES:
            if any(not _vide(v) for v in valeurs[MAX_COLONNES:]):
                bilan.tronque = True
            del valeurs[MAX_COLONNES:]
        bilan.cellules += sum(1 for v in valeurs if not _vide(v))
        if bilan.cellules > MAX_CELLULES:
            bilan.tronque = True
            return
        yield numero, valeurs


def _ouvrir_xlsx(chemin: Path, formules: bool, seule: Optional[str], bilan: _Bilan):
    import openpyxl

    # un objet fichier, jamais un nom : openpyxl refuse un nom en .xls (un .xlsx renommé)
    with open(chemin, "rb") as fichier:
        classeur = openpyxl.load_workbook(fichier, read_only=True, data_only=not formules)
        try:
            for feuille in classeur.worksheets:
                if seule is not None and feuille.title != seule:
                    continue
                if (feuille.max_row or 0) > MAX_LIGNES:     # dimension déclarée par le fichier
                    bilan.tronque = True
                brutes = feuille.iter_rows(min_row=1, max_row=MAX_LIGNES + 1, max_col=_LARGEUR_LUE,
                                           values_only=True)
                yield feuille.title, _lignes_bornees(brutes, bilan)
        finally:
            classeur.close()


def _ouvrir_xls(chemin: Path, formules: bool, seule: Optional[str], bilan: _Bilan):
    import xlrd

    try:
        classeur = xlrd.open_workbook(str(chemin), on_demand=True)
    except xlrd.biffh.XLRDError:
        if zipfile.is_zipfile(chemin):          # un .xlsx renommé en .xls
            yield from _ouvrir_xlsx(chemin, formules, seule, bilan)
            return
        raise
    try:
        for i in range(classeur.nsheets):
            feuille = classeur.sheet_by_index(i)
            if seule is None or feuille.name == seule:
                if feuille.nrows > MAX_LIGNES:
                    bilan.tronque = True
                brutes = (feuille.row_values(r, 0, _LARGEUR_LUE)
                          for r in range(min(feuille.nrows, MAX_LIGNES + 1)))
                yield feuille.name, _lignes_bornees(brutes, bilan)
            classeur.unload_sheet(i)
    finally:
        classeur.release_resources()


def _ouvrir_ods(chemin: Path, formules: bool, seule: Optional[str], bilan: _Bilan):
    from services.lecture_ods import _MAX_LIGNES_ODS, onglets_ods

    for nom, lignes in onglets_ods(chemin, lignes_vides_exactes=True):
        if seule is None or nom == seule:
            lues = list(lignes)
            if len(lues) >= _MAX_LIGNES_ODS:      # le lecteur s'est arrêté à sa borne
                bilan.tronque = True
            yield nom, _lignes_bornees(iter(lues), bilan)


def _feuilles(chemin, formules: bool, seule: Optional[str], bilan: _Bilan):
    """(nom, lignes numérotées) de chaque feuille (ou de la feuille ``seule``)."""
    chemin = Path(chemin)
    controler_tableur(chemin)       # bombe de décompression : jamais ouverte
    suffixe = chemin.suffix.lower()
    ouvrir = {".xlsx": _ouvrir_xlsx, ".xlsm": _ouvrir_xlsx, ".xls": _ouvrir_xls,
              ".ods": _ouvrir_ods}.get(suffixe)
    if ouvrir is None:
        raise ValueError("format de tableur non pris en charge")
    yield from ouvrir(chemin, formules, seule, bilan)


# ─── Index des lignes ─────────────────────────────────────────────────────────

class _Ligne(NamedTuple):
    feuille: int                              # index dans ``_Index.noms``
    numero: int                               # numéro de ligne du tableur
    flux: str                                 # flux alphanumérique des cellules, à la suite
    cellules: tuple[tuple[int, int, int], ...]   # (colonne, début, fin exclu) dans ``flux``


class _Index(NamedTuple):
    noms: list[str]
    lignes: list[_Ligne]
    tronque: bool


def _texte(valeur: Any) -> str:
    """Le texte d'une cellule tel que l'extraction du DCE l'a donné à l'analyse (``str``)."""
    return str(valeur)


def _indexer(chemin) -> _Index:
    bilan = _Bilan()
    noms: list[str] = []
    lignes: list[_Ligne] = []
    for nom, rangees in _feuilles(chemin, True, None, bilan):
        noms.append(nom)
        for numero, valeurs in rangees:
            morceaux: list[str] = []
            cellules: list[tuple[int, int, int]] = []
            debut = 0
            for colonne, valeur in enumerate(valeurs):
                if valeur is None:
                    continue
                flux, _ = forme_flux(_texte(valeur))
                if flux:
                    cellules.append((colonne, debut, debut + len(flux)))
                    morceaux.append(flux)
                    debut += len(flux)
            if morceaux:
                lignes.append(_Ligne(len(noms) - 1, numero, "".join(morceaux), tuple(cellules)))
    return _Index(noms, lignes, bilan.tronque)


# ─── Recherche ────────────────────────────────────────────────────────────────

def _occurrences(index: _Index, flux: str) -> list[tuple[_Ligne, list[str]]]:
    """Les lignes qui portent ``flux`` en entier, avec les colonnes concernées."""
    trouvees: list[tuple[_Ligne, list[str]]] = []
    court = len(flux) < MIN_FLUX_PASSAGE
    for ligne in index.lignes:
        if court:
            # un code, une unité : seulement le contenu entier d'une cellule
            colonnes = [lettre(c) for c, d, f in ligne.cellules if ligne.flux[d:f] == flux]
            if colonnes:
                trouvees.append((ligne, colonnes))
            continue
        k = ligne.flux.find(flux)
        if k < 0:
            continue
        fin = k + len(flux)
        trouvees.append((ligne, [lettre(c) for c, d, f in ligne.cellules if d < fin and f > k]))
    return trouvees


def _ancre_de(index: _Index, doc_id: str, extrait: Optional[str]) -> dict:
    flux, _ = forme_flux(extrait or "")
    if len(flux) < MIN_FLUX_MOT:
        return ancre_aucune(doc_id)
    trouvees = _occurrences(index, flux)
    if len(trouvees) == 1 and not index.tronque:
        ligne, colonnes = trouvees[0]
        return _ancre_passage(doc_id, index.noms[ligne.feuille], ligne.numero, colonnes)
    if trouvees:
        feuilles = {ligne.feuille for ligne, _ in trouvees}
        if len(feuilles) == 1:
            return _ancre_page(doc_id, index.noms[feuilles.pop()])
        return ancre_aucune(doc_id)
    # rien d'un seul tenant : un extrait de plusieurs lignes du tableur, toutes dans la même feuille
    segments = [s for s in re.split(r"[\r\n]+", extrait or "") if len(forme_flux(s)[0]) >= MIN_FLUX_PASSAGE]
    if len(segments) >= 2:
        feuilles_communes = None
        for segment in segments:
            presentes = {ligne.feuille for ligne, _ in _occurrences(index, forme_flux(segment)[0])}
            feuilles_communes = presentes if feuilles_communes is None else feuilles_communes & presentes
        if feuilles_communes and len(feuilles_communes) == 1:
            return _ancre_page(doc_id, index.noms[next(iter(feuilles_communes))])
    return ancre_aucune(doc_id)


def calculer_ancres(chemin: str, doc_id: str, demandes: Sequence[dict]) -> dict[str, Optional[dict]]:
    """Les ancres de toutes les exigences d'un tableur, en UNE lecture du fichier (c'est la
    fonction qu'exécute le processus plafonné). ``demandes`` : des ``{"id", "extrait", ...}``.
    Rend ``{id: ancre}``. Une exception de lecture (fichier illisible, bombe) est relancée."""
    index = _indexer(chemin)
    resultat: dict[str, Optional[dict]] = {}
    for demande in demandes:
        try:
            resultat[demande["id"]] = _ancre_de(index, doc_id, demande.get("extrait"))
        except Exception:
            resultat[demande["id"]] = None
    return resultat


# ─── La fenêtre affichée ──────────────────────────────────────────────────────

_MOTS_ENTETE = re.compile(
    r"\b(?:designation|designations|descriptif|description|libelle|intitule|ouvrage|prestation"
    r"|unite|u|quantite|qte|qt|prix|montant|total|ref|reference|article|poste|n)\b")


def _afficher(valeur: Any) -> str:
    """La valeur d'une cellule en texte, telle que le tableur la montre (jamais de calcul)."""
    if valeur is None:
        return ""
    if isinstance(valeur, bool):
        return "Vrai" if valeur else "Faux"
    if isinstance(valeur, float):
        texte = str(int(valeur)) if valeur.is_integer() and abs(valeur) < 1e15 else format(valeur, ".10g")
    elif isinstance(valeur, datetime.datetime):
        texte = valeur.date().isoformat() if valeur.time() == datetime.time(0) else valeur.isoformat(" ")
    elif isinstance(valeur, datetime.date):
        texte = valeur.isoformat()
    else:
        texte = str(valeur)
    texte = texte.strip()
    return texte if len(texte) <= _TAILLE_CELLULE else texte[:_TAILLE_CELLULE - 1] + "…"


def _est_entete(valeurs: Sequence) -> bool:
    """Une ligne d'en-tête : au moins deux cellules, toutes du texte, dont un mot d'en-tête."""
    textes = [v for v in valeurs if not _vide(v)]
    if len(textes) < 2 or not all(isinstance(v, str) for v in textes):
        return False
    norm = [" ".join(re.sub(r"[^a-z0-9]+", " ", _sans_accents(v)).split()) for v in textes]
    return any(_MOTS_ENTETE.search(n) for n in norm)


def _sans_accents(texte: str) -> str:
    import unicodedata

    nfkd = unicodedata.normalize("NFKD", texte.lower())
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def fenetre(chemin: str, feuille: str, ligne: Optional[int], avant: int = 30, apres: int = 30,
            max_colonnes: int = 16) -> Optional[dict]:
    """Une fenêtre du tableur autour de ``ligne`` (ou le haut de la feuille si ``ligne`` est None),
    en texte : ``{"feuille", "ligne", "colonnes": ["A", ...], "entete": {"n", "cellules"} | None,
    "lignes": [{"n", "cellules", "cible"}], "colonnes_tronquees": bool}``.

    Les colonnes sans aucune valeur dans la fenêtre sont retirées ; au-delà de ``max_colonnes``,
    les colonnes de la cellule citée sont gardées d'abord, puis les plus à gauche. Valeurs lues
    telles que le fichier les a enregistrées (``data_only`` : pas de formule recalculée).
    None si la feuille n'existe pas. Lève ce que lève la lecture du fichier."""
    bas = max(1, ligne - avant) if ligne else 1
    haut = (ligne + apres) if ligne else 1 + avant + apres
    bilan = _Bilan()
    zone: dict[int, list] = {}
    entete: Optional[tuple[int, list]] = None
    trouvee = False
    for nom, rangees in _feuilles(chemin, False, feuille, bilan):
        trouvee = True
        for numero, valeurs in rangees:
            if numero > max(haut, _ZONE_ENTETE):
                break
            if entete is None and numero <= _ZONE_ENTETE and (not ligne or numero < ligne) \
                    and _est_entete(valeurs):
                entete = (numero, valeurs)
            if bas <= numero <= haut and any(not _vide(v) for v in valeurs):
                zone[numero] = valeurs
    if not trouvee:
        return None
    if entete is not None:
        zone.pop(entete[0], None)
    lignes_brutes = [entete[1]] if entete else []
    lignes_brutes += list(zone.values())
    utilisees = sorted({c for valeurs in lignes_brutes for c, v in enumerate(valeurs) if not _vide(v)})
    tronquees = len(utilisees) > max_colonnes
    if tronquees:
        citees = sorted(c for c, v in enumerate(zone.get(ligne, [])) if not _vide(v))[:max_colonnes] if ligne else []
        autres = [c for c in utilisees if c not in citees][:max_colonnes - len(citees)]
        utilisees = sorted(citees + autres)

    def cellules(valeurs: Sequence) -> list[str]:
        return [_afficher(valeurs[c]) if c < len(valeurs) else "" for c in utilisees]

    return {
        "feuille": feuille,
        "ligne": ligne,
        "colonnes": [lettre(c) for c in utilisees],
        "entete": {"n": entete[0], "cellules": cellules(entete[1])} if entete else None,
        "lignes": [{"n": n, "cellules": cellules(v), "cible": n == ligne} for n, v in sorted(zone.items())],
        "colonnes_tronquees": tronquees,
    }
