import io
import logging
import os
import posixpath
import zipfile
from pathlib import Path
from typing import Iterator, List, Optional, Tuple, Union

from services.pdf_lecture import ERREUR_MOT_DE_PASSE, PdfIllisible, lire_document

from services.garde_archives import controler_archive_office, controler_tableur

logger = logging.getLogger(__name__)

# Une pièce : ses octets, ou le chemin du fichier déposé. Par le chemin, un gros PDF (au-delà
# de 16 Mio, voir pdf_lecture), un .docx ou un .xlsx est lu au fil de l'eau, sans être chargé
# en entier (dépôts de plusieurs centaines de Mo, extraits à 8 en parallèle, serveur de 4 Go).
Source = Union[bytes, str, "os.PathLike[str]"]

# Seuil « document scanné » : moins de N caractères par page lue en moyenne
# = pas de couche texte exploitable (PDF image). Déterministe, testé.
_SCANNED_AVG_CHARS_PER_PAGE = 15

_WARN_SCANNED = (
    "Document scanné (aucune couche texte lisible) — "
    "exigences non extraites de ce fichier"
)
_WARN_UNREADABLE = (
    "Document non lisible (fichier corrompu ou format non reconnu) — "
    "texte non extrait de ce fichier"
)
_WARN_DOC = (
    "Format .doc (ancien Word) non pris en charge — texte non extrait : "
    "convertissez ce fichier en PDF ou DOCX"
)
_WARN_PASSWORD = (
    "PDF protégé par un mot de passe — texte non extrait : "
    "déposez une version sans mot de passe"
)
_PAGES_ILLISIBLES_CITEES = 8


def _warn_pages_illisibles(numeros: Tuple[int, ...]) -> str:
    """« Page 12 illisible — … » : les pages lues qui ne se chargent pas, nommées."""
    suite = " — le texte des autres pages est extrait"
    if len(numeros) == 1:
        return f"Page {numeros[0]} illisible{suite}"
    if len(numeros) <= _PAGES_ILLISIBLES_CITEES:
        liste = ", ".join(str(n) for n in numeros[:-1]) + f" et {numeros[-1]}"
        return f"Pages {liste} illisibles{suite}"
    liste = ", ".join(str(n) for n in numeros[:_PAGES_ILLISIBLES_CITEES])
    return f"Pages {liste}… illisibles ({len(numeros)} pages){suite}"


class DocumentProcessor:
    """Extract text from PDF, DOCX, XLSX/XLS (BIFF inclus) files."""

    def extract(
        self,
        content: Source,
        filename: str,
        max_pages: Optional[int] = None,
        on_page=None,
    ) -> Tuple[str, Optional[int]]:
        """Contrat historique : (extracted_text, page_count).

        Always returns a string (never None) — empty string on failure."""
        text, pages, _warning = self.extract_ex(content, filename, max_pages, on_page)
        return text, pages

    def extract_ex(
        self,
        content: Source,
        filename: str,
        max_pages: Optional[int] = None,
        on_page=None,
        on_pages=None,
    ) -> Tuple[str, Optional[int], Optional[str]]:
        """(texte, pages, warning) — plus AUCUN échec d'extraction silencieux.

        content: les octets de la pièce, ou le chemin du fichier (préféré : lu au fil de
        l'eau, jamais chargé en entier).
        Le warning est un message UI PAR FICHIER (audit écart #3) : renseigné
        dès qu'un format censé porter du texte n'en a pas livré (scanné,
        corrompu, protégé par un mot de passe, pages illisibles, .doc ancien).
        None = extraction saine. Ne lève jamais.

        max_pages: if set, only extract text from the first N pages (PDF only).
        on_page: optional callback (pages_done, total_pages) — progression
        RÉELLE intra-document pour les gros PDFs (barre d'upload).
        on_pages: optional callback (textes de TOUTES les pages, vides comprises)
        appelé quand un PDF a été lu en entier — l'ancrage des sources les
        reprend au lieu de relire le fichier."""
        lower = filename.lower()

        # .doc (ancien Word, conteneur OLE2) : accepté à l'upload mais
        # inextrayable sans LibreOffice — signalé, plus jamais silencieux.
        if lower.endswith(".doc"):
            return "", None, _WARN_DOC

        # Types jamais porteurs de texte (images, plans…) → pas de warning.
        if not any(lower.endswith(ext) for ext in (".pdf", ".docx", ".xlsx", ".xls", ".txt", ".ods")):
            return "", None, None

        try:
            text, pages, illisibles = self._do_extract(content, lower, max_pages, on_page, on_pages)
        except PdfIllisible as e:
            logger.warning(f"Extraction failed for {filename}: {e}")
            return "", None, _WARN_PASSWORD if e.code == ERREUR_MOT_DE_PASSE else _WARN_UNREADABLE
        except Exception as e:
            logger.warning(f"Extraction failed for {filename}: {e}")
            return "", None, _WARN_UNREADABLE

        if lower.endswith(".pdf"):
            # Un PDF sans aucune page n'a rien livré : non lisible (plan H §6, n° 1 —
            # PyMuPDF rendait 0 page sans avertissement pour un PDF tronqué).
            if not pages:
                return "", None, _WARN_UNREADABLE
            # PDF « scanné » : des pages mais quasi aucun texte → warning explicite.
            pages_read = min(max_pages, pages) if max_pages else pages
            if len(text.strip()) < _SCANNED_AVG_CHARS_PER_PAGE * max(pages_read, 1):
                return text, pages, _WARN_SCANNED
            if illisibles:
                return text, pages, _warn_pages_illisibles(illisibles)

        return text, pages, None

    def _do_extract(
        self, content: Source, lower: str, max_pages: Optional[int], on_page=None, on_pages=None
    ) -> Tuple[str, Optional[int], Tuple[int, ...]]:
        """(texte, pages, numéros des pages illisibles)."""
        if lower.endswith(".pdf"):
            return self._extract_pdf(content, max_pages, on_page, on_pages)
        elif lower.endswith(".docx"):
            return texte_docx(content), None, ()
        elif lower.endswith(".xls"):
            return self._extract_xls_biff(content), None, ()
        elif lower.endswith((".xlsx", ".ods")):
            return self._extract_xlsx(content), None, ()
        elif lower.endswith(".txt"):
            return _octets(content).decode("utf-8", errors="replace"), None, ()
        return "", None, ()

    def _extract_pdf(
        self, content: Source, max_pages: Optional[int] = None, on_page=None, on_pages=None
    ) -> Tuple[str, int, Tuple[int, ...]]:
        """PDFium (services/pdf_lecture) : texte au format de PyMuPDF, page par page."""
        def progression(fait: int, a_lire: int) -> None:
            try:
                on_page(fait, a_lire)
            except Exception:
                pass  # la progression ne casse jamais l'extraction

        lu = lire_document(content, max_pages=max_pages,
                           on_page=progression if on_page is not None else None)
        if on_pages is not None and len(lu.pages) == lu.total:
            try:
                on_pages(lu.pages)   # une entrée par page, vides comprises (numéros exacts)
            except Exception:
                pass  # le cache de l'ancrage ne casse jamais l'extraction
        return "\n\n".join(t for t in lu.pages if t), lu.total, lu.illisibles

    def _extract_xls_biff(self, content: Source) -> str:
        """Vieux format Excel (BIFF, avant 2007) — xlrd (audit risque #1 :
        openpyxl le rejetait en silence). Fallback openpyxl pour les .xls
        qui sont en réalité des xlsx renommés."""
        controler_tableur(content)   # bombe de décompression : jamais ouverte
        try:
            import xlrd

            wb = xlrd.open_workbook(file_contents=_octets(content))
            lines = []
            for sheet in wb.sheets():
                for rx in range(sheet.nrows):
                    row_text = "\t".join(
                        str(c.value) for c in sheet.row(rx)
                        if c.value not in (None, "")
                    )
                    if row_text.strip():
                        lines.append(row_text)
            return "\n".join(lines)
        except Exception:
            return self._extract_xlsx(content)

    def _extract_xlsx(self, content: Source) -> str:
        import openpyxl

        controler_tableur(content)   # bombe de décompression : jamais ouverte (garde W, bloc 1)
        # Un objet fichier, jamais un nom : openpyxl refuse un nom en .xls (le repli des
        # .xls qui sont des .xlsx renommés). read_only : lu au fil de l'eau.
        if isinstance(content, (bytes, bytearray)):
            fichier = io.BytesIO(content)
        else:
            fichier = open(content, "rb")
        try:
            wb = openpyxl.load_workbook(fichier, read_only=True)
            try:
                lines = []
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows(values_only=True):
                        row_text = "\t".join(str(c) for c in row if c is not None)
                        if row_text.strip():
                            lines.append(row_text)
                return "\n".join(lines)
            finally:
                wb.close()
        finally:
            fichier.close()


def _octets(content: Source) -> bytes:
    """Les octets d'une pièce (formats lus en entier : .xls, .txt)."""
    if isinstance(content, (bytes, bytearray)):
        return bytes(content)
    return Path(content).read_bytes()


# ─── Word (.docx) : le corps dans l'ordre, tableaux compris (B3) ────────────────
#
# python-docx ne lisait que `doc.paragraphs` : ni les tableaux (critères pondérés,
# liste des pièces, tableau des lots, cellules d'un cadre de réponse), ni les
# contrôles de contenu, ni les insertions suivies, ni les zones de texte. Il
# charge aussi toutes les parties du paquet en mémoire, images comprises.
#
# Ici, seule la partie principale est lue, directement dans l'archive, bloc par
# bloc (lxml.iterparse) : la mémoire reste celle d'un paragraphe ou d'un
# tableau, quelle que soit la taille du fichier.
#
# Forme du texte (la même pour l'extraction au dépôt et pour lot_detector) :
#   - un paragraphe = une ligne ; les paragraphes vides sont sautés ;
#   - une ligne de tableau = une ligne, cellules séparées par une tabulation
#     (comme les tableurs ci-dessus) ; une cellule fusionnée n'est lue qu'une
#     fois ; une ligne à une seule cellule (encadré de formulaire) garde ses
#     paragraphes, un par ligne ;
#   - un tableau imbriqué ou une zone de texte suit la ligne qui le porte.

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_CORPS = _W + "body"
_PARAGRAPHE = _W + "p"
_TABLEAU = _W + "tbl"
_LIGNE = _W + "tr"
_CELLULE = _W + "tc"
_PROPRIETES_CELLULE = _W + "tcPr"
_FUSION_VERTICALE = _W + "vMerge"
_FUSION_HORIZONTALE = _W + "hMerge"
_CONTROLE = _W + "sdt"
_CONTENU_CONTROLE = _W + "sdtContent"
_XML_PERSO = _W + "customXml"
_ZONE_DE_TEXTE = _W + "txbxContent"
_TEXTE = _W + "t"
_TABULATIONS = frozenset({_W + "tab", _W + "ptab"})
_SAUT = _W + "br"
_RETOUR = _W + "cr"
_TRAIT_INSECABLE = _W + "noBreakHyphen"
_VAL = _W + "val"
_TYPE = _W + "type"

# Sous-arbres sans texte affiché : propriétés (dont les taquets de tabulation de
# w:pPr), texte supprimé ou déplacé ailleurs (révisions suivies), codes de
# champ, et la copie de repli d'un objet (mc:Fallback répète la zone de texte
# déjà lue dans mc:Choice).
_SANS_TEXTE = frozenset({
    _W + "pPr", _W + "rPr", _W + "sdtPr", _W + "sdtEndPr",
    _W + "del", _W + "moveFrom", _W + "instrText", _W + "delText",
    "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback",
})

_RELATION_DOCUMENT = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument"
)

# Taille décompressée maximale de la partie principale. Mesuré sur de vrais
# formulaires Word : 11 à 26 octets de XML par caractère de texte ; un CCTP de
# 500 pages (≈ 1,5 million de caractères) ferait donc ≈ 25 Mo. Au-delà de
# 128 Mo, c'est une bombe de décompression (quelques Mo dans l'archive, des Go
# une fois ouverts) : le fichier est signalé non lisible.
_DOCX_PARTIE_MAX = 128 * 1024 * 1024
_DOCX_RELATIONS_MAX = 1024 * 1024


def texte_docx(source) -> str:
    """Texte du corps d'un .docx, dans l'ordre du document (B3).

    `source` : octets, chemin ou fichier binaire. Lève une exception si le
    fichier n'est pas un .docx lisible : l'appelant la transforme en
    avertissement (`extract_ex`) ou en texte vide (`lot_detector`)."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    controler_archive_office(source)   # répertoire et parties XML jaugés avant ouverture
    with zipfile.ZipFile(source) as archive:
        info = archive.getinfo(_partie_principale(archive))
        if info.file_size > _DOCX_PARTIE_MAX:
            raise ValueError(
                f"partie principale de {info.file_size // (1024 * 1024)} Mo une fois "
                f"décompressée (max {_DOCX_PARTIE_MAX // (1024 * 1024)} Mo)"
            )
        # zipfile s'arrête à la taille annoncée : un en-tête falsifié ne fait
        # pas lire davantage, il rend le flux illisible.
        with archive.open(info) as flux:
            return "\n".join(_lignes_corps(flux))


def _parseur():
    from lxml import etree

    return etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False)


def _partie_principale(archive: zipfile.ZipFile) -> str:
    """Nom de la partie principale, d'après _rels/.rels (repli : word/document.xml)."""
    from lxml import etree

    try:
        info = archive.getinfo("_rels/.rels")
    except KeyError:
        return "word/document.xml"
    if info.file_size > _DOCX_RELATIONS_MAX:
        return "word/document.xml"
    for relation in etree.fromstring(archive.read(info), parser=_parseur()):
        cible = relation.get("Target")
        if relation.get("Type") == _RELATION_DOCUMENT and cible:
            return posixpath.normpath(cible.lstrip("/"))
    return "word/document.xml"


def _lignes_corps(flux) -> Iterator[str]:
    from lxml import etree

    blocs = etree.iterparse(
        flux, events=("end",), tag=(_PARAGRAPHE, _TABLEAU, _CONTROLE, _XML_PERSO),
        resolve_entities=False, no_network=True, load_dtd=False,
    )
    for _evenement, bloc in blocs:
        parent = bloc.getparent()
        if parent is None or parent.tag != _CORPS:
            continue  # bloc imbriqué : lu avec le bloc du corps qui le contient
        yield from _lignes_bloc(bloc)
        # Bloc lu : libéré, avec tout ce qui le précède dans le corps.
        bloc.clear()
        while bloc.getprevious() is not None:
            del parent[0]


def _contenu(element) -> Iterator:
    """Enfants de `element`, enveloppes transparentes ouvertes (contrôle de
    contenu w:sdt → w:sdtContent, w:customXml)."""
    for enfant in element:
        if enfant.tag == _CONTROLE:
            corps = enfant.find(_CONTENU_CONTROLE)
            if corps is not None:
                yield from _contenu(corps)
        elif enfant.tag == _XML_PERSO:
            yield from _contenu(enfant)
        else:
            yield enfant


def _lignes_conteneur(conteneur) -> Iterator[str]:
    """Lignes d'une cellule, d'une zone de texte ou d'un contenu de contrôle."""
    for bloc in _contenu(conteneur):
        yield from _lignes_bloc(bloc)


def _lignes_bloc(bloc) -> Iterator[str]:
    if bloc.tag == _PARAGRAPHE:
        texte, zones = _texte_paragraphe(bloc)
        if texte.strip():
            yield texte
        for zone in zones:
            yield from _lignes_conteneur(zone)
    elif bloc.tag == _TABLEAU:
        yield from _lignes_tableau(bloc)
    elif bloc.tag == _CONTROLE:
        corps = bloc.find(_CONTENU_CONTROLE)
        if corps is not None:
            yield from _lignes_conteneur(corps)
    elif bloc.tag == _XML_PERSO:
        yield from _lignes_conteneur(bloc)


def _texte_paragraphe(paragraphe) -> Tuple[str, list]:
    """(texte affiché du paragraphe, zones de texte qu'il porte)."""
    morceaux: List[str] = []
    zones: list = []
    _parcourir(paragraphe, morceaux, zones)
    return "".join(morceaux), zones


def _parcourir(element, morceaux: List[str], zones: list) -> None:
    # Mêmes équivalences que python-docx pour les éléments d'un passage de texte.
    for enfant in element:
        balise = enfant.tag
        if balise == _TEXTE:
            if enfant.text:
                morceaux.append(enfant.text)
        elif balise in _TABULATIONS:
            morceaux.append("\t")
        elif balise == _SAUT:
            # saut de ligne ; un saut de page ou de colonne ne donne rien
            if enfant.get(_TYPE, "textWrapping") == "textWrapping":
                morceaux.append("\n")
        elif balise == _RETOUR:
            morceaux.append("\n")
        elif balise == _TRAIT_INSECABLE:
            morceaux.append("-")
        elif balise == _ZONE_DE_TEXTE:
            zones.append(enfant)
        elif not isinstance(balise, str) or balise in _SANS_TEXTE:
            continue  # commentaire XML, propriétés, texte supprimé…
        else:
            _parcourir(enfant, morceaux, zones)  # passage, lien, contrôle, insertion…


def _lignes_tableau(tableau) -> Iterator[str]:
    for ligne in _contenu(tableau):
        if ligne.tag != _LIGNE:
            continue
        cellules: List[List[str]] = []
        a_la_suite: List[str] = []  # tableaux imbriqués et zones de texte de la ligne
        for cellule in _contenu(ligne):
            if cellule.tag != _CELLULE or _suite_de_fusion(cellule):
                continue
            paragraphes: List[str] = []
            for bloc in _contenu(cellule):
                if bloc.tag == _PARAGRAPHE:
                    texte, zones = _texte_paragraphe(bloc)
                    if texte.strip():
                        paragraphes.append(texte.strip())
                    for zone in zones:
                        a_la_suite.extend(_lignes_conteneur(zone))
                elif bloc.tag == _TABLEAU:
                    a_la_suite.extend(_lignes_tableau(bloc))
            if paragraphes:
                cellules.append(paragraphes)
        if len(cellules) == 1:
            yield from cellules[0]  # encadré : ses paragraphes, un par ligne
        elif cellules:
            yield "\t".join(
                " ".join(paragraphes).replace("\t", " ").replace("\n", " ")
                for paragraphes in cellules
            )
        yield from a_la_suite


def _suite_de_fusion(cellule) -> bool:
    """Cellule qui prolonge une fusion : son texte est dans la première cellule."""
    proprietes = cellule.find(_PROPRIETES_CELLULE)
    if proprietes is None:
        return False
    for balise in (_FUSION_VERTICALE, _FUSION_HORIZONTALE):
        fusion = proprietes.find(balise)
        if fusion is not None and fusion.get(_VAL, "continue") == "continue":
            return True
    return False
