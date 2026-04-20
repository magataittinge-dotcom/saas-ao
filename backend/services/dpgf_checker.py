"""
DPGF verification: check that a filled DPGF spreadsheet is complete and consistent.
Supports xlsx (openpyxl), xls (xlrd), ods (odfpy).
"""

import logging
import re
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


def check_dpgf(file_path: Path) -> dict:
    """
    Verify a DPGF file and return a structured report.

    Returns:
        {
            "valid": bool,
            "warnings": list[str],
            "total_ht": float | None,
            "nb_lignes": int,
            "nb_lignes_remplies": int,
            "nb_lignes_vides": int,
        }
    """
    suffix = file_path.suffix.lower()

    try:
        if suffix in (".xlsx", ".xlsm"):
            return _check_xlsx(file_path)
        elif suffix == ".xls":
            return _check_xls(file_path)
        elif suffix == ".ods":
            return _check_ods(file_path)
        elif suffix == ".pdf":
            return {
                "valid": True,
                "warnings": ["Fichier PDF : vérification automatique limitée (pas de lecture des cellules)"],
                "total_ht": None,
                "nb_lignes": 0,
                "nb_lignes_remplies": 0,
                "nb_lignes_vides": 0,
            }
        else:
            return {
                "valid": False,
                "warnings": [f"Format non supporté : {suffix}"],
                "total_ht": None,
                "nb_lignes": 0,
                "nb_lignes_remplies": 0,
                "nb_lignes_vides": 0,
            }
    except Exception as e:
        logger.warning(f"DPGF check failed for {file_path}: {e}")
        return {
            "valid": False,
            "warnings": [f"Erreur de lecture du fichier : {e}"],
            "total_ht": None,
            "nb_lignes": 0,
            "nb_lignes_remplies": 0,
            "nb_lignes_vides": 0,
        }


# ── Price column detection heuristics ────────────────────────────────────────

_PRICE_HEADERS = re.compile(
    r"montant|total|prix\s*(unitaire|u\.?p\.?|ht|ttc)?|"
    r"p\.?u\.?\s*(ht)?|"
    r"sous.?total|"
    r"ht|"
    r"amount",
    re.IGNORECASE,
)

_SKIP_ROW = re.compile(
    r"^(total|sous.?total|tva|ttc|ht|remise|avance|r[eé]capitulat)",
    re.IGNORECASE,
)


def _is_number(val) -> bool:
    """Check if a cell value is a numeric price."""
    if val is None:
        return False
    if isinstance(val, (int, float)):
        return True
    if isinstance(val, str):
        clean = val.strip().replace(" ", "").replace(",", ".").replace("€", "").replace("\xa0", "")
        if not clean:
            return False
        try:
            float(clean)
            return True
        except ValueError:
            return False
    return False


def _to_float(val) -> Optional[float]:
    """Convert cell value to float."""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, str):
        clean = val.strip().replace(" ", "").replace(",", ".").replace("€", "").replace("\xa0", "")
        try:
            return float(clean)
        except ValueError:
            return None
    return None


def _analyze_rows(rows: list[list], sheet_name: str = "") -> dict:
    """
    Analyze rows from a spreadsheet sheet.
    rows: list of lists, each inner list = cell values for one row.
    """
    if not rows:
        return _empty_result()

    # Step 1: find header row and price column(s)
    header_row_idx = None
    price_cols: list[int] = []

    for i, row in enumerate(rows[:15]):  # search in first 15 rows
        for j, cell in enumerate(row):
            if cell and isinstance(cell, str) and _PRICE_HEADERS.search(cell):
                if header_row_idx is None:
                    header_row_idx = i
                price_cols.append(j)

    if header_row_idx is None:
        # Fallback: look for rightmost numeric columns
        return _fallback_check(rows, sheet_name)

    # Deduplicate and sort price columns
    price_cols = sorted(set(price_cols))

    # Use the LAST price column (usually "Montant total HT")
    # If multiple, prefer the rightmost one
    main_price_col = price_cols[-1] if price_cols else None

    if main_price_col is None:
        return _fallback_check(rows, sheet_name)

    # Step 2: iterate data rows (after header)
    warnings = []
    nb_lignes = 0
    nb_remplies = 0
    nb_vides = 0
    total_sum = 0.0
    found_total_row = None

    for i in range(header_row_idx + 1, len(rows)):
        row = rows[i]
        if not row or all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue  # skip fully empty rows

        # Get first text cell to check if it's a label
        first_text = ""
        for c in row:
            if c and isinstance(c, str) and c.strip():
                first_text = c.strip()
                break

        # Skip total/summary rows
        if first_text and _SKIP_ROW.match(first_text):
            val = row[main_price_col] if main_price_col < len(row) else None
            f = _to_float(val)
            if f is not None:
                found_total_row = f
            continue

        # This looks like a data row — check if it has a description
        has_description = bool(first_text)
        if not has_description:
            continue  # not a real line item

        nb_lignes += 1
        price_val = row[main_price_col] if main_price_col < len(row) else None

        if _is_number(price_val):
            nb_remplies += 1
            f = _to_float(price_val)
            if f is not None:
                total_sum += f
        else:
            nb_vides += 1
            prefix = f"[{sheet_name}] " if sheet_name else ""
            warnings.append(f"{prefix}Ligne {i + 1} : montant vide ({first_text[:50]})")

    # Step 3: check total consistency
    total_ht = round(total_sum, 2) if nb_remplies > 0 else None

    if found_total_row is not None and total_ht is not None:
        ecart = abs(found_total_row - total_ht)
        if ecart > 1.0:  # more than 1 euro difference
            warnings.append(
                f"Total incohérent : somme des lignes = {total_ht:.2f} €, "
                f"total affiché = {found_total_row:.2f} € (écart de {ecart:.2f} €)"
            )

    valid = nb_vides == 0 and len(warnings) == 0

    return {
        "valid": valid,
        "warnings": warnings[:20],  # cap warnings
        "total_ht": total_ht,
        "nb_lignes": nb_lignes,
        "nb_lignes_remplies": nb_remplies,
        "nb_lignes_vides": nb_vides,
    }


def _fallback_check(rows: list[list], sheet_name: str = "") -> dict:
    """Fallback when no price header found: just count rows with numbers."""
    nb_lignes = 0
    nb_remplies = 0
    total = 0.0

    for row in rows:
        if not row or all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue
        has_text = any(isinstance(c, str) and c.strip() for c in row)
        has_number = any(_is_number(c) for c in row)
        if has_text:
            nb_lignes += 1
            if has_number:
                nb_remplies += 1
                nums = [_to_float(c) for c in row if _is_number(c)]
                total += max(n for n in nums if n is not None) if nums else 0

    return {
        "valid": True,
        "warnings": [f"En-têtes de prix non détectés — vérification partielle"] if nb_lignes > 0 else [],
        "total_ht": round(total, 2) if nb_remplies > 0 else None,
        "nb_lignes": nb_lignes,
        "nb_lignes_remplies": nb_remplies,
        "nb_lignes_vides": nb_lignes - nb_remplies,
    }


def _empty_result() -> dict:
    return {
        "valid": False,
        "warnings": ["Fichier vide ou illisible"],
        "total_ht": None,
        "nb_lignes": 0,
        "nb_lignes_remplies": 0,
        "nb_lignes_vides": 0,
    }


# ── Format-specific readers ──────────────────────────────────────────────────

def _check_xlsx(file_path: Path) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
    try:
        results = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(list(row))
            if rows:
                results.append(_analyze_rows(rows, sheet_name))

        return _merge_results(results)
    finally:
        wb.close()


def _check_xls(file_path: Path) -> dict:
    import xlrd
    wb = xlrd.open_workbook(str(file_path))
    results = []
    for sheet_name in wb.sheet_names():
        ws = wb.sheet_by_name(sheet_name)
        rows = []
        for row_idx in range(ws.nrows):
            rows.append([ws.cell_value(row_idx, c) for c in range(ws.ncols)])
        if rows:
            results.append(_analyze_rows(rows, sheet_name))
    return _merge_results(results)


def _check_ods(file_path: Path) -> dict:
    from odf.opendocument import load as ods_load
    from odf.table import Table, TableRow, TableCell
    from odf.text import P

    doc = ods_load(str(file_path))
    sheets = doc.spreadsheet.getElementsByType(Table)
    results = []
    for sheet in sheets:
        sheet_name = sheet.getAttribute("name") or ""
        rows = []
        for row_el in sheet.getElementsByType(TableRow):
            cells = []
            for cell in row_el.getElementsByType(TableCell):
                ps = cell.getElementsByType(P)
                val = " ".join(str(p) for p in ps).strip() if ps else None
                if val:
                    # Try to convert to number
                    f = _to_float(val)
                    cells.append(f if f is not None else val)
                else:
                    cells.append(None)
            rows.append(cells)
        if rows:
            results.append(_analyze_rows(rows, sheet_name))
    return _merge_results(results)


def _merge_results(results: list[dict]) -> dict:
    """Merge results from multiple sheets into one report."""
    if not results:
        return _empty_result()

    if len(results) == 1:
        return results[0]

    # Merge across sheets
    warnings = []
    total_ht = 0.0
    nb_lignes = 0
    nb_remplies = 0
    nb_vides = 0
    any_total = False

    for r in results:
        warnings.extend(r["warnings"])
        nb_lignes += r["nb_lignes"]
        nb_remplies += r["nb_lignes_remplies"]
        nb_vides += r["nb_lignes_vides"]
        if r["total_ht"] is not None:
            total_ht += r["total_ht"]
            any_total = True

    valid = nb_vides == 0 and all(
        w.startswith("En-têtes") or "vérification partielle" in w
        for w in warnings
    )

    return {
        "valid": valid and nb_lignes > 0,
        "warnings": warnings[:20],
        "total_ht": round(total_ht, 2) if any_total else None,
        "nb_lignes": nb_lignes,
        "nb_lignes_remplies": nb_remplies,
        "nb_lignes_vides": nb_vides,
    }
