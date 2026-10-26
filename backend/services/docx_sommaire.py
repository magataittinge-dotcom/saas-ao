"""Sommaire du mémoire : champ TOC Word + résultat « en cache » rempli en deux passes.

Pourquoi deux passes : LibreOffice, sans interface, ne met pas à jour un champ
TOC créé par programme ; il imprime le résultat en cache tel quel. Et les
numéros de page ne se connaissent qu'une fois le document mis en page.

  passe 1 : le .docx porte un sommaire complet (titres connus à la
            construction) dont les numéros de page sont vides ; LibreOffice en
            fait un PDF ;
  entre   : on lit dans ce PDF la page de chaque titre (signets du PDF) ;
  passe 2 : on écrit ces numéros dans le .docx et LibreOffice refait le PDF.
            La mise en page ne bouge pas : seul un chiffre s'ajoute à des lignes
            déjà là.

Word, lui, recalcule le champ à l'ouverture (`w:updateFields`, champ « dirty »)
et à la touche F9. Le résultat en cache sert surtout au PDF.
"""
from __future__ import annotations

import io
import logging
import re
import unicodedata
import zipfile
from typing import Any

from docx.document import Document as _Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree

logger = logging.getLogger(__name__)

NIVEAU_MAX = 2                       # le sommaire reprend les titres 1 et 2
INSTRUCTION_TOC = f'TOC \\o "1-{NIVEAU_MAX}" \\h \\z \\u'
_LARGEUR_UTILE_CM = 17.0
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W = f"{{{W_NS}}}"


# ─── Construction (dans le document en cours d'écriture) ───────────────────────

def _style_toc(doc: _Document, niveau: int, police: str = "Calibri"):
    nom = f"TOC {niveau}"
    try:
        style = doc.styles[nom]
    except KeyError:
        style = doc.styles.add_style(nom, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), police)
    pf = style.paragraph_format
    pf.space_after = Pt(2)
    pf.space_before = Pt(6 if niveau == 1 else 0)
    pf.left_indent = Cm(0 if niveau == 1 else 0.6)
    # numéro de page collé à droite, avec des points de conduite
    pf.tab_stops.add_tab_stop(Cm(_LARGEUR_UTILE_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    style.font.bold = niveau == 1
    return style


def _titres(doc: _Document, apres) -> list[tuple[int, str, Any]]:
    """(niveau, texte, paragraphe) des titres 1 et 2 placés après `apres`."""
    trouves: list[tuple[int, str, Any]] = []
    vu = False
    for p in doc.paragraphs:
        if p._p is apres:
            vu = True
            continue
        if not vu:
            continue
        nom = p.style.name if p.style is not None else ""
        m = re.fullmatch(r"Heading ([1-9])", nom)
        if m and int(m.group(1)) <= NIVEAU_MAX and p.text.strip():
            trouves.append((int(m.group(1)), p.text.strip(), p))
    return trouves


def _fld(par, type_: str, *, sale: bool = False):
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), type_)
    if sale:
        e.set(qn("w:dirty"), "true")      # Word recalcule ce champ à l'ouverture
    par.add_run()._r.append(e)


def inserer_sommaire(doc: _Document, ancre) -> bool:
    """Remplace le paragraphe `ancre` (posé après la couverture) par le
    sommaire, construit d'après les titres 1 et 2 écrits après lui. Renvoie
    False (et retire l'ancre) s'il n'y a aucun titre."""
    titres = _titres(doc, ancre)
    if not titres:
        ancre.getparent().remove(ancre)
        return False
    _style_toc(doc, 1)
    _style_toc(doc, 2)

    # La couverture est suivie d'une page de sommaire, puis des parties : un
    # saut de page avant le titre du sommaire et avant le premier titre.
    titre_toc = doc.add_paragraph()
    titre_toc.paragraph_format.page_break_before = True
    titre_toc.paragraph_format.space_after = Pt(12)
    r = titre_toc.add_run("Sommaire")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
    titres[0][2].paragraph_format.page_break_before = True

    entrees = []
    for k, (niveau, texte, _) in enumerate(titres):
        par = doc.add_paragraph(style=f"TOC {niveau}")
        if k == 0:
            _fld(par, "begin", sale=True)
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = f" {INSTRUCTION_TOC} "
            par.add_run()._r.append(instr)
            _fld(par, "separate")
        par.add_run(texte)
        par.add_run("\t")
        par.add_run("")                   # numéro de page : écrit à la passe 2
        if k == len(titres) - 1:
            _fld(par, "end")
        entrees.append(par._p)

    for el in [titre_toc._p, *entrees]:   # ces paragraphes ont été ajoutés en fin de corps
        ancre.addprevious(el)
    ancre.getparent().remove(ancre)

    # Word : mise à jour des champs à l'ouverture (avant <w:compat>, ordre du schéma)
    reglages = doc.settings.element
    if reglages.find(qn("w:updateFields")) is None:
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        for nom in ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:rsids"):
            suivant = reglages.find(qn(nom))
            if suivant is not None:
                suivant.addprevious(upd)
                break
        else:
            reglages.append(upd)
    return True


# ─── Passe 1 → passe 2 ─────────────────────────────────────────────────────────

def _normaliser(texte: str) -> str:
    texte = unicodedata.normalize("NFKC", texte)
    return re.sub(r"\s+", " ", texte).strip().casefold()


def _lire_document_xml(docx_bytes: bytes) -> etree._Element:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        return etree.fromstring(z.read("word/document.xml"))


def _texte(p: etree._Element) -> str:
    return "".join(t.text or "" for t in p.iter(f"{_W}t"))


def _style_id(p: etree._Element) -> str:
    st = p.find(f"{_W}pPr/{_W}pStyle")
    return st.get(f"{_W}val", "") if st is not None else ""


def titres_du_document(docx_bytes: bytes) -> list[tuple[int, str]]:
    """Tous les titres (niveau, texte) du corps, dans l'ordre."""
    racine = _lire_document_xml(docx_bytes)
    titres = []
    for p in racine.iter(f"{_W}p"):
        m = re.fullmatch(r"Heading([1-9])", _style_id(p))
        if m and _texte(p).strip():
            titres.append((int(m.group(1)), _texte(p).strip()))
    return titres


def _signets_pdf(pdf_bytes: bytes) -> list[tuple[str, int]]:
    """(titre, numéro de page à partir de 1) de chaque signet, à plat, dans l'ordre."""
    from pypdf import PdfReader

    lecteur = PdfReader(io.BytesIO(pdf_bytes))
    resultat: list[tuple[str, int]] = []

    def parcourir(noeuds):
        for n in noeuds:
            if isinstance(n, list):
                parcourir(n)
            else:
                resultat.append((str(n.title), lecteur.get_destination_page_number(n) + 1))

    parcourir(lecteur.outline)
    return resultat


def pages_des_titres(pdf_bytes: bytes, titres: list[tuple[int, str]]) -> list[int | None]:
    """Page de chaque titre dans le PDF de la passe 1, par ses signets. Un titre
    sans signet reçoit None."""
    signets = _signets_pdf(pdf_bytes)
    pages: list[int | None] = []
    curseur = 0
    for _, texte in titres:
        cible = _normaliser(texte)
        trouve = None
        for k in range(curseur, len(signets)):
            if _normaliser(signets[k][0]) == cible:
                trouve = signets[k][1]
                curseur = k + 1
                break
        pages.append(trouve)
    return pages


def renseigner_sommaire(docx_bytes: bytes, pages: dict[int, int]) -> bytes:
    """Écrit les numéros de page dans les lignes du sommaire.
    `pages` : rang de l'entrée du sommaire (0, 1, 2…) → numéro de page."""
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        racine = etree.fromstring(z.read("word/document.xml"))
        rang = 0
        for p in racine.iter(f"{_W}p"):
            if not re.fullmatch(r"TOC[1-9]", _style_id(p)):
                continue
            runs = p.findall(f"{_W}r")
            for k, run in enumerate(runs):
                if run.find(f"{_W}tab") is not None and k + 1 < len(runs):
                    page_run = runs[k + 1]
                    for t in page_run.findall(f"{_W}t"):
                        page_run.remove(t)
                    if rang in pages:
                        t = etree.SubElement(page_run, f"{_W}t")
                        t.text = str(pages[rang])
                    break
            rang += 1
        nouveau = etree.tostring(racine, xml_declaration=True, encoding="UTF-8", standalone=True)
        sortie = io.BytesIO()
        with zipfile.ZipFile(sortie, "w", zipfile.ZIP_DEFLATED) as out:
            for info in z.infolist():
                donnees = nouveau if info.filename == "word/document.xml" else z.read(info.filename)
                out.writestr(info, donnees)
    return sortie.getvalue()


def completer_sommaire(docx_bytes: bytes, pdf_bytes: bytes) -> bytes | None:
    """.docx de la passe 2 (numéros de page écrits) d'après le PDF de la passe 1,
    ou None s'il n'y a pas de sommaire à remplir ou si les pages sont
    introuvables."""
    try:
        racine = _lire_document_xml(docx_bytes)
        entrees_toc = [p for p in racine.iter(f"{_W}p") if re.fullmatch(r"TOC[1-9]", _style_id(p))]
        if not entrees_toc:
            return None
        titres = titres_du_document(docx_bytes)
        pages_titres = pages_des_titres(pdf_bytes, titres)
        # le sommaire reprend les titres de niveau <= NIVEAU_MAX, dans l'ordre
        pages = {}
        rang = 0
        for (niveau, _), page in zip(titres, pages_titres):
            if niveau > NIVEAU_MAX:
                continue
            if page is not None:
                pages[rang] = page
            rang += 1
        if not pages:
            return None
        return renseigner_sommaire(docx_bytes, pages)
    except Exception:  # noqa: BLE001 — jamais bloquer l'export : le PDF 1 reste valable
        logger.warning("Sommaire : numéros de page non calculés", exc_info=True)
        return None
