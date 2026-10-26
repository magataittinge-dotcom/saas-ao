"""Convertit le content_json d'un mémoire technique en fichier .docx formaté."""
from __future__ import annotations

import io
import logging
import re
from contextvars import ContextVar
from dataclasses import dataclass, field
from datetime import date
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

from services import memoire_markdown as md
from services.docx_sommaire import inserer_sommaire

# ─── Passages « À compléter » ──────────────────────────────────────────────────
#
# Un « [À COMPLÉTER : X] » restant est une donnée d'entreprise que l'IA n'avait
# pas : il ne part jamais en silence ni entre crochets, il devient un passage
# surligné « À compléter : X » et il est compté (l'écran dira « N passages à
# compléter avant le dépôt »). Même traitement pour un « [📍 … ] » d'un ancien
# mémoire, dont le prompt demandait une consigne de carte.

_A_COMPLETER_RE = re.compile(
    r"\[\s*[ÀA]\s*COMPL[ÉE]TER(?:\s+PAR\s+L['’]ENTREPRISE)?\s*(?::\s*([^\]]{0,300}))?\]", re.IGNORECASE,
)
_PIN_RE = re.compile(r"\[📍[^\]]{0,600}\]")
_DEBUT_PASSAGE, _FIN_PASSAGE = "\ue000", "\ue001"   # balises privées, jamais imprimées
_PASSAGE_RE = re.compile(f"{_DEBUT_PASSAGE}([^{_FIN_PASSAGE}]*){_FIN_PASSAGE}")


@dataclass(frozen=True)
class PassageACompleter:
    section: str    # titre de la section où il se trouve
    libelle: str    # ce qui est attendu ; vide si l'IA n'a rien précisé


@dataclass
class ExportDocx:
    octets: bytes
    a_completer: list[PassageACompleter] = field(default_factory=list)

    @property
    def nombre(self) -> int:
        return len(self.a_completer)


_COLLECTE: ContextVar[list[PassageACompleter] | None] = ContextVar("docx_a_completer", default=None)
_SECTION: ContextVar[str] = ContextVar("docx_section", default="")


def _libelle_pin(texte: str) -> str:
    """« [📍 INSÉREZ ICI UNE CARTE DE LOCALISATION — … ] » → « carte de localisation »."""
    avant_tiret = texte.strip("[]").split("—")[0]
    nettoye = re.sub(r"^\W*(?:ins[ée]rez ici\s+)?(?:une?\s+)?", "", avant_tiret, flags=re.IGNORECASE)
    return nettoye.strip().lower()


def _marquer_passages(texte: str, *, retirer_marqueur_carte: bool = False) -> str:
    """Remplace chaque passage à compléter par une balise privée portant le
    libellé, et l'enregistre dans l'ordre du document. `retirer_marqueur_carte` :
    la carte est dessinée, l'ancienne consigne « [📍 … ] » n'a plus lieu d'être."""
    if retirer_marqueur_carte:
        texte = _PIN_RE.sub("", texte)
    collecte = _COLLECTE.get()

    def enregistrer(libelle: str) -> str:
        libelle = " ".join(libelle.split())
        if collecte is not None:
            collecte.append(PassageACompleter(_SECTION.get(), libelle))
        return f"{_DEBUT_PASSAGE}{libelle}{_FIN_PASSAGE}"

    # un seul balayage, dans l'ordre du texte
    def remplacer(m: re.Match[str]) -> str:
        if m.group(0).startswith("[📍"):
            return enregistrer(_libelle_pin(m.group(0)))
        return enregistrer(m.group(1) or "")

    return re.sub(f"{_PIN_RE.pattern}|{_A_COMPLETER_RE.pattern}", remplacer, texte, flags=re.IGNORECASE)

logger = logging.getLogger(__name__)


# ─── Couleurs ──────────────────────────────────────────────────────────────────
_BLUE_H1  = RGBColor(0x1E, 0x40, 0xAF)
_BLUE_H2  = RGBColor(0x1D, 0x4E, 0xD8)
_GREY_TXT = RGBColor(0x1F, 0x2D, 0x3D)
_GREY_META= RGBColor(0x6B, 0x72, 0x80)
_GREY_PG  = RGBColor(0x9C, 0xA3, 0xAF)

_PART_A: list[tuple[str, str]] = [
    ("implantation",          "1. Implantation géographique"),
    ("historique",            "2. Historique"),
    ("engagement_qualitatif", "3. Engagement qualitatif"),
    ("activites",             "4. Nos activités"),
    ("organigramme",          "5. Organigramme"),
    ("roles_missions",        "6. Rôles et missions de l'équipe d'encadrement"),
    ("moyens_informatiques",  "7. Moyens informatiques"),
    ("vehicules",             "8. Véhicules"),
    ("materiel",              "9. Matériel"),
    ("references",            "10. Références chantiers"),
    ("fournisseurs",          "11. Fournisseurs"),
]
_PART_B: list[tuple[str, str]] = [
    ("demarrage",        "1. Démarrage du chantier"),
    ("interlocuteur",    "2. Interlocuteur dédié"),
    ("qualite_ouvrages", "3. Qualité des ouvrages"),
    ("respect_planning", "4. Respect du planning"),
    ("securite",         "5. Dispositions relatives à la sécurité"),
    ("dechets",          "6. Traitement des déchets"),
    ("environnement",    "7. Environnement"),
]
_PART_C: list[tuple[str, str]] = [
    ("methodologie",             "1. Méthodologie détaillée"),
    ("effectifs",                "2. Effectifs dédiés au chantier"),
    ("materiels",                "3. Matériels dédiés"),
    ("hygiene_securite",         "4. Hygiène et sécurité"),
    ("mesures_environnementales","5. Mesures environnementales"),
    ("gpa",                      "6. Garantie de parfait achèvement (GPA)"),
    ("delai",                    "7. Délai de travaux"),
]


# ─── Date en français, sans dépendre de la locale du système ───────────────────

_MOIS_FR = ("janvier", "février", "mars", "avril", "mai", "juin", "juillet",
            "août", "septembre", "octobre", "novembre", "décembre")


def _date_fr(jour: date) -> str:
    """« octobre 2026 ». `strftime("%B")` suit la locale du processus (anglais
    sur un serveur sans locale française) : on n'y touche pas."""
    return f"{_MOIS_FR[jour.month - 1]} {jour.year}"


def _aujourdhui() -> date:
    return date.today()


# ─── Helpers document ──────────────────────────────────────────────────────────

def _champ(para, instruction: str, resultat: str = "1", *, taille: float = 9,
           couleur: RGBColor = _GREY_PG):
    """Champ Word complet (begin / instruction / separate / résultat en cache /
    end), un run par étape, mis en forme comme le texte voisin : le résultat en
    cache s'affiche tel quel dans LibreOffice, Word le recalcule."""
    def run(enfant=None, texte: str | None = None):
        r = para.add_run(texte) if texte is not None else para.add_run()
        r.font.size = Pt(taille)
        r.font.color.rgb = couleur
        if enfant is not None:
            r._r.append(enfant)
        return r

    debut = OxmlElement("w:fldChar")
    debut.set(qn("w:fldCharType"), "begin")
    run(debut)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run(sep)
    run(texte=resultat)
    fin = OxmlElement("w:fldChar")
    fin.set(qn("w:fldCharType"), "end")
    run(fin)


def _effacer_tabulations_du_style(para):
    """Les styles En-tête / Pied de page du modèle portent des tabulations
    (centre 8,25 cm, droite 16,5 cm) qui décalent tout texte après une
    tabulation : on les efface sur le paragraphe."""
    ts = para.paragraph_format.tab_stops
    ts.add_tab_stop(Twips(4680), WD_TAB_ALIGNMENT.CLEAR)
    ts.add_tab_stop(Twips(9360), WD_TAB_ALIGNMENT.CLEAR)


def _add_footer_page_numbers(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    para = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _effacer_tabulations_du_style(para)

    r_before = para.add_run("Page ")
    r_before.font.size = Pt(9)
    r_before.font.color.rgb = _GREY_PG
    _champ(para, "PAGE")
    r_sep = para.add_run(" / ")
    r_sep.font.size = Pt(9)
    r_sep.font.color.rgb = _GREY_PG
    _champ(para, "NUMPAGES")


def _setup_doc(org_name: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)   # A4 (le modèle python-docx est en Letter)
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.0)
    section.right_margin  = Cm(2.0)

    # Couverture sans en-tête ni pied : première page distincte, laissée vide.
    # Les en-têtes / pieds ci-dessous valent pour les pages suivantes.
    section.different_first_page_header_footer = True

    # Header : nom entreprise à droite
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _effacer_tabulations_du_style(hp)
    hr = hp.add_run(org_name)
    hr.font.name = "Calibri"
    hr.font.size = Pt(9)
    hr.font.color.rgb = _GREY_META

    _add_footer_page_numbers(section)

    _styles(doc)
    # Pas de police embarquée : le Word déclare Calibri (présente dans tout Word). Le PDF
    # sort en Carlito grâce à `pdf_export`, qui la montre à LibreOffice.
    return doc


def _police_style(style, police: str = "Calibri"):
    """Police posée sur les 4 attributs (ascii, hAnsi, cs, eastAsia) : sinon
    Word et LibreOffice retombent sur la police du thème (Cambria/Calibri) pour
    les titres."""
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), police)


def _styles(doc: Document):
    """Styles Word natifs : le patron retouche dans Word sans casser la mise en
    page (modifier « Titre 2 » change tous les titres 2)."""
    st = doc.styles
    normal = st["Normal"]
    normal.font.size = Pt(11)
    normal.font.color.rgb = _GREY_TXT
    normal.paragraph_format.space_after = Pt(4)
    _police_style(normal)
    langue = OxmlElement("w:lang")
    langue.set(qn("w:val"), "fr-FR")
    normal.element.get_or_add_rPr().append(langue)

    for nom, taille, couleur, avant, apres in (
        ("Heading 1", 16, _BLUE_H1, 20, 8),
        ("Heading 2", 13, _BLUE_H2, 12, 4),
        ("Heading 3", 11.5, _GREY_TXT, 10, 3),
        ("Heading 4", 11, _GREY_META, 8, 2),
        ("Heading 5", 11, _GREY_META, 6, 2),
        ("Heading 6", 11, _GREY_META, 6, 2),
    ):
        h = st[nom]
        h.font.size = Pt(taille)
        h.font.bold = True
        h.font.italic = nom in ("Heading 5", "Heading 6")
        h.font.color.rgb = couleur
        _police_style(h)
        h.paragraph_format.space_before = Pt(avant)
        h.paragraph_format.space_after = Pt(apres)
        h.paragraph_format.keep_with_next = True

    for nom in ("List Bullet", "List Bullet 2", "List Number", "List Number 2", "Caption"):
        _police_style(st[nom])
    for nom in ("List Bullet", "List Bullet 2", "List Number", "List Number 2"):
        st[nom].paragraph_format.space_after = Pt(2)


def _centered_line(doc: Document, text: str, *, size: int, color, bold: bool = False,
                   space_after: int = 6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def _cover_page(
    doc: Document,
    project_name: str,
    org_name: str,
    *,
    logo_image: bytes | None = None,
    lot_name: str | None = None,
    maitre_ouvrage: str | None = None,
    org_address: str | None = None,
    org_siret: str | None = None,
):
    """C8b — page de garde sobre : logo (si présent), marché, lot, MOA,
    date, coordonnées entreprise. Sans logo → texte propre, jamais de
    placeholder d'image cassé."""
    for _ in range(2):
        doc.add_paragraph()

    if logo_image:
        try:
            p_logo = doc.add_paragraph()
            p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_logo.add_run().add_picture(io.BytesIO(logo_image), width=Cm(4.5))
            p_logo.paragraph_format.space_after = Pt(18)
        except Exception:
            # Logo illisible → page de garde texte propre, pas d'image cassée.
            pass

    _centered_line(doc, "MÉMOIRE TECHNIQUE", size=13, color=_GREY_META, space_after=14)
    _centered_line(doc, project_name.upper(), size=24, color=_BLUE_H1, bold=True, space_after=10)
    if lot_name:
        _centered_line(doc, lot_name, size=14, color=RGBColor(0x03, 0x69, 0xA1), space_after=8)
    if maitre_ouvrage:
        _centered_line(
            doc, f"Maître d'ouvrage : {maitre_ouvrage}", size=11,
            color=RGBColor(0x47, 0x55, 0x69), space_after=16,
        )

    _centered_line(doc, org_name, size=13, color=RGBColor(0x37, 0x41, 0x51))
    if org_address:
        _centered_line(doc, org_address, size=10, color=_GREY_PG, space_after=2)
    if org_siret:
        _centered_line(doc, f"SIRET {org_siret}", size=10, color=_GREY_PG, space_after=2)

    _centered_line(doc, _date_fr(_aujourdhui()), size=11, color=_GREY_PG)

    # Pas de saut de page ici : le sommaire (ou, à défaut, le premier titre)
    # démarre sa page par `page_break_before`, sans paragraphe vide en tête.


def _h1(doc: Document, text: str):
    doc.add_heading(text, level=1)


def _h2(doc: Document, text: str):
    _SECTION.set(text)
    doc.add_heading(text, level=2)


# ─── Corps : Markdown de l'IA → styles Word natifs ─────────────────────────────

_LARGEUR_UTILE_CM = 17.0   # A4 21 cm - 2 × 2 cm de marge
_TEXTE_TABLEAU_PT = 9.5


def _ajouter_segments(p, segments: list[md.Segment], *, taille: float | None = None,
                      couleur: RGBColor | None = None, gras: bool = False):
    for seg in segments:
        position = 0
        morceaux: list[tuple[str, bool]] = []
        for m in _PASSAGE_RE.finditer(seg.texte):
            if m.start() > position:
                morceaux.append((seg.texte[position:m.start()], False))
            libelle = m.group(1)
            morceaux.append((f"À compléter : {libelle}" if libelle else "À compléter", True))
            position = m.end()
        if position < len(seg.texte):
            morceaux.append((seg.texte[position:], False))
        for texte, a_completer in morceaux:
            texte = texte.replace(_DEBUT_PASSAGE, "").replace(_FIN_PASSAGE, "")
            if not texte:
                continue
            r = p.add_run(texte)
            if seg.gras or gras:
                r.bold = True
            if seg.italique:
                r.italic = True
            if taille is not None:
                r.font.size = Pt(taille)
            if a_completer:
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                r.font.color.rgb = _GREY_TXT      # lisible sur le surlignage, même en en-tête de tableau
            elif couleur is not None:
                r.font.color.rgb = couleur


def _nouvelle_numerotation(doc: Document) -> int | None:
    """Crée une instance de numérotation qui repart à 1, branchée sur la
    définition abstraite du style « List Number ». Sans cela, toutes les listes
    numérotées du document se suivent (1 à 3, puis 4 à 6…)."""
    try:
        style_ppr = doc.styles["List Number"].element.pPr
        num_id_style = style_ppr.numPr.numId.val
        numbering = doc.part.numbering_part.element
        abstract_id = None
        for num in numbering.findall(qn("w:num")):
            if int(num.get(qn("w:numId"))) == int(num_id_style):
                abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        if abstract_id is None:
            return None
        nouvel_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(nouvel_id))
        ref = OxmlElement("w:abstractNumId")
        ref.set(qn("w:val"), abstract_id)
        num.append(ref)
        override = OxmlElement("w:lvlOverride")
        override.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:startOverride")
        start.set(qn("w:val"), "1")
        override.append(start)
        num.append(override)
        numbering.append(num)
        return nouvel_id
    except (KeyError, AttributeError, ValueError):
        return None


def _rendre_liste(doc: Document, liste: md.Liste):
    num_id = _nouvelle_numerotation(doc) if liste.numerotee else None
    base = "List Number" if liste.numerotee else "List Bullet"
    for el in liste.elements:
        style = f"{base} 2" if el.profondeur else base
        p = doc.add_paragraph(style=style)
        if num_id is not None and not el.profondeur:
            p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = num_id
            p._p.pPr.numPr.get_or_add_ilvl().val = 0
        _ajouter_segments(p, el.segments)


def _largeurs_colonnes(tableau: md.Tableau) -> list[float]:
    """Largeurs (cm) proportionnelles au contenu, bornées : ni colonne
    écrasée, ni colonne qui mange la page."""
    n = len(tableau.entetes)
    poids = []
    for j in range(n):
        textes = [tableau.entetes[j]] + [ligne[j] for ligne in tableau.lignes]
        longueur = max((sum(len(s.texte) for s in segs) for segs in textes), default=1)
        poids.append(min(max(longueur, 6), 38))
    total = sum(poids)
    return [_LARGEUR_UTILE_CM * w / total for w in poids]


def _ombrer(cellule, couleur_hex: str):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), couleur_hex)
    cellule._tc.get_or_add_tcPr().append(shd)


def _rendre_tableau(doc: Document, tableau: md.Tableau):
    """Vrai tableau Word : en-tête répété sur chaque page, lignes non coupées."""
    n = len(tableau.entetes)
    if n == 0:
        return
    largeurs = _largeurs_colonnes(tableau)
    t = doc.add_table(rows=1, cols=n)
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tbl_pr = t._tbl.tblPr
    marges = OxmlElement("w:tblCellMar")
    for cote, valeur in (("top", 50), ("left", 90), ("bottom", 50), ("right", 90)):
        e = OxmlElement(f"w:{cote}")
        e.set(qn("w:w"), str(valeur))
        e.set(qn("w:type"), "dxa")
        marges.append(e)
    tbl_pr.append(marges)
    bordures = OxmlElement("w:tblBorders")
    for cote in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{cote}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), "D1D5DB")
        bordures.append(e)
    tbl_pr.append(bordures)
    for col, largeur in zip(t.columns, largeurs):
        col.width = Cm(largeur)

    def remplir(cellule, segments, largeur, *, entete=False):
        cellule.width = Cm(largeur)
        p = cellule.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        _ajouter_segments(
            p, segments, taille=_TEXTE_TABLEAU_PT, gras=entete,
            couleur=RGBColor(0xFF, 0xFF, 0xFF) if entete else None,
        )

    for j, segs in enumerate(tableau.entetes):
        cellule = t.rows[0].cells[j]
        remplir(cellule, segs, largeurs[j], entete=True)
        _ombrer(cellule, "1E40AF")
    ligne0 = t.rows[0]._tr.get_or_add_trPr()
    ligne0.append(OxmlElement("w:cantSplit"))
    ligne0.append(OxmlElement("w:tblHeader"))
    for k, ligne in enumerate(tableau.lignes):
        row = t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for j, segs in enumerate(ligne):
            remplir(row.cells[j], segs, largeurs[j])
            if k % 2 == 1:
                _ombrer(row.cells[j], "F3F4F6")
    # respiration après le tableau
    espace = doc.add_paragraph()
    espace.paragraph_format.space_after = Pt(2)


def _body(doc: Document, text: str, *, skip_map_placeholder: bool = False, niveau_base: int = 2):
    """Convertit le Markdown d'une section en styles Word natifs (titres, listes
    à puces et numérotées, gras, italique, tableaux). `niveau_base` : niveau du
    titre qui précède la section ; un `#` de l'IA descend d'un cran dessous."""
    text = _marquer_passages(text or "", retirer_marqueur_carte=skip_map_placeholder)
    for bloc in md.analyser(text):
        if isinstance(bloc, md.Titre):
            niveau = min(niveau_base + bloc.niveau, 6)
            p = doc.add_heading("", level=niveau)
            _ajouter_segments(p, bloc.segments)
        elif isinstance(bloc, md.Liste):
            _rendre_liste(doc, bloc)
        elif isinstance(bloc, md.Tableau):
            _rendre_tableau(doc, bloc)
        else:
            p = doc.add_paragraph()
            _ajouter_segments(p, bloc.segments)


def _add_map_image(doc: Document, map_image: bytes, caption: str):
    """Insert a map PNG image with a caption below it."""
    img_stream = io.BytesIO(map_image)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p_img.add_run()
    run.add_picture(img_stream, width=Cm(15))
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Calibri"
    r_cap.font.size = Pt(9)
    r_cap.font.italic = True
    r_cap.font.color.rgb = _GREY_META
    p_cap.paragraph_format.space_after = Pt(12)


def _ajouter_visuel(doc: Document, svg: str, png: bytes, largeur_cm: float,
                    texte_alt: str, legende: str | None = None, paysage: bool = False):
    """Pose un visuel (SVG vectoriel + PNG de repli + texte alternatif) et sa légende.
    `paysage` : le visuel occupe une page A4 paysage, puis le document revient en portrait."""
    from services.visuels.word import inserer_svg_png

    if paysage:
        _changer_orientation(doc, paysage=True)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = bool(legende)
    inserer_svg_png(p_img, svg, png, largeur_cm, texte_alt)
    if legende:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run(legende)
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.color.rgb = _GREY_META
        p_cap.paragraph_format.space_after = Pt(12)
    if paysage:
        _changer_orientation(doc, paysage=False)


def _changer_orientation(doc: Document, *, paysage: bool):
    """Nouvelle section (nouvelle page) en A4 paysage ou portrait."""
    from docx.enum.section import WD_ORIENT, WD_SECTION

    s = doc.add_section(WD_SECTION.NEW_PAGE)
    portrait_l, portrait_h = Cm(21.0), Cm(29.7)
    s.orientation = WD_ORIENT.LANDSCAPE if paysage else WD_ORIENT.PORTRAIT
    s.page_width, s.page_height = (portrait_h, portrait_l) if paysage else (portrait_l, portrait_h)
    s.different_first_page_header_footer = False   # pas de « première page » sur ces sections


def _add_part(
    doc: Document,
    title: str,
    part: dict[str, Any],
    sections: list[tuple[str, str]],
    *,
    part_key: str = "",
    map_image: bytes | None = None,
    map_caption: str | None = None,
    visuels: list | None = None,
):
    from services.visuels.modele import retirer_gabarits

    _h1(doc, title)
    for key, label in sections:
        text = (part or {}).get(key, "")
        ses_visuels = [v for v in (visuels or []) if v.section == (part_key, key)]
        if (text and text.strip()) or ses_visuels:
            _h2(doc, label)
            text = retirer_gabarits(text or "", ses_visuels)
            if key == "implantation" and map_image:
                _body(doc, text, skip_map_placeholder=True)
                _add_map_image(doc, map_image, map_caption or "")
            else:
                _body(doc, text)
            for v in ses_visuels:
                if v.tableau:
                    from services.visuels.word import ajouter_tableau
                    ajouter_tableau(doc, v.tableau, v.legende)
                    continue
                try:
                    png = v.png          # un visuel qui ne se rastérise pas disparaît, l'export continue
                except Exception:
                    logger.warning("Visuel %s omis : PNG de repli impossible", v.section, exc_info=True)
                    continue
                _ajouter_visuel(doc, v.svg, png, v.largeur_cm, v.texte_alt, v.legende, v.paysage)


# ─── Point d'entrée public ─────────────────────────────────────────────────────

def build_memoire_docx(
    content_json: dict[str, Any],
    project_name: str,
    org_name: str,
    *,
    map_image: bytes | None = None,
    map_caption: str | None = None,
    visuels: list | None = None,
    logo_image: bytes | None = None,
    lot_name: str | None = None,
    maitre_ouvrage: str | None = None,
    org_address: str | None = None,
    org_siret: str | None = None,
) -> bytes:
    """Retourne les bytes du .docx prêt à être streamé en réponse HTTP."""
    doc = _setup_doc(org_name)
    _cover_page(
        doc, project_name, org_name,
        logo_image=logo_image, lot_name=lot_name, maitre_ouvrage=maitre_ouvrage,
        org_address=org_address, org_siret=org_siret,
    )

    ancre_sommaire = doc.add_paragraph()   # le sommaire prendra cette place, une fois les titres écrits

    preambule = content_json.get("preambule", "")
    if preambule and preambule.strip():
        _SECTION.set("Préambule")
        _h1(doc, "PRÉAMBULE")
        _body(doc, preambule, niveau_base=1)

    _add_part(doc, "PARTIE A — PRÉSENTATION GÉNÉRALE",       content_json.get("partie_a", {}), _PART_A, part_key="partie_a", map_image=map_image, map_caption=map_caption, visuels=visuels)
    _add_part(doc, "PARTIE B — PRÉSENTATION DE LA PRESTATION", content_json.get("partie_b", {}), _PART_B, part_key="partie_b", visuels=visuels)
    _add_part(doc, "PARTIE C — MÉTHODOLOGIE MISE EN ŒUVRE",   content_json.get("partie_c", {}), _PART_C, part_key="partie_c", visuels=visuels)

    inserer_sommaire(doc, ancre_sommaire._p)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_memoire_docx_ex(*args: Any, **kwargs: Any) -> ExportDocx:
    """Comme `build_memoire_docx` (mêmes paramètres), mais renvoie aussi les
    passages « À compléter » du document (nombre et libellés, dans l'ordre),
    pour que l'écran dise « N passages à compléter avant le dépôt »."""
    collecte: list[PassageACompleter] = []
    jeton_collecte = _COLLECTE.set(collecte)
    jeton_section = _SECTION.set("")
    try:
        octets = build_memoire_docx(*args, **kwargs)
    finally:
        _COLLECTE.reset(jeton_collecte)
        _SECTION.reset(jeton_section)
    return ExportDocx(octets, collecte)


# ─── Compatibilité avec routers/export.py ──────────────────────────────────────

class DocxExporter:
    """Façade conservée pour la rétrocompatibilité avec routers/export.py."""

    def export(self, content_json: dict, organization) -> bytes:
        org_name = organization.name if organization else "Entreprise"
        project_name = content_json.get("_project_name", "Projet")
        return build_memoire_docx(content_json, project_name, org_name)
